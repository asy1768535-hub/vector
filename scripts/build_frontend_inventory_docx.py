from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "48-frontend-button-function-inventory.md"
SCREENSHOT = ROOT / "docs" / "48-frontend-console-login.png"
OUTPUT = ROOT / "docs" / "48-frontend-button-function-inventory.docx"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
TABLE_CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}
BODY_FONT = "Calibri"
EAST_ASIA_FONT = "Microsoft YaHei"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 48, 58)
MUTED = RGBColor(92, 105, 113)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "EAF5F1"


def set_run_font(run, name=BODY_FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, margins=TABLE_CELL_MARGINS):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in margins.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="C8D0D8", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_num_definition(document, kind):
    numbering = document.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_sep, text, fld_char_end])


def set_page_furniture(document):
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("向量知识库  ·  前端功能清单")
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), 2)
    header.add_run("\t发布候选验收")
    set_run_font(header.runs[-1], size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    footer_run = footer.add_run("第 ")
    set_run_font(footer_run, size=8.5, color=MUTED)
    add_field(footer, "PAGE")
    set_run_font(footer.runs[-1], size=8.5, color=MUTED)
    tail = footer.add_run(" 页")
    set_run_font(tail, size=8.5, color=MUTED)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True


def add_inline_runs(paragraph, text, size=11, color=INK):
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    token_re = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|__[^_]+__)")
    position = 0
    for match in token_re.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=max(size - 0.5, 8), color=DARK_BLUE)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=color)


def add_body_paragraph(document, text, style=None):
    paragraph = document.add_paragraph(style=style or "Normal")
    add_inline_runs(paragraph, text)
    return paragraph


def add_table(document, rows):
    if not rows:
        return
    column_count = len(rows[0])
    table = document.add_table(rows=1, cols=column_count)
    widths = [2700, 6660] if column_count == 2 else [1700, 3600, 4060]
    set_table_geometry(table, widths)
    header = table.rows[0]
    repeat_table_header(header)
    for index, value in enumerate(rows[0]):
        set_cell_shading(header.cells[index], LIGHT_BLUE)
        paragraph = header.cells[index].paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        add_inline_runs(paragraph, value, size=8.7, color=DARK_BLUE)
        for run in paragraph.runs:
            run.bold = True

    for row_values in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline_runs(paragraph, value, size=8.7, color=INK)
            if index == 0:
                for run in paragraph.runs:
                    run.bold = True
    set_table_geometry(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code_block(document, lines):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_border(cell, color="D8DEE5", size="4")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=8.5, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_and_add_markdown(document, lines, bullet_num_id, decimal_num_id):
    index = 0
    paragraph_lines = []
    first_h1_skipped = False

    def flush_paragraph():
        nonlocal paragraph_lines
        if paragraph_lines:
            add_body_paragraph(document, " ".join(line.strip() for line in paragraph_lines))
            paragraph_lines = []

    while index < len(lines):
        line = lines[index].rstrip("\n")
        if not line.strip():
            flush_paragraph()
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if level == 1 and not first_h1_skipped:
                first_h1_skipped = True
                index += 1
                continue
            paragraph = document.add_paragraph(style=f"Heading {level}")
            if level == 1:
                paragraph.paragraph_format.page_break_before = True
            add_inline_runs(paragraph, title, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level])
            for run in paragraph.runs:
                run.bold = True
            index += 1
            continue

        if line.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip("\n"))
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue

        if line.startswith("|"):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = []
            for table_line in table_lines:
                parts = [part.strip() for part in table_line.strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", part) for part in parts):
                    continue
                rows.append(parts)
            add_table(document, rows)
            continue

        list_match = re.match(r"^\s*(-|\*)\s+(.+)$", line)
        decimal_match = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if list_match or decimal_match:
            flush_paragraph()
            paragraph = document.add_paragraph(style="Normal")
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.25
            apply_num(paragraph, bullet_num_id if list_match else decimal_num_id)
            add_inline_runs(paragraph, (list_match or decimal_match).group(2 if list_match else 1))
            index += 1
            continue

        paragraph_lines.append(line.strip())
        index += 1

    flush_paragraph()


def add_title_page(document):
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(42)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(10)
    run = kicker.add_run("发布候选验收资料")
    set_run_font(run, size=11, color=BLUE, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("项目功能清单")
    set_run_font(run, size=30, color=DARK_BLUE, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("前端按钮、页面权限与后端 API 对照")
    set_run_font(run, size=15, color=MUTED)

    metadata = [
        ("项目", "向量知识库"),
        ("文档版本", "48 · 前端按钮功能清单"),
        ("验收日期", "2026-08-17"),
        ("运行页面", "http://127.0.0.1:8101/console/"),
        ("当前分支", "codex/release-candidate"),
    ]
    table = document.add_table(rows=0, cols=2)
    set_table_geometry(table, [1800, 7560])
    for label, value in metadata:
        cells = table.add_row().cells
        set_cell_shading(cells[0], LIGHT_BLUE)
        for idx, text in enumerate((label, value)):
            paragraph = cells[idx].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline_runs(paragraph, text, size=9.5, color=DARK_BLUE if idx == 0 else INK)
            if idx == 0:
                for run in paragraph.runs:
                    run.bold = True
    set_table_geometry(table, [1800, 7560])

    document.add_paragraph().paragraph_format.space_after = Pt(8)
    callout = document.add_table(rows=1, cols=1)
    set_table_geometry(callout, [CONTENT_WIDTH_DXA])
    cell = callout.cell(0, 0)
    set_cell_shading(cell, PALE_GREEN)
    set_cell_border(cell, color="B8DCCE", size="6")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    add_inline_runs(
        paragraph,
        "验收结论：前端可见功能已按按钮和页面整理；当前服务已启动，登录页可访问。文档同时明确了纯前端操作、权限边界、真实 API 以及没有前端按钮但仍需保留的 MCP、Worker 和健康检查能力。",
        size=10.5,
        color=INK,
    )

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run("阅读方式：先看权限速记，再按页面查找按钮；API 路径用于开发、验收和运维定位。")
    set_run_font(run, size=9.5, color=MUTED, italic=True)
    document.add_page_break()


def add_runtime_screenshot(document):
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.page_break_before = True
    add_inline_runs(heading, "附录：当前运行页面截图", size=16, color=BLUE)
    for run in heading.runs:
        run.bold = True

    add_body_paragraph(document, "截图来自当前本地运行实例的真实浏览器页面。页面处于未登录状态，没有填写或展示任何账号、密码或 API Key。")
    caption = document.add_paragraph()
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run("图 1  ·  向量知识库登录页（http://127.0.0.1:8101/console/）")
    set_run_font(run, size=9.5, color=MUTED, bold=True)
    if SCREENSHOT.exists():
        picture = document.add_picture(str(SCREENSHOT), width=Inches(6.3))
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc_pr = picture._inline.docPr
        doc_pr.set("title", "向量知识库登录页")
        doc_pr.set("descr", "向量知识库本地管理后台的未登录页面，左侧展示智能问答、文档治理和检索能力，右侧为企业邮箱和密码登录表单。")
    else:
        add_body_paragraph(document, "截图文件未找到，无法嵌入。")


def main():
    document = Document()
    configure_styles(document)
    set_page_furniture(document)
    bullet_num_id = add_num_definition(document, "bullet")
    decimal_num_id = add_num_definition(document, "decimal")
    add_title_page(document)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    parse_and_add_markdown(document, lines, bullet_num_id, decimal_num_id)
    add_runtime_screenshot(document)
    document.core_properties.title = "项目功能清单：前端按钮、页面权限与后端 API 对照"
    document.core_properties.subject = "向量知识库发布候选验收资料"
    document.core_properties.author = "Vector Knowledgebase"
    document.core_properties.comments = "Generated from docs/48-frontend-button-function-inventory.md"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
