from __future__ import annotations

import copy
import csv
import hashlib
import io
import json
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from app.config import settings
from app.models.library import Library
from app.services import (
    docx_extract,
    mineru_pdf,
    pdf_extract,
    pdf_preflight,
    pdf_routing,
    splitter,
    video_transcription,
    xlsx_extract,
)
from app.services.import_uploads import IMAGE_IMPORT_EXTENSIONS
from app.services.parser_units import (
    build_ocr_parser_units,
    build_parser_unit,
    excel_column_name,
    ocr_result_text_and_blocks,
    parser_provenance,
    text_range,
)
from app.services.video_transcription import AUDIO_IMPORT_EXTENSIONS, VIDEO_IMPORT_EXTENSIONS


@dataclass(frozen=True, slots=True)
class ParsedImport:
    normalized_text: str
    chunks: list[dict]
    splitter_name: str
    segments: list[dict] = field(default_factory=list)


RESOURCE_LIMIT_ERROR = "import content exceeds parser resource limits"

# These limits protect the two structured formats without changing the shared
# application settings or the chunked staging upload contract.
MAX_JSON_DEPTH = 64
MAX_JSON_NODES = 100_000
MAX_JSON_INPUT_BYTES = 50 * 1024 * 1024
MAX_CSV_INPUT_BYTES = 50 * 1024 * 1024
MAX_CSV_ROWS = 100_000
MAX_CSV_COLUMNS = 1_024
MAX_CSV_CELLS = 1_000_000
MAX_FANOUT_DOCUMENTS = 1_000
MAX_NORMALIZED_TEXT_CHARS = 10_000_000
MAX_TEXT_INPUT_BYTES = 64 * 1024 * 1024


class ImportResourceLimitError(ValueError):
    """Raised when structured input exceeds a parser resource budget."""

    status_code = 413

    def __init__(self, *_args: object, **_kwargs: object) -> None:
        super().__init__(RESOURCE_LIMIT_ERROR)


def rebind_converted_doc(
    parsed: ParsedImport,
    *,
    converter_version: str,
) -> ParsedImport:
    """Bind converted DOCX parser output back to the original DOC identity."""
    segments = copy.deepcopy(parsed.segments)

    def converted_parser(downstream: object) -> dict[str, str]:
        payload = downstream if isinstance(downstream, dict) else {}
        downstream_name = str(payload.get("name") or "unknown")
        return parser_provenance(
            f"libreoffice-doc-to-docx+{downstream_name}"[:128],
            "v1",
            {
                "converter_version": converter_version,
                "downstream_parser": payload,
            },
        )

    for segment in segments:
        downstream = segment.get("parser") or {}
        parser = converted_parser(downstream)
        segment["source_kind"] = "doc"
        segment["parser"] = parser
        units = [segment.get("parser_unit"), *(segment.get("structured_units") or [])]
        for unit in units:
            if not isinstance(unit, dict):
                continue
            unit_parser = converted_parser(unit.get("parser"))
            unit["source_kind"] = "doc"
            unit["parser"] = unit_parser
            source = unit.get("source")
            if isinstance(source, dict):
                source.pop("page", None)
    return ParsedImport(
        normalized_text=parsed.normalized_text,
        chunks=copy.deepcopy(parsed.chunks),
        splitter_name=parsed.splitter_name,
        segments=segments,
    )


def validate_json_resource_budget(value: object) -> None:
    """Validate JSON shape before building parser units or ingest documents."""
    stack = [(value, 0)]
    nodes = 0
    while stack:
        node, depth = stack.pop()
        nodes += 1
        if depth > MAX_JSON_DEPTH or nodes > MAX_JSON_NODES:
            raise ImportResourceLimitError
        if isinstance(node, dict):
            stack.extend((child, depth + 1) for child in node.values())
        elif isinstance(node, list):
            stack.extend((child, depth + 1) for child in node)


def validate_json_input_size(size_bytes: int) -> None:
    if size_bytes > MAX_JSON_INPUT_BYTES:
        raise ImportResourceLimitError


def validate_csv_input_size(size_bytes: int) -> None:
    if size_bytes > MAX_CSV_INPUT_BYTES:
        raise ImportResourceLimitError


def validate_csv_row_budget(
    row_number: int,
    row: object,
    total_cells: int,
) -> int:
    """Return the new cell count after validating one parsed CSV row."""
    try:
        column_count = len(row)  # type: ignore[arg-type]
    except TypeError as exc:
        raise ImportResourceLimitError from exc
    if row_number > MAX_CSV_ROWS or column_count > MAX_CSV_COLUMNS:
        raise ImportResourceLimitError
    total_cells += column_count
    if total_cells > MAX_CSV_CELLS:
        raise ImportResourceLimitError
    return total_cells


def validate_fanout_document_count(count: int) -> None:
    if count > MAX_FANOUT_DOCUMENTS:
        raise ImportResourceLimitError


def validate_normalized_text_budget(total_chars: int) -> None:
    if total_chars > MAX_NORMALIZED_TEXT_CHARS:
        raise ImportResourceLimitError


def _read_text(path: Path) -> str:
    if path.stat().st_size > MAX_TEXT_INPUT_BYTES:
        raise ImportResourceLimitError
    data = path.read_bytes()
    encodings = (
        ("utf-16",) if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ()
    ) + ("utf-8-sig", "gb18030", "big5")
    for encoding in encodings:
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("text encoding is unsupported")
    text = text.replace("\x00", "")
    validate_normalized_text_budget(len(text))
    return text


class _VisibleHtmlText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden_depth = 0

    def handle_starttag(self, tag: str, _attrs) -> None:
        if tag.lower() in {"script", "style", "template", "noscript"}:
            self.hidden_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "template", "noscript"}:
            self.hidden_depth = max(0, self.hidden_depth - 1)

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if not self.hidden_depth and value:
            self.parts.append(value)


def _parse_html(path: Path, library: Library) -> ParsedImport:
    parser = _VisibleHtmlText()
    parser.feed(_read_text(path))
    parser.close()
    return _structured_text("\n".join(parser.parts), library, source_type="html")


def _parse_pptx(path: Path, library: Library) -> ParsedImport:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > settings.xlsx_max_zip_entries:
                raise ImportResourceLimitError
            compressed = sum(info.compress_size for info in infos)
            uncompressed = sum(info.file_size for info in infos)
            if (
                compressed > settings.xlsx_max_zip_compressed_bytes
                or uncompressed > settings.xlsx_max_zip_uncompressed_bytes
                or any(info.file_size > settings.xlsx_max_zip_entry_bytes for info in infos)
                or any(
                    info.file_size
                    and info.file_size / max(1, info.compress_size)
                    > settings.xlsx_max_zip_compression_ratio
                    for info in infos
                )
            ):
                raise ImportResourceLimitError
    except zipfile.BadZipFile:
        raise ValueError("PPTX is invalid or corrupted") from None

    try:
        from pptx import Presentation

        presentation = Presentation(path)
    except Exception:  # noqa: BLE001
        raise ValueError("PPTX is invalid or corrupted") from None

    slides: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        values: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = str(getattr(shape, "text", "")).strip()
                if value:
                    values.append(value)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    line = " | ".join(cell.text.strip() for cell in row.cells).strip(" |")
                    if line:
                        values.append(line)
        if values:
            slides.append(f"【第 {index} 页】\n" + "\n".join(values))
    return _structured_text("\n\n".join(slides), library, source_type="pptx")


def _structured_text(
    text: str,
    library: Library,
    *,
    source_type: str,
    segments: list[dict] | None = None,
) -> ParsedImport:
    chunks = splitter.split_structured_text(
        text,
        chunk_size=library.chunk_size,
        chunk_overlap=library.chunk_overlap,
        splitter="text",
        base_location={"type": source_type},
    )
    if not chunks:
        raise ValueError("file contains no importable text")
    if segments is None:
        parser = parser_provenance("builtin-text", "v1", {"source_kind": source_type})
        segments = [{
            "kind": "prose",
            "text": text,
            "source_kind": source_type,
            "ordinal": 0,
            "unit_key": f"{source_type}:segment:0",
            "parser": parser,
            "location": {"type": source_type},
            "parser_unit": build_parser_unit(
                source_kind=source_type,
                unit_kind="section",
                ordinal=0,
                unit_key=f"{source_type}:segment:0",
                parser=parser,
                location={"type": source_type},
                source_text=text,
                source_start=0,
                source_end=len(text),
            ),
        }]
    return ParsedImport(text, chunks, "text", segments)


def _parse_image(
    path: Path,
    library: Library,
    *,
    file_name: str | None = None,
) -> ParsedImport:
    from app.services import ocr as ocr_service

    if path.stat().st_size > settings.image_ocr_max_input_bytes:
        raise ImportResourceLimitError

    ocr_enabled = (
        library.ocr_enabled
        if library.ocr_enabled is not None
        else settings.ocr_enabled
    )
    if not ocr_enabled:
        raise ValueError("image import requires OCR; enable OCR for this library")
    if not ocr_service.is_available():
        raise ValueError("image OCR dependency is unavailable")

    ocr_text, blocks = ocr_result_text_and_blocks(
        ocr_service.ocr_image_blocks(path.read_bytes())
    )
    ocr_text = ocr_text.replace("\x00", "").strip()
    source_name = Path(file_name).name if file_name else path.name
    source_name = source_name.replace("\x00", "").strip() or path.name
    file_name_text = f"文件名：{source_name}"
    normalized_text = (
        f"{file_name_text}\n\n{ocr_text}"
        if ocr_text
        else file_name_text
    )
    validate_normalized_text_budget(len(normalized_text))

    file_name_parser = parser_provenance(
        "upload-metadata",
        "v1",
        {"field": "file_name", "source_kind": "image"},
    )
    file_name_key = "image:file-name"
    segments = [{
        "kind": "prose",
        "text": file_name_text,
        "source_kind": "image",
        "ordinal": 0,
        "unit_key": file_name_key,
        "parser": file_name_parser,
        "location": {"type": "image"},
        "quality": {"extraction_mode": "native"},
        "parser_unit": build_parser_unit(
            source_kind="image",
            unit_kind="section",
            ordinal=0,
            unit_key=file_name_key,
            parser=file_name_parser,
            location={"type": "image"},
            source={"file_name": source_name},
            source_text=normalized_text,
            source_start=0,
            source_end=len(file_name_text),
            quality={"extraction_mode": "native"},
        ),
        "structured_units": [],
    }]

    if not ocr_text:
        return _structured_text(
            normalized_text,
            library,
            source_type="image",
            segments=segments,
        )

    parser = parser_provenance("rapidocr", "v1", {"source_kind": "image"})
    unit_key = "image:ocr"
    location = {"type": "image"}
    ocr_start = len(file_name_text) + 2
    quality = {
        "extraction_mode": "ocr",
        "native_text_present": False,
        "ocr_blocks": blocks,
    }
    segment = {
        "kind": "prose",
        "text": ocr_text,
        "source_kind": "image",
        "ordinal": 1,
        "unit_key": unit_key,
        "parser": parser,
        "location": location,
        "quality": quality,
    }
    segment["parser_unit"] = build_parser_unit(
        source_kind="image",
        unit_kind="section",
        ordinal=1,
        unit_key=unit_key,
        parser=parser,
        location=location,
        source={"file_name": source_name},
        source_text=normalized_text,
        source_start=ocr_start,
        source_end=len(normalized_text),
        quality={"extraction_mode": "ocr", "native_text_present": False},
    )
    segment["structured_units"] = build_ocr_parser_units(
        blocks,
        source_kind="image",
        parser=parser,
        parent_key=unit_key,
        extraction_mode="ocr",
        unit_text=ocr_text,
        normalized_text=normalized_text,
        normalized_offset=ocr_start,
    )
    segments.append(segment)
    return _structured_text(
        normalized_text,
        library,
        source_type="image",
        segments=segments,
    )


def _parse_csv(
    path: Path,
    library: Library,
    *,
    delimiter: str = ",",
    source_type: str = "csv",
) -> ParsedImport:
    validate_csv_input_size(path.stat().st_size)
    rows: list[str] = []
    segments: list[dict] = []
    total_cells = 0
    normalized_chars = 0
    parser = parser_provenance("python-csv", "v1", {"delimiter": delimiter})
    with io.StringIO(_read_text(path), newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for row_number, row in enumerate(reader, start=1):
            total_cells = validate_csv_row_budget(row_number, row, total_cells)
            row = [cell.strip().replace("\r\n", "\n").replace("\r", "\n") for cell in row]
            line = " | ".join(row).strip(" |")
            if line:
                normalized_chars += len(line) + (1 if rows else 0)
                validate_normalized_text_budget(normalized_chars)
                rows.append(line)
                cells: list[dict] = []
                row_key = f"csv:row:{row_number}"
                for column, cell in enumerate(row, start=1):
                    if not cell:
                        continue
                    cells.append({
                        **build_parser_unit(
                            source_kind="csv",
                            unit_kind="cell",
                            ordinal=column - 1,
                            unit_key=f"{row_key}:cell:{excel_column_name(column)}{row_number}",
                            parser=parser,
                            parent_key=row_key,
                            location={"type": "csv_row", "row": row_number},
                            source={
                                "row": {"start": row_number, "end": row_number},
                                "column": {"start": column, "end": column},
                                "cell": {
                                    "start": f"{excel_column_name(column)}{row_number}",
                                    "end": f"{excel_column_name(column)}{row_number}",
                                },
                            },
                        ),
                        "value": cell,
                    })
                segment = {
                    "kind": "prose",
                    "text": line,
                    "source_kind": "csv",
                    "ordinal": len(segments),
                    "unit_key": row_key,
                    "parser": parser,
                    "location": {"type": "csv_row", "row": row_number},
                    "structured_units": cells,
                }
                segment["parser_unit"] = build_parser_unit(
                    source_kind="csv",
                    unit_kind="row",
                    ordinal=len(segments),
                    unit_key=row_key,
                    parser=parser,
                    location=segment["location"],
                )
                segments.append(segment)
    normalized_text = "\n".join(rows)
    offset = 0
    for segment in segments:
        end = offset + len(segment["text"])
        segment["parser_unit"]["source"]["text"] = {
            "start": offset,
            "end": end,
            "ranges": [text_range(normalized_text, offset, end)],
        }
        offset = end + 1
    return _structured_text(
        normalized_text,
        library,
        source_type=source_type,
        segments=segments,
    )


def _json_unit_key(pointer: str) -> str:
    if not pointer:
        return "json:root"
    return f"json:pointer:{hashlib.sha256(pointer.encode('utf-8')).hexdigest()[:24]}"


def _bounded_json_node(value) -> tuple[object, str]:
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    )
    if len(encoded.encode("utf-8")) <= 4096:
        return value, encoded
    digest = hashlib.sha256(encoded.encode("utf-8")).hexdigest()
    summary = {
        "truncated": True,
        "sha256": digest,
        "json_value_kind": type(value).__name__,
    }
    return summary, f"<truncated {type(value).__name__} sha256={digest}>"


def _json_pointer_units(value, *, parser: dict[str, str]) -> list[dict]:
    units: list[dict] = []

    def walk(node, pointer: str, parent_key: str | None) -> None:
        unit_key = _json_unit_key(pointer)
        json_value, text = _bounded_json_node(node)
        units.append({
            **build_parser_unit(
                source_kind="json",
                unit_kind="structured_unit",
                ordinal=len(units),
                unit_key=unit_key,
                parser=parser,
                source={"json_pointer": pointer},
                parent_key=parent_key,
            ),
            "json_value_kind": type(node).__name__,
            "json_value": json_value,
            "text": text,
        })
        if isinstance(node, dict):
            for key, child in node.items():
                escaped = str(key).replace("~", "~0").replace("/", "~1")
                child_pointer = f"{pointer}/{escaped}" if pointer else f"/{escaped}"
                walk(child, child_pointer, unit_key)
        elif isinstance(node, list):
            for index, child in enumerate(node):
                child_pointer = f"{pointer}/{index}" if pointer else f"/{index}"
                walk(child, child_pointer, unit_key)

    walk(value, "", "json:document")
    return units


def _parse_json(path: Path, library: Library) -> ParsedImport:
    validate_json_input_size(path.stat().st_size)
    try:
        with path.open("r", encoding="utf-8-sig") as handle:
            value = json.load(handle)
    except RecursionError as exc:
        raise ImportResourceLimitError from exc
    validate_json_resource_budget(value)
    try:
        text = json.dumps(value, ensure_ascii=False, indent=2)
    except RecursionError as exc:
        raise ImportResourceLimitError from exc
    validate_normalized_text_budget(len(text))
    parser = parser_provenance("python-json", "v1", {"indent": 2, "ensure_ascii": False})
    segment = {
        "kind": "prose",
        "text": text,
        "source_kind": "json",
        "ordinal": 0,
        "unit_key": "json:document",
        "parser": parser,
        "location": {"type": "json"},
        "structured_units": _json_pointer_units(value, parser=parser),
    }
    segment["parser_unit"] = build_parser_unit(
        source_kind="json",
        unit_kind="section",
        ordinal=0,
        unit_key="json:document",
        parser=parser,
        location=segment["location"],
        source_text=text,
        source_start=0,
        source_end=len(text),
    )
    return _structured_text(text, library, source_type="json", segments=[segment])


def _parse_video(path: Path, library: Library, *, file_name: str) -> ParsedImport:
    transcript = video_transcription.transcribe_video(path, file_name=file_name)
    normalized_text = "\n".join(segment.text for segment in transcript.segments) or transcript.text
    segments: list[dict] = []
    offset = 0
    for ordinal, transcript_segment in enumerate(transcript.segments):
        text = transcript_segment.text
        location = {
            "type": "video",
            "start_seconds": transcript_segment.start_seconds,
            "end_seconds": transcript_segment.end_seconds,
        }
        parser = parser_provenance("video-asr", "v1", {"source_kind": "video"})
        segments.append({
            "kind": "prose",
            "text": text,
            "source_kind": "video",
            "ordinal": ordinal,
            "unit_key": f"video:segment:{ordinal}",
            "parser": parser,
            "location": location,
            "parser_unit": build_parser_unit(
                source_kind="video",
                unit_kind="timestamp",
                ordinal=ordinal,
                unit_key=f"video:segment:{ordinal}",
                parser=parser,
                location=location,
                source={
                    "timestamp": {
                        "start_seconds": transcript_segment.start_seconds,
                        "end_seconds": transcript_segment.end_seconds,
                    }
                },
                source_text=normalized_text,
                source_start=offset,
                source_end=offset + len(text),
            ),
        })
        offset += len(text) + 1
    if not segments:
        return _structured_text(normalized_text, library, source_type="video")
    return _structured_text(
        normalized_text,
        library,
        source_type="video",
        segments=segments,
    )


def _parse_audio(path: Path, library: Library, *, file_name: str) -> ParsedImport:
    transcript = video_transcription.transcribe_audio(path, file_name=file_name)
    normalized_text = "\n".join(segment.text for segment in transcript.segments) or transcript.text
    return _structured_text(normalized_text, library, source_type="audio")


def mineru_pdf_enabled(library: Library) -> bool:
    """Return whether the fail-closed MinerU rollout selects this library."""
    allowlist = {
        slug.strip()
        for slug in settings.mineru_pdf_library_slugs.split(",")
        if slug.strip()
    }
    return bool(settings.mineru_pdf_base_url.strip() and library.slug in allowlist)


def build_pdf_import_source(
    data: bytes | Path,
    library: Library,
    *,
    structured_ocr: bool = True,
) -> dict:
    """Select one PDF parser identically for every import entry point."""
    mineru_authorized = mineru_pdf_enabled(library)
    preflight = pdf_preflight.preflight_pdf(
        data,
        min_text_chars=settings.pdf_ocr_min_text_chars,
    )
    routing = pdf_routing.choose_pdf_route(
        preflight,
        mineru_authorized=mineru_authorized,
    )
    if routing["selection"] == "mineru":
        source = mineru_pdf.parse_pdf_remote(
            data,
            base_url=settings.mineru_pdf_base_url,
            expected_version=settings.mineru_pdf_expected_version,
            timeout_seconds=settings.mineru_pdf_timeout_seconds,
            max_response_bytes=settings.mineru_pdf_max_response_bytes,
            chunk_size=library.chunk_size,
            chunk_overlap=library.chunk_overlap,
            total_pages=preflight.get("total_pages"),
        )
    else:
        from app.services import ocr as ocr_service

        ocr_enabled = (
            library.ocr_enabled
            if library.ocr_enabled is not None
            else settings.ocr_enabled
        )
        if ocr_enabled and ocr_service.is_available():
            ocr_callback = (
                ocr_service.ocr_image_blocks if structured_ocr else ocr_service.ocr_image
            )
        else:
            ocr_callback = None
        source = pdf_extract.build_pdf_source(
            data,
            chunk_size=library.chunk_size,
            chunk_overlap=library.chunk_overlap,
            ocr_enabled=bool(ocr_enabled),
            ocr=ocr_callback,
            min_text_chars=settings.pdf_ocr_min_text_chars,
            render_dpi=settings.pdf_ocr_render_dpi,
            max_ocr_pages=settings.pdf_ocr_max_pages,
            preflight_report=preflight,
        )

    segments = source.get("segments")
    if not isinstance(segments, list):
        raise ValueError("PDF parser returned invalid segments")
    source["routing"] = pdf_routing.attach_pdf_routing_unit(segments, routing)
    return source


def parse_import_file(
    path: Path,
    library: Library,
    *,
    file_name: str | None = None,
) -> ParsedImport:
    suffix = Path(file_name).suffix.lower() if file_name else path.suffix.lower()
    if suffix in VIDEO_IMPORT_EXTENSIONS:
        return _parse_video(path, library, file_name=file_name or path.name)
    if suffix in AUDIO_IMPORT_EXTENSIONS:
        return _parse_audio(path, library, file_name=file_name or path.name)
    if suffix in IMAGE_IMPORT_EXTENSIONS:
        return _parse_image(path, library, file_name=file_name)
    if suffix == ".pdf":
        source = build_pdf_import_source(path, library)
        return ParsedImport(
            source["normalized_text"],
            source["chunks"],
            "text",
            source.get("segments", []),
        )
    if suffix == ".docx":
        from app.services import ocr as ocr_service

        ocr_enabled = (
            library.ocr_enabled
            if library.ocr_enabled is not None
            else settings.ocr_enabled
        )
        ocr_callback = (
            ocr_service.ocr_image_blocks
            if ocr_enabled and ocr_service.is_available()
            else None
        )
        table_aware = (
            library.docx_table_aware
            if library.docx_table_aware is not None
            else settings.docx_table_aware
        )
        if table_aware:
            source = splitter.build_structured_source_from_segments(
                docx_extract.extract_docx_segments(path, ocr=ocr_callback),
                chunk_size=library.chunk_size,
                chunk_overlap=library.chunk_overlap,
            )
            if not source["chunks"]:
                raise ValueError("DOCX contains no importable text")
            return ParsedImport(
                source["normalized_text"], source["chunks"], "docx", source.get("segments", [])
            )
        return _structured_text(
            docx_extract.extract_docx_text(path, ocr=ocr_callback),
            library,
            source_type="docx",
        )
    if suffix in {".xlsx", ".xls"}:
        source = splitter.build_structured_source_from_segments(
            xlsx_extract.extract_xlsx_segments(path, is_xls=suffix == ".xls"),
            chunk_size=library.chunk_size,
            chunk_overlap=library.chunk_overlap,
        )
        if not source["chunks"]:
            raise ValueError("spreadsheet contains no importable text")
        return ParsedImport(source["normalized_text"], source["chunks"], "docx", source.get("segments", []))
    if suffix == ".csv":
        return _parse_csv(path, library)
    if suffix == ".tsv":
        return _parse_csv(path, library, delimiter="\t", source_type="tsv")
    if suffix == ".json":
        return _parse_json(path, library)
    if suffix == ".pptx":
        return _parse_pptx(path, library)
    if suffix in {".htm", ".html"}:
        return _parse_html(path, library)
    if suffix in {
        ".cfg",
        ".conf",
        ".ini",
        ".log",
        ".markdown",
        ".md",
        ".rst",
        ".txt",
        ".xml",
        ".yaml",
        ".yml",
    }:
        return _structured_text(_read_text(path), library, source_type="text")
    source_name = file_name or path.name
    display_suffix = suffix or "无后缀"
    return _structured_text(
        "\n".join(
            (
                f"文件名：{source_name}",
                f"文件格式：{display_suffix}",
                "文件说明：该文件已安全保存为原文件；当前系统不解析其正文，可在“我的文件”中下载。",
            )
        ),
        library,
        source_type="file_metadata",
    )
