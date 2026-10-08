"""xlsx/xls 表格摄入：每个工作表 → 一个 table segment（首行作表头）。

产出与 docx_extract 同构的 table segment（kind/heading/caption/header/rows），
交给 splitter.chunk_segments 做表格感知切分。.xlsx 用 openpyxl；.xls 走 xlrd（best-effort）。
"""
from __future__ import annotations

import io
import zipfile
from contextlib import ExitStack
from math import isfinite
from pathlib import Path
from typing import Any

from app.config import settings
from app.services.parser_units import (
    build_parser_unit,
    excel_column_name,
    parser_provenance,
)

_MAX_PHYSICAL_ROWS = 2_000_000


class XlsxLimitError(ValueError):
    """Raised when an XLSX exceeds a configured resource limit."""


def _open_binary_source(data: bytes | Path):
    return data.open("rb") if isinstance(data, Path) else io.BytesIO(data)


def _validate_xlsx_archive(data: bytes | Path) -> None:
    try:
        with _open_binary_source(data) as source, zipfile.ZipFile(source) as archive:
            infos = archive.infolist()
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        raise ValueError("invalid XLSX archive") from exc

    if len(infos) > settings.xlsx_max_zip_entries:
        raise XlsxLimitError("XLSX archive exceeds the entry limit")

    compressed_total = 0
    uncompressed_total = 0
    for info in infos:
        compressed_size = info.compress_size
        uncompressed_size = info.file_size
        if compressed_size < 0 or uncompressed_size < 0:
            raise ValueError("invalid XLSX archive metadata")
        compressed_total += compressed_size
        uncompressed_total += uncompressed_size
        if compressed_total > settings.xlsx_max_zip_compressed_bytes:
            raise XlsxLimitError("XLSX archive exceeds the compressed-size limit")
        if uncompressed_total > settings.xlsx_max_zip_uncompressed_bytes:
            raise XlsxLimitError("XLSX archive exceeds the uncompressed-size limit")
        if uncompressed_size > settings.xlsx_max_zip_entry_bytes:
            raise XlsxLimitError("XLSX archive entry exceeds the size limit")
        if uncompressed_size and (
            not compressed_size
            or uncompressed_size / compressed_size > settings.xlsx_max_zip_compression_ratio
        ):
            raise XlsxLimitError("XLSX archive entry exceeds the compression-ratio limit")


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def _check_cell_text(value: Any, total_text_chars: int) -> int:
    text = _cell_text(value)
    if len(text) > settings.xlsx_max_cell_text_chars:
        raise XlsxLimitError("XLSX cell text exceeds the length limit")
    total_text_chars += len(text)
    if total_text_chars > settings.xlsx_max_text_chars:
        raise XlsxLimitError("XLSX cell text exceeds the total text limit")
    return total_text_chars


def _row_to_numbered_line(cells: tuple[Any, ...]) -> str:
    non_empty_indices: list[tuple[int, str]] = []
    for col_idx, cell in enumerate(cells, start=1):
        text = _cell_text(cell).strip()
        if text:
            non_empty_indices.append((col_idx, text))

    if not non_empty_indices:
        return ""

    parts: list[str] = []
    prev_col = 0
    SPARSE_EMPTY_THRESHOLD = 8

    for col_idx, text in non_empty_indices:
        if prev_col == 0:
            if col_idx - 1 > SPARSE_EMPTY_THRESHOLD:
                parts.append(f"[第{col_idx}列/{excel_column_name(col_idx)}] {text}")
            else:
                for _ in range(col_idx - 1):
                    parts.append("")
                parts.append(text)
        else:
            empty_span = col_idx - prev_col - 1
            if empty_span > SPARSE_EMPTY_THRESHOLD:
                parts.append(f"[第{col_idx}列/{excel_column_name(col_idx)}] {text}")
            else:
                for _ in range(empty_span):
                    parts.append("")
                parts.append(text)
        prev_col = col_idx

    return " | ".join(parts).strip(" |")


def _rows_to_numbered_lines(rows) -> list[tuple[int, str]]:
    """二维单元格 → [(row_number, 'a | b | c')]，去行尾空单元格、跳过空行。"""
    out: list[tuple[int, str]] = []
    for row_number, cells in rows:
        line = _row_to_numbered_line(cells)
        if line:
            out.append((row_number, line))
    return out


def _cell_units_for_row(
    sheet_name: str,
    row_number: int,
    cells: tuple[Any, ...],
    *,
    source_kind: str,
    parser: dict[str, str],
    parent_key: str,
    formula_cells: tuple[Any, ...] = (),
    ordinal_start: int = 0,
) -> list[dict]:
    units: list[dict] = []
    for column, value in enumerate(cells, start=1):
        formula_value = formula_cells[column - 1] if column <= len(formula_cells) else None
        formula = formula_value if isinstance(formula_value, str) and formula_value.startswith("=") else None
        if value is None and formula is None:
            continue
        cell_name = f"{excel_column_name(column)}{row_number}"
        unit = build_parser_unit(
            source_kind=source_kind,
            unit_kind="cell",
            ordinal=ordinal_start + len(units),
            unit_key=f"{parent_key}:cell:{cell_name}",
            parser=parser,
            source={
                "sheet": {"name": sheet_name},
                "row": {"start": row_number, "end": row_number},
                "column": {"start": column, "end": column},
                "cell": {"start": cell_name, "end": cell_name},
            },
            parent_key=parent_key,
        )
        unit["value"] = value
        unit["formula"] = formula
        units.append(unit)
    return units


def _cell_units(
    sheet_name: str,
    rows: list[tuple[Any, ...]],
    *,
    source_kind: str,
    parser: dict[str, str],
    parent_key: str,
    formula_rows: list[tuple[Any, ...]] | None = None,
) -> list[dict]:
    units: list[dict] = []
    for row_number, cells in enumerate(rows, start=1):
        formula_cells = formula_rows[row_number - 1] if formula_rows and row_number <= len(formula_rows) else ()
        for column, value in enumerate(cells, start=1):
            formula_value = formula_cells[column - 1] if column <= len(formula_cells) else None
            formula = formula_value if isinstance(formula_value, str) and formula_value.startswith("=") else None
            if value is None and formula is None:
                continue
            cell_name = f"{excel_column_name(column)}{row_number}"
            unit = build_parser_unit(
                source_kind=source_kind,
                unit_kind="cell",
                ordinal=len(units),
                unit_key=f"{parent_key}:cell:{cell_name}",
                parser=parser,
                source={
                    "sheet": {"name": sheet_name},
                    "row": {"start": row_number, "end": row_number},
                    "column": {"start": column, "end": column},
                    "cell": {"start": cell_name, "end": cell_name},
                },
                parent_key=parent_key,
            )
            unit["value"] = value
            unit["formula"] = formula
            units.append(unit)
    return units


def _seg(
    name: str,
    numbered_rows: list[tuple[int, str]],
    *,
    ordinal: int,
    source_kind: str,
    parser: dict[str, str],
    structured_units: list[dict],
) -> dict:
    row_numbers = [n for n, _line in numbered_rows]
    rows = [line for _n, line in numbered_rows]
    return {
        "kind": "table",
        "heading": name,
        "caption": name,
        "header": rows[0],
        "rows": rows,
        "row_numbers": row_numbers,
        "source_kind": source_kind,
        "ordinal": ordinal,
        "unit_key": f"{source_kind}:sheet:{ordinal}",
        "parser": parser,
        "structured_units": structured_units,
        "location": {"type": "sheet", "sheet": name},
    }


def _validate_sheet_dimensions(worksheets) -> None:
    for worksheet in worksheets:
        max_row = worksheet.max_row or 0
        # Pre-check guards against corrupt or malformed row counts beyond Excel limits.
        # Real cell count and row count are validated accurately during streaming in _read_xlsx
        # to avoid false-positive rejections on phantom grid dimensions caused by whole-sheet styles.
        if max_row > _MAX_PHYSICAL_ROWS:
            raise XlsxLimitError("XLSX worksheet exceeds the row limit")


def _iter_xlsx_rows(worksheet):
    """Read this worksheet without changing openpyxl's process-wide parser.

    Some exporters label text such as NULL as numeric. Treat only those cell
    values as strings, so a date style cannot send the text through from_excel.
    Keep openpyxl's parser and read-only row assembly for every other cell.
    """
    from openpyxl.worksheet._reader import VALUE_TAG, WorkSheetParser, _cast_number, from_excel

    class NumericTextParser(WorkSheetParser):
        def parse_cell(self, element):
            if element.get("t", "n") == "n":
                value = element.findtext(VALUE_TAG, None) or None
                if value is not None:
                    try:
                        numeric_value = _cast_number(value)
                    except (ValueError, TypeError):
                        element.set("t", "str")
                    else:
                        if isinstance(numeric_value, float) and not isfinite(numeric_value):
                            element.set("t", "str")
                        elif int(element.get("s") or 0) in self.date_formats:
                            try:
                                from_excel(
                                    numeric_value,
                                    self.epoch,
                                    timedelta=int(element.get("s") or 0) in self.timedelta_formats,
                                )
                            except (OverflowError, ValueError):
                                element.set("t", "str")
            return super().parse_cell(element)

    workbook = worksheet.parent
    next_row = 1
    empty_row = ()
    with worksheet._get_source() as source:
        parser = NumericTextParser(
            source,
            worksheet._shared_strings,
            data_only=workbook.data_only,
            epoch=workbook.epoch,
            date_formats=workbook._date_formats,
            timedelta_formats=workbook._timedelta_formats,
        )
        for row_number, cells in parser.parse():
            if not 1 <= row_number <= _MAX_PHYSICAL_ROWS:
                raise XlsxLimitError("XLSX worksheet exceeds the physical row bounds")
            while next_row < row_number:
                yield empty_row
                next_row += 1
            if next_row <= row_number:
                last_value_column = max((cell["column"] for cell in cells if cell["value"] is not None), default=0)
                yield worksheet._get_row(cells, 1, last_value_column, values_only=True) if last_value_column else empty_row
                next_row = row_number + 1


def _read_xlsx(data: bytes | Path) -> list[dict]:
    _validate_xlsx_archive(data)
    import openpyxl

    def open_source():
        return _open_binary_source(data)

    with ExitStack() as stack:
        source = stack.enter_context(open_source())
        formula_source = stack.enter_context(open_source())
        wb = openpyxl.load_workbook(source, read_only=True, data_only=True, keep_links=False)
        stack.callback(wb.close)
        formula_wb = openpyxl.load_workbook(
            formula_source, read_only=True, data_only=False, keep_links=False
        )
        stack.callback(formula_wb.close)
        if len(wb.worksheets) > settings.xlsx_max_sheets:
            raise XlsxLimitError("XLSX workbook exceeds the sheet limit")
        if len(formula_wb.worksheets) > settings.xlsx_max_sheets:
            raise XlsxLimitError("XLSX workbook exceeds the sheet limit")
        _validate_sheet_dimensions(wb.worksheets)
        _validate_sheet_dimensions(formula_wb.worksheets)

        segs: list[dict] = []
        parser = parser_provenance("openpyxl", "v1", {"data_only": True})
        total_rows = 0
        total_cells = 0
        total_text_chars = 0
        for ordinal, (ws, formula_ws) in enumerate(zip(wb.worksheets, formula_wb.worksheets, strict=True)):
            numbered_rows: list[tuple[int, str]] = []
            structured_units: list[dict] = []
            for row_number, (data_cells, formula_cells) in enumerate(
                zip(_iter_xlsx_rows(ws), _iter_xlsx_rows(formula_ws), strict=True),
                start=1,
            ):
                d_len = len(data_cells)
                while d_len > 0 and data_cells[d_len - 1] is None:
                    d_len -= 1
                f_len = len(formula_cells)
                while f_len > 0 and formula_cells[f_len - 1] is None:
                    f_len -= 1

                row_width = max(d_len, f_len)
                if row_width == 0:
                    continue

                total_rows += 1
                if total_rows > settings.xlsx_max_rows:
                    raise XlsxLimitError("XLSX workbook exceeds the row limit")
                total_cells += row_width
                if total_cells > settings.xlsx_max_cells:
                    raise XlsxLimitError("XLSX workbook exceeds the cell limit")

                trimmed_data = data_cells[:row_width] + (None,) * max(0, row_width - len(data_cells))
                trimmed_formula = formula_cells[:row_width]

                for value in trimmed_data:
                    total_text_chars = _check_cell_text(value, total_text_chars)
                for value in trimmed_formula:
                    total_text_chars = _check_cell_text(value, total_text_chars)
                line = _row_to_numbered_line(trimmed_data)
                if line:
                    numbered_rows.append((row_number, line))
                structured_units.extend(
                    _cell_units_for_row(
                        ws.title,
                        row_number,
                        trimmed_data,
                        source_kind="xlsx",
                        parser=parser,
                        parent_key=f"xlsx:sheet:{ordinal}",
                        formula_cells=trimmed_formula,
                        ordinal_start=len(structured_units),
                    )
                )
            if numbered_rows:
                segs.append(
                    _seg(
                        ws.title,
                        numbered_rows,
                        ordinal=ordinal,
                        source_kind="xlsx",
                        parser=parser,
                        structured_units=structured_units,
                    )
                )
        return segs


def _read_xls(data: bytes) -> list[dict]:
    try:
        import xlrd  # noqa: F401
    except Exception as exc:  # noqa: BLE001
        raise ValueError(".xls 需要 xlrd（pip install xlrd），或先另存为 .xlsx") from exc
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    segs: list[dict] = []
    parser = parser_provenance("xlrd", "v1", {"data_only": True})
    for ordinal, sh in enumerate(book.sheets()):
        raw_rows = [tuple(sh.row_values(r)) for r in range(sh.nrows)]
        rows = _rows_to_numbered_lines(enumerate(raw_rows, start=1))
        if rows:
            segs.append(
                _seg(
                    sh.name,
                    rows,
                    ordinal=ordinal,
                    source_kind="xls",
                    parser=parser,
                    structured_units=_cell_units(
                        sh.name,
                        raw_rows,
                        source_kind="xls",
                        parser=parser,
                        parent_key=f"xls:sheet:{ordinal}",
                    ),
                )
            )
    return segs


def extract_xlsx_segments(data: bytes | Path, *, is_xls: bool = False) -> list[dict]:
    """读电子表格 → table segment 列表（空表跳过）。"""
    if is_xls:
        size_bytes = data.stat().st_size if isinstance(data, Path) else len(data)
        if size_bytes > settings.xls_max_input_bytes:
            raise XlsxLimitError("XLS input exceeds the size limit")
        return _read_xls(data.read_bytes() if isinstance(data, Path) else data)
    return _read_xlsx(data)
