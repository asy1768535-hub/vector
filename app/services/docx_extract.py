r"""docx 结构化抽取：按文档顺序产出「段落 / 表格」segment，并维护章节路径。

为什么结构化：表格里常有关键信息（技术选型、人员签认…），若拍平进大段落 chunk 会被
稀释、检索排不上（见评测 #8）。这里把内容拆成 segment，交给 splitter.chunk_segments
做表格感知切分——每个表格单独成块并带上「章节 / 表名 / 表头」上下文。

segment 结构（普通 dict）：
  - {"kind":"prose","heading":<章节路径>,"text":...}     段落 + 图片 OCR 文本
  - {"kind":"table","heading":...,"caption":...,"header":<首行>,"rows":[<"a | b | c">...]}

heading 检测双信号：段落样式名含「标题」/「Heading」 或 文本匹配 `^\d+(\.\d+)*\s+`。
"""
from __future__ import annotations

import io
import re
import zipfile
from typing import Any
from pathlib import Path
from typing import Callable, Optional

from docx import Document
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.parser_units import (
    build_ocr_parser_units,
    build_parser_unit,
    excel_column_name,
    ocr_result_text_and_blocks,
    parser_provenance,
)
from app.services.ocr import OcrResourceLimitError

_HEADING_NUM = re.compile(r"^\d+(\.\d+)*\s+\S")
_LEVEL_CN = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6}

# Module-local limits keep parser safety independent from import/staging settings.
DOCX_MAX_ZIP_ENTRIES = 4096
DOCX_MAX_ZIP_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
DOCX_MAX_ZIP_ENTRY_BYTES = 64 * 1024 * 1024
DOCX_MAX_ZIP_COMPRESSION_RATIO = 100.0
DOCX_MAX_IMAGES = 128
DOCX_MAX_IMAGE_BYTES = 128 * 1024 * 1024
DOCX_MAX_SEGMENTS = 10_000
DOCX_MAX_STRUCTURED_UNITS = 100_000
DOCX_MAX_TABLES = 1_000
DOCX_MAX_TABLE_ROWS = 100_000
DOCX_MAX_TABLE_CELLS = 1_000_000


class DocxResourceLimitError(ValueError):
    """DOCX parser resource budget was exceeded."""


def _resource_limit(name: str) -> None:
    raise DocxResourceLimitError(f"DOCX resource limit exceeded: {name}")


def _preflight_docx(data: bytes | Path) -> None:
    """Inspect ZIP metadata before python-docx opens or expands the document."""
    source = data if isinstance(data, Path) else io.BytesIO(data)
    try:
        with zipfile.ZipFile(source) as archive:
            infos = archive.infolist()
            if len(infos) > DOCX_MAX_ZIP_ENTRIES:
                _resource_limit("zip entries")

            total_uncompressed = 0
            image_count = 0
            image_bytes = 0
            for info in infos:
                uncompressed = max(0, int(info.file_size))
                compressed = max(0, int(info.compress_size))
                total_uncompressed += uncompressed
                if total_uncompressed > DOCX_MAX_ZIP_UNCOMPRESSED_BYTES:
                    _resource_limit("zip uncompressed bytes")
                if uncompressed > DOCX_MAX_ZIP_ENTRY_BYTES:
                    _resource_limit("zip entry bytes")
                ratio = float("inf") if uncompressed and not compressed else (
                    uncompressed / compressed if compressed else 1.0
                )
                if ratio > DOCX_MAX_ZIP_COMPRESSION_RATIO:
                    _resource_limit("zip compression ratio")

                name = info.filename.replace("\\", "/").lower()
                if name.startswith("word/media/") and not info.is_dir():
                    image_count += 1
                    image_bytes += uncompressed
                    if image_count > DOCX_MAX_IMAGES:
                        _resource_limit("embedded image count")
                    if image_bytes > DOCX_MAX_IMAGE_BYTES:
                        _resource_limit("embedded image bytes")
    except (OSError, zipfile.BadZipFile):
        # Let python-docx preserve its existing malformed-document error path.
        return


class _DocxBudget:
    def __init__(self) -> None:
        self.segments = 0
        self.units = 0
        self.tables = 0
        self.table_rows = 0
        self.table_cells = 0
        self.images = 0
        self.image_bytes = 0

    def add_segment(self) -> None:
        self.segments += 1
        if self.segments > DOCX_MAX_SEGMENTS:
            _resource_limit("segments")

    def add_units(self, count: int = 1) -> None:
        self.units += count
        if self.units > DOCX_MAX_STRUCTURED_UNITS:
            _resource_limit("structured units")

    def add_table(self) -> None:
        self.tables += 1
        if self.tables > DOCX_MAX_TABLES:
            _resource_limit("tables")

    def add_table_row(self) -> None:
        self.table_rows += 1
        if self.table_rows > DOCX_MAX_TABLE_ROWS:
            _resource_limit("table rows")

    def add_table_cell(self) -> None:
        self.table_cells += 1
        if self.table_cells > DOCX_MAX_TABLE_CELLS:
            _resource_limit("table cells")

    def add_image(self, size: int) -> None:
        self.images += 1
        self.image_bytes += max(0, size)
        if self.images > DOCX_MAX_IMAGES:
            _resource_limit("embedded image count")
        if self.image_bytes > DOCX_MAX_IMAGE_BYTES:
            _resource_limit("embedded image bytes")


def _iter_block_items(parent):
    """按文档顺序产出 Paragraph / Table（python-docx 默认把两者分到不同集合）。"""
    body = parent.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _table_row_cells(
    table: Table, *, budget: _DocxBudget | None = None
) -> list[tuple[int, list[tuple[int, str]], str]]:
    records: list[tuple[int, list[tuple[int, str]], str]] = []
    for row_number, row in enumerate(table.rows, start=1):
        if budget:
            budget.add_table_row()
        dedup: list[tuple[int, str]] = []
        for column, cell in enumerate(row.cells, start=1):
            if budget:
                budget.add_table_cell()
            value = cell.text.strip()
            if not dedup or dedup[-1][1] != value:
                dedup.append((column, value))
        line = " | ".join(value for _column, value in dedup).strip(" |")
        if line:
            records.append((row_number, dedup, line))
    return records


def _table_rows(table: Table) -> list[str]:
    """表格转行列表：行内 ' | ' 拼接，去相邻重复单元格（合并单元格），跳过空行。"""
    return [line for _row_number, _cells, line in _table_row_cells(table)]


def _table_parser_units(
    table: Table,
    *,
    table_key: str,
    table_index: int,
    parser: dict[str, str],
    records: list[tuple[int, list[tuple[int, str]], str]] | None = None,
    budget: _DocxBudget | None = None,
) -> list[dict]:
    table_source = {"table": {"index": max(0, table_index - 1)}}
    units: list[dict] = []
    row_records = records if records is not None else _table_row_cells(table, budget=budget)
    for row_number, cells, line in row_records:
        row_key = f"{table_key}:row:{row_number}"
        row_source = {
            **table_source,
            "row": {"start": row_number, "end": row_number},
        }
        if budget:
            budget.add_units()
        row_unit = build_parser_unit(
            source_kind="docx",
            unit_kind="row",
            ordinal=len(units),
            unit_key=row_key,
            parser=parser,
            source=row_source,
            parent_key=table_key,
        )
        row_unit["text"] = line
        units.append(row_unit)
        for column, value in cells:
            if budget:
                budget.add_units()
            cell_name = f"{excel_column_name(column)}{row_number}"
            cell_key = f"{row_key}:cell:{cell_name}"
            cell_unit = build_parser_unit(
                source_kind="docx",
                unit_kind="cell",
                ordinal=len(units),
                unit_key=cell_key,
                parser=parser,
                source={
                    **row_source,
                    "column": {"start": column, "end": column},
                    "cell": {"start": cell_name, "end": cell_name},
                },
                parent_key=row_key,
            )
            cell_unit["value"] = value
            cell_unit["text"] = value
            units.append(cell_unit)
    return units


def _paragraph_image_blobs(paragraph: Paragraph):
    """取段落里内嵌图片的字节流（按出现顺序）。无图返回空列表。"""
    for blip in paragraph._p.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        try:
            yield paragraph.part.related_parts[rid].blob
        except KeyError:
            continue


def _heading_info(paragraph: Paragraph) -> Optional[tuple[str, int]]:
    """是标题则返回 (标题文本, 层级)，否则 None。"""
    txt = paragraph.text.strip()
    if not txt:
        return None
    style = (paragraph.style.name if paragraph.style else "") or ""
    is_style = ("标题" in style) or ("heading" in style.lower())
    m = _HEADING_NUM.match(txt)
    if not (is_style or m):
        return None
    if m:
        level = txt.split()[0].count(".") + 1      # "2.2 …" → 2 段 → level 2
    elif "标题" in style:
        level = next((v for k, v in _LEVEL_CN.items() if k in style), 1)
    else:
        level = 1
    return txt, level


def extract_docx_segments(
    data: bytes | Path, ocr: Optional[Callable[[bytes], Any]] = None
) -> list[dict]:
    """按文档顺序抽成 prose/table segment 列表，并维护章节路径。"""
    _preflight_docx(data)
    doc = Document(data if isinstance(data, Path) else io.BytesIO(data))
    budget = _DocxBudget()
    segs: list[dict] = []
    heading_stack: list[str] = []          # heading_stack[i] = 第 i+1 级标题
    prose_buf: list[str] = []
    prose_ocr_blocks: list[dict[str, Any]] = []
    prose_has_native_text = False
    prose_has_ocr_text = False
    prose_start: int | None = None
    prose_end: int | None = None
    paragraph_index = 0
    table_index = 0
    parser = parser_provenance("python-docx", "v1", {"ocr_enabled": ocr is not None})

    def cur_heading() -> str:
        return " / ".join(h for h in heading_stack if h)

    def flush_prose() -> None:
        nonlocal prose_buf, prose_start, prose_end, prose_ocr_blocks
        nonlocal prose_has_native_text, prose_has_ocr_text
        if prose_buf:
            text = "\n".join(prose_buf).strip()
            if text:
                location = {
                    "type": "paragraph",
                    "start_paragraph": prose_start,
                    "end_paragraph": prose_end,
                    "heading": cur_heading(),
                }
                if prose_has_native_text and prose_has_ocr_text:
                    extraction_mode = "mixed"
                elif prose_has_ocr_text:
                    extraction_mode = "ocr"
                elif prose_has_native_text:
                    extraction_mode = "native"
                else:
                    extraction_mode = "not_applicable"
                quality = {
                    "extraction_mode": extraction_mode,
                    "native_text_present": prose_has_native_text,
                    "ocr_blocks": list(prose_ocr_blocks),
                }
                budget.add_segment()
                budget.add_units()
                segment = {
                    "kind": "prose",
                    "heading": cur_heading(),
                    "text": text,
                    "source_kind": "docx",
                    "ordinal": len(segs),
                    "unit_key": f"docx:segment:{len(segs)}",
                    "parser": parser,
                    "location": location,
                    "quality": quality,
                }
                segment["parser_unit"] = build_parser_unit(
                    source_kind="docx",
                    unit_kind="section",
                    ordinal=segment["ordinal"],
                    unit_key=segment["unit_key"],
                    parser=parser,
                    location=location,
                    section_path=[part for part in cur_heading().split(" / ") if part],
                    quality={
                        "extraction_mode": extraction_mode,
                        "native_text_present": prose_has_native_text,
                    },
                )
                if prose_ocr_blocks:
                    budget.add_units(len(prose_ocr_blocks))
                    segment["structured_units"] = build_ocr_parser_units(
                        prose_ocr_blocks,
                        source_kind="docx",
                        parser=parser,
                        parent_key=segment["unit_key"],
                        extraction_mode=extraction_mode if extraction_mode in {"ocr", "mixed"} else "ocr",
                        unit_text=text,
                    )
                segs.append(segment)
        prose_buf = []
        prose_ocr_blocks = []
        prose_start = None
        prose_end = None
        prose_has_native_text = False
        prose_has_ocr_text = False

    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            info = _heading_info(block)
            if info:
                flush_prose()
                title, level = info
                del heading_stack[level - 1:]              # 截断到上一级
                while len(heading_stack) < level - 1:      # 跳级时补空位
                    heading_stack.append("")
                heading_stack.append(title)
                continue
            t = block.text.strip()
            if t:
                paragraph_index += 1
                if prose_start is None:
                    prose_start = paragraph_index
                prose_end = paragraph_index
                prose_buf.append(t)
                prose_has_native_text = True
            if ocr is not None:
                for blob in _paragraph_image_blobs(block):
                    budget.add_image(len(blob))
                    try:
                        ot, blocks = ocr_result_text_and_blocks(ocr(blob))
                    except OcrResourceLimitError:
                        raise
                    except Exception:  # noqa: BLE001
                        ot, blocks = "", []
                    if ot and ot.strip():
                        prose_buf.append(ot.strip())
                        prose_has_ocr_text = True
                        prose_ocr_blocks.extend(blocks)
        else:  # Table
            caption = prose_buf[-1] if prose_buf else cur_heading()
            flush_prose()
            budget.add_table()
            records = _table_row_cells(block, budget=budget)
            rows = [line for _row_number, _cells, line in records]
            if rows:
                table_index += 1
                budget.add_segment()
                budget.add_units()
                location = {"type": "table", "table_index": table_index, "heading": cur_heading()}
                segment = {
                    "kind": "table", "heading": cur_heading(),
                    "caption": caption, "header": rows[0], "rows": rows,
                    "source_kind": "docx",
                    "ordinal": len(segs),
                    "unit_key": f"docx:segment:{len(segs)}",
                    "parser": parser,
                    "location": location,
                }
                segment["parser_unit"] = build_parser_unit(
                    source_kind="docx",
                    unit_kind="table",
                    ordinal=segment["ordinal"],
                    unit_key=segment["unit_key"],
                    parser=parser,
                    location=location,
                    section_path=[part for part in cur_heading().split(" / ") if part],
                )
                segment["structured_units"] = _table_parser_units(
                    block,
                    table_key=segment["unit_key"],
                    table_index=table_index,
                    parser=parser,
                    records=records,
                    budget=budget,
                )
                segs.append(segment)
    flush_prose()
    return segs


def extract_docx_text(
    data: bytes | Path, ocr: Optional[Callable[[bytes], Any]] = None
) -> str:
    """从 docx 抽取扁平正文（段落 + 表格 + 可选图片 OCR，按文档顺序）。

    供 content_hash / 不需要结构的场景用；表格感知切分请走 extract_docx_segments。
    """
    parts: list[str] = []
    for s in extract_docx_segments(data, ocr):
        parts.append(s["text"] if s["kind"] == "prose" else "\n".join(s["rows"]))
    return "\n".join(parts)
