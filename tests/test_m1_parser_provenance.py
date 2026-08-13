from __future__ import annotations

import io
import hashlib
import json
import types
from types import SimpleNamespace

import docx
import openpyxl
import pytest

from app.services import pdf_extract, splitter, xlsx_extract
from app.services.docx_extract import extract_docx_segments
from app.services.import_parsing import (
    ImportResourceLimitError,
    parse_import_file,
)
from app.services.parser_units import (
    build_ocr_parser_units,
    build_parser_unit,
    parser_provenance,
    validate_parser_unit_contract,
    validate_parser_unit_hierarchy,
)
from app.services.pdf_extract import PdfOcrUnavailableError, build_pdf_source


def _library(**overrides):
    values = {
        "chunk_size": 200,
        "chunk_overlap": 0,
        "ocr_enabled": False,
        "docx_table_aware": True,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def _fake_reader(page_texts):
    pages = [types.SimpleNamespace(extract_text=(lambda text=text: text)) for text in page_texts]
    return types.SimpleNamespace(pages=pages)


def _pdf_kwargs(**overrides):
    values = {
        "ocr_enabled": False,
        "ocr": None,
        "min_text_chars": 5,
        "render_dpi": 200,
        "max_ocr_pages": 5,
    }
    values.update(overrides)
    return values


def _flatten_segments(segments):
    units = []
    for segment in segments:
        units.append(segment["parser_unit"])
        units.extend(segment.get("structured_units") or [])
    return units


def test_native_pdf_segments_have_page_provenance_without_polluting_chunks(monkeypatch):
    monkeypatch.setattr(pdf_extract, "_open_reader", lambda _data: _fake_reader(["native page text"]))

    source = build_pdf_source(b"pdf", chunk_size=80, chunk_overlap=0, **_pdf_kwargs())

    assert source["chunks"]
    assert all("parser_unit" not in chunk for chunk in source["chunks"])
    unit = source["segments"][0]["parser_unit"]
    assert unit["source_kind"] == "pdf"
    assert unit["source"]["page"] == {"start": 1, "end": 1}
    assert unit["quality"]["extraction_mode"] == "native"
    text_span = unit["source"]["text"]
    assert source["normalized_text"][text_span["start"] : text_span["end"]] == source["segments"][0]["text"]
    assert text_span["ranges"][0]["sha256"] == hashlib.sha256(
        source["segments"][0]["text"].encode("utf-8")
    ).hexdigest()
    validate_parser_unit_contract(unit)
    validate_parser_unit_hierarchy(_flatten_segments(source["segments"]))


def test_scanned_and_mixed_pdf_preserve_ocr_quality(monkeypatch):
    monkeypatch.setattr(pdf_extract, "_open_reader", lambda _data: _fake_reader(["", "native text here"]))
    monkeypatch.setattr(pdf_extract, "_render_page_png", lambda *_args: b"png")

    ocr_result = [{"text": "scanned line", "confidence": 0.91, "bbox": [1, 2, 30, 40]}]
    source = build_pdf_source(
        b"pdf",
        chunk_size=80,
        chunk_overlap=0,
        **_pdf_kwargs(ocr_enabled=True, ocr=lambda _data: ocr_result),
    )

    ocr_unit = next(s["parser_unit"] for s in source["segments"] if s["location"]["page"] == 1)
    assert ocr_unit["quality"]["extraction_mode"] == "ocr"
    assert source["segments"][0]["quality"]["ocr_blocks"] == [
        {"text": "scanned line", "confidence": 0.91, "bbox": [1.0, 2.0, 30.0, 40.0]}
    ]
    line_unit = source["segments"][0]["structured_units"][0]
    assert line_unit["unit_kind"] == "image_region"
    assert line_unit["source"]["page"] == {"start": 1, "end": 1}
    assert line_unit["source"]["bbox"]["coordinate_system"] == "image_pixels"
    assert line_unit["quality"] == {"extraction_mode": "ocr", "ocr_confidence": 0.91}
    assert "ocr_blocks" not in line_unit
    validate_parser_unit_contract(line_unit)
    validate_parser_unit_hierarchy(_flatten_segments(source["segments"]))

    monkeypatch.setattr(
        pdf_extract,
        "_open_reader",
        lambda _data: types.SimpleNamespace(
            pages=[
                types.SimpleNamespace(
                    extract_text=lambda: "native text here",
                    images=[types.SimpleNamespace(data=b"embedded")],
                )
            ]
        ),
    )
    mixed = build_pdf_source(
        b"pdf",
        chunk_size=80,
        chunk_overlap=0,
        **_pdf_kwargs(ocr_enabled=True, ocr=lambda _data: ocr_result),
    )
    assert mixed["segments"][0]["parser_unit"]["quality"]["extraction_mode"] == "mixed"


def test_unparsed_visual_pdf_and_missing_ocr_dependency_are_explicit(monkeypatch):
    page = types.SimpleNamespace(
        extract_text=lambda: "native text here",
        images=[object()],
    )
    monkeypatch.setattr(pdf_extract, "_open_reader", lambda _data: types.SimpleNamespace(pages=[page]))
    source = build_pdf_source(b"pdf", chunk_size=80, chunk_overlap=0, **_pdf_kwargs(ocr_enabled=True, ocr=lambda _data: "unused"))
    assert source["segments"][0]["location"]["extraction_mode"] == "native_visual_unparsed"
    assert source["segments"][0]["quality"]["extraction_mode"] == "unparsed"
    assert source["segments"][0]["parser_unit"]["quality"]["unparsed_reason"] == "native_visual_content_without_ocr"

    monkeypatch.setattr(pdf_extract, "_open_reader", lambda _data: _fake_reader([""]))
    with pytest.raises(PdfOcrUnavailableError):
        pdf_extract.extract_pdf_text(b"pdf", **_pdf_kwargs(ocr_enabled=True, ocr=None))


def _build_docx_fixture(*, include_image: bool = False, include_native: bool = True) -> bytes:
    document = docx.Document()
    document.add_heading("Section", level=1)
    if include_native:
        document.add_paragraph("Paragraph text")
    if include_image:
        image = pytest.importorskip("PIL.Image")
        image_bytes = io.BytesIO()
        image.new("RGB", (20, 20), (255, 255, 255)).save(image_bytes, format="PNG")
        image_bytes.seek(0)
        document.add_picture(image_bytes)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "A"
    table.rows[1].cells[1].text = "B"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_segments_preserve_heading_table_and_embedded_ocr_metadata():
    segments = extract_docx_segments(
        _build_docx_fixture(include_image=True),
        ocr=lambda _data: [{"text": "image text", "confidence": 0.8, "bbox": [0, 0, 5, 6]}],
    )

    prose = next(segment for segment in segments if segment["kind"] == "prose")
    table = next(segment for segment in segments if segment["kind"] == "table")
    assert prose["parser_unit"]["section_path"] == ["Section"]
    assert prose["quality"]["extraction_mode"] == "mixed"
    assert "ocr_blocks" not in prose["parser_unit"]["quality"]
    validate_parser_unit_contract(prose["parser_unit"])
    assert "image text" in prose["text"]
    assert prose["quality"]["ocr_blocks"] == [
        {"text": "image text", "confidence": 0.8, "bbox": [0.0, 0.0, 5.0, 6.0]}
    ]
    validate_parser_unit_contract(prose["structured_units"][0])
    assert table["parser_unit"]["source"]["table"] == {"index": 0}
    assert all(unit["unit_kind"] == "image_region" for unit in prose.get("structured_units", []))
    validate_parser_unit_hierarchy(_flatten_segments(segments))


@pytest.mark.parametrize(
    ("include_native", "include_image", "expected_mode"),
    [(True, False, "native"), (False, True, "ocr"), (True, True, "mixed")],
)
def test_docx_prose_extraction_mode_reflects_native_and_ocr_content(
    include_native, include_image, expected_mode
):
    segments = extract_docx_segments(
        _build_docx_fixture(include_image=include_image, include_native=include_native),
        ocr=lambda _data: [{"text": "image text", "confidence": 0.8, "bbox": [0, 0, 5, 6]}],
    )
    prose = next(segment for segment in segments if segment["kind"] == "prose")
    assert prose["quality"]["extraction_mode"] == expected_mode


def _build_xlsx_fixture() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Formula"])
    sheet.append(["A", "=1+1"])
    second = workbook.create_sheet("Second")
    second.append(["Name", "Value"])
    second.append(["B", "2"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_xlsx_segments_include_sheet_row_cell_and_formula_provenance():
    segments = xlsx_extract.extract_xlsx_segments(_build_xlsx_fixture())
    segment = segments[0]

    assert segment["source_kind"] == "xlsx"
    assert [item["ordinal"] for item in segments] == [0, 1]
    cell = next(unit for unit in segment["structured_units"] if unit["source"]["cell"]["start"] == "B2")
    assert cell["source"]["row"] == {"start": 2, "end": 2}
    assert cell["source"]["column"] == {"start": 2, "end": 2}
    assert cell["formula"] == "=1+1"
    source = splitter.build_structured_source_from_segments(segments, chunk_size=80, chunk_overlap=0)
    assert source["segments"][0]["parser_unit"]["source"]["sheet"] == {"name": "Data"}
    assert [item["parser_unit"]["unit_key"] for item in source["segments"]] == [
        "xlsx:sheet:0",
        "xlsx:sheet:1",
    ]
    validate_parser_unit_hierarchy(_flatten_segments(source["segments"]))


def test_parser_unit_adapter_rejects_m0_incompatible_quality():
    with pytest.raises(ValueError, match="EvidenceLocatorV1 contract"):
        validate_parser_unit_contract(
            {
                "version": "v1",
                "unit_kind": "section",
                "source_kind": "pdf",
                "ordinal": 0,
                "unit_key": "pdf:segment:0",
                "parser": parser_provenance("fixture", "v1"),
                "source": {"page": {"start": 1, "end": 1}},
                "quality": {"extraction_mode": "native_visual_unparsed"},
            }
        )


def test_parser_unit_hierarchy_rejects_duplicate_dangling_self_and_cycle():
    parser = parser_provenance("fixture", "v1")

    def unit(key, parent=None):
        return build_parser_unit(
            source_kind="text",
            unit_kind="section",
            ordinal=0,
            unit_key=key,
            parser=parser,
            parent_key=parent,
        )

    with pytest.raises(ValueError, match="duplicate"):
        validate_parser_unit_hierarchy([unit("root"), unit("root")])
    with pytest.raises(ValueError, match="dangling"):
        validate_parser_unit_hierarchy([unit("child", "missing")])
    with pytest.raises(ValueError, match="itself"):
        validate_parser_unit_hierarchy([unit("self", "self")])
    with pytest.raises(ValueError, match="cycle"):
        validate_parser_unit_hierarchy([unit("a", "b"), unit("b", "a")])


def test_ocr_line_without_proven_bbox_is_not_promoted_to_image_region():
    assert build_ocr_parser_units(
        [{"text": "unlocated line", "confidence": 0.5}],
        source_kind="pdf",
        parser=parser_provenance("pypdf", "v1"),
        parent_key="pdf:page:1",
        extraction_mode="ocr",
        unit_text="unlocated line",
        page=1,
    ) == []


def test_xls_parser_is_available_as_pure_parser_but_dispatch_remains_unsupported(monkeypatch, tmp_path):
    parser = parser_provenance("fixture-xls", "v1")
    segment = {
        "kind": "table",
        "heading": "Legacy",
        "caption": "Legacy",
        "header": "Name",
        "rows": ["Name", "A"],
        "row_numbers": [1, 2],
        "source_kind": "xls",
        "ordinal": 0,
        "parser": parser,
        "location": {"type": "sheet", "sheet": "Legacy"},
        "structured_units": [],
    }
    monkeypatch.setattr(xlsx_extract, "_read_xls", lambda _data: [segment])
    assert xlsx_extract.extract_xlsx_segments(b"xls", is_xls=True)[0]["source_kind"] == "xls"

    path = tmp_path / "legacy.xls"
    path.write_bytes(b"xls")
    with pytest.raises(ValueError, match="unsupported file type"):
        parse_import_file(path, _library())


def test_csv_and_json_segments_keep_row_cell_and_pointer_units(tmp_path):
    csv_path = tmp_path / "rows.csv"
    csv_path.write_text("name,value\nA,1\n", encoding="utf-8")
    csv_result = parse_import_file(csv_path, _library())
    csv_units = csv_result.segments[0]["structured_units"]
    assert csv_result.chunks
    assert all("parser_unit" not in chunk for chunk in csv_result.chunks)
    assert csv_units[0]["source"]["cell"] == {"start": "A1", "end": "A1"}
    row_unit = csv_result.segments[0]["parser_unit"]
    assert csv_result.normalized_text[row_unit["source"]["text"]["start"] : row_unit["source"]["text"]["end"]] == csv_result.segments[0]["text"]
    assert row_unit["source"]["text"]["ranges"][0]["sha256"] == hashlib.sha256(
        csv_result.segments[0]["text"].encode("utf-8")
    ).hexdigest()
    validate_parser_unit_hierarchy(_flatten_segments(csv_result.segments))

    json_path = tmp_path / "records.json"
    json_path.write_text(json.dumps({"records": [{"name": "A"}]}), encoding="utf-8")
    json_result = parse_import_file(json_path, _library())
    pointers = {unit["source"]["json_pointer"] for unit in json_result.segments[0]["structured_units"]}
    assert "/records/0/name" in pointers
    validate_parser_unit_hierarchy(_flatten_segments(json_result.segments))


def test_csv_rows_have_exact_normalized_ranges_for_empty_rows_and_quoted_newlines(tmp_path):
    csv_path = tmp_path / "rows.csv"
    csv_path.write_text('name,value\n\n"A\n1",two\nB,three\n', encoding="utf-8")

    result = parse_import_file(csv_path, _library())

    assert result.normalized_text == "name | value\nA\r\n1 | two\nB | three"
    assert [segment["parser_unit"]["unit_key"] for segment in result.segments] == [
        "csv:row:1",
        "csv:row:3",
        "csv:row:4",
    ]
    for segment in result.segments:
        span = segment["parser_unit"]["source"]["text"]
        assert result.normalized_text[span["start"] : span["end"]] == segment["text"]
        assert span["ranges"][0]["sha256"] == hashlib.sha256(segment["text"].encode("utf-8")).hexdigest()
    assert result.segments[1]["structured_units"][0]["source"]["cell"] == {"start": "A3", "end": "A3"}
    assert result.segments[2]["structured_units"][1]["source"]["cell"] == {"start": "B4", "end": "B4"}
    validate_parser_unit_hierarchy(_flatten_segments(result.segments))


def test_json_parser_rejects_depth_and_node_budgets_without_echoing_input(tmp_path, monkeypatch):
    from app.services import import_parsing

    path = tmp_path / "too-deep.json"
    path.write_text(json.dumps({"secret": {"nested": "value"}}), encoding="utf-8")
    monkeypatch.setattr(import_parsing, "MAX_JSON_DEPTH", 1)

    with pytest.raises(ImportResourceLimitError) as exc_info:
        parse_import_file(path, _library())

    assert str(exc_info.value) == import_parsing.RESOURCE_LIMIT_ERROR
    assert "secret" not in str(exc_info.value)

    monkeypatch.setattr(import_parsing, "MAX_JSON_DEPTH", 64)
    monkeypatch.setattr(import_parsing, "MAX_JSON_NODES", 2)
    with pytest.raises(ImportResourceLimitError):
        parse_import_file(path, _library())


def test_json_parser_rejects_input_bytes_before_json_load(tmp_path, monkeypatch):
    from app.services import import_parsing

    path = tmp_path / "too-large-input.json"
    payload = b'{"secret":"private"}'
    path.write_bytes(payload)
    monkeypatch.setattr(import_parsing, "MAX_JSON_INPUT_BYTES", len(payload) - 1)
    load_calls = 0

    def fail_json_load(*_args, **_kwargs):
        nonlocal load_calls
        load_calls += 1
        raise AssertionError("json.load must not run")

    monkeypatch.setattr(import_parsing.json, "load", fail_json_load)
    with pytest.raises(ImportResourceLimitError) as exc_info:
        parse_import_file(path, _library())

    assert load_calls == 0
    assert str(exc_info.value) == import_parsing.RESOURCE_LIMIT_ERROR
    assert "secret" not in str(exc_info.value)


def test_json_parser_rejects_normalized_text_budget_without_partial_units(tmp_path, monkeypatch):
    from app.services import import_parsing

    path = tmp_path / "too-large.json"
    path.write_text(json.dumps({"secret": "classified value"}), encoding="utf-8")
    monkeypatch.setattr(import_parsing, "MAX_NORMALIZED_TEXT_CHARS", 8)

    with pytest.raises(ImportResourceLimitError) as exc_info:
        parse_import_file(path, _library())

    assert str(exc_info.value) == import_parsing.RESOURCE_LIMIT_ERROR
    assert "classified" not in str(exc_info.value)


def test_csv_parser_rejects_row_column_cell_and_text_budgets(tmp_path, monkeypatch):
    from app.services import import_parsing

    path = tmp_path / "too-wide.csv"
    path.write_text("a,b,c\nsecret,2,3\n", encoding="utf-8")
    monkeypatch.setattr(import_parsing, "MAX_CSV_COLUMNS", 2)

    with pytest.raises(ImportResourceLimitError):
        parse_import_file(path, _library())

    monkeypatch.setattr(import_parsing, "MAX_CSV_COLUMNS", 1024)
    monkeypatch.setattr(import_parsing, "MAX_CSV_CELLS", 3)
    with pytest.raises(ImportResourceLimitError):
        parse_import_file(path, _library())

    monkeypatch.setattr(import_parsing, "MAX_CSV_CELLS", 1_000_000)
    monkeypatch.setattr(import_parsing, "MAX_CSV_ROWS", 1)
    with pytest.raises(ImportResourceLimitError):
        parse_import_file(path, _library())

    monkeypatch.setattr(import_parsing, "MAX_CSV_ROWS", 100_000)
    monkeypatch.setattr(import_parsing, "MAX_NORMALIZED_TEXT_CHARS", 4)
    with pytest.raises(ImportResourceLimitError):
        parse_import_file(path, _library())


def test_csv_parser_rejects_input_bytes_before_csv_reader(tmp_path, monkeypatch):
    from app.services import import_parsing

    path = tmp_path / "too-large-input.csv"
    payload = b"text\nsecret private value\n"
    path.write_bytes(payload)
    monkeypatch.setattr(import_parsing, "MAX_CSV_INPUT_BYTES", len(payload) - 1)
    reader_calls = 0

    def fail_csv_reader(*_args, **_kwargs):
        nonlocal reader_calls
        reader_calls += 1
        raise AssertionError("csv.reader must not run")

    monkeypatch.setattr(import_parsing.csv, "reader", fail_csv_reader)
    with pytest.raises(ImportResourceLimitError) as exc_info:
        parse_import_file(path, _library())

    assert reader_calls == 0
    assert str(exc_info.value) == import_parsing.RESOURCE_LIMIT_ERROR
    assert "secret" not in str(exc_info.value)


def test_docx_table_units_form_table_row_cell_tree_and_deduplicate_merged_values():
    document = docx.Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(0, 2).text = "Kind"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "1"
    table.cell(1, 2).text = "x"
    table.cell(2, 0).text = "B"
    table.cell(2, 1).text = "2"
    table.cell(2, 2).text = "y"
    merged = table.cell(2, 0).merge(table.cell(2, 1))
    merged.text = "B"
    output = io.BytesIO()
    document.save(output)

    segments = extract_docx_segments(output.getvalue())
    table_segment = next(segment for segment in segments if segment["kind"] == "table")
    units = table_segment["structured_units"]

    assert [unit["unit_kind"] for unit in units] == [
        "row", "cell", "cell", "cell",
        "row", "cell", "cell", "cell",
        "row", "cell", "cell",
    ]
    assert units[0]["parent_key"] == table_segment["unit_key"]
    assert units[1]["parent_key"] == units[0]["unit_key"]
    assert units[1]["source"]["table"] == {"index": 0}
    assert units[1]["source"]["row"] == {"start": 1, "end": 1}
    assert units[1]["source"]["cell"] == {"start": "A1", "end": "A1"}
    cell_values = [unit["value"] for unit in units if unit["unit_kind"] == "cell"]
    assert cell_values[-2:] == ["B", "y"]
    assert cell_values.count("B") == 1
    validate_parser_unit_hierarchy(_flatten_segments(segments))

    source = splitter.build_structured_source_from_segments(segments, chunk_size=200, chunk_overlap=0)
    assert source["normalized_text"] == "Name | Value | Kind\nA | 1 | x\nB | y"
    assert source["chunks"]


def test_same_heading_segments_get_distinct_keys():
    parser = parser_provenance("fixture", "v1")
    source = splitter.build_structured_source_from_segments(
        [
            {"kind": "prose", "text": "first", "heading": "Same", "source_kind": "docx", "parser": parser},
            {"kind": "prose", "text": "second", "heading": "Same", "source_kind": "docx", "parser": parser},
        ],
        chunk_size=80,
        chunk_overlap=0,
    )
    keys = [segment["parser_unit"]["unit_key"] for segment in source["segments"]]
    assert keys == ["docx:segment:0", "docx:segment:1"]
    validate_parser_unit_hierarchy(_flatten_segments(source["segments"]))


def test_malformed_json_and_direct_image_upload_are_not_silently_supported(tmp_path):
    bad_json = tmp_path / "bad.json"
    bad_json.write_text("{not-json", encoding="utf-8")
    with pytest.raises(json.JSONDecodeError):
        parse_import_file(bad_json, _library())

    image = tmp_path / "image.png"
    image.write_bytes(b"not-an-image")
    with pytest.raises(ValueError, match="unsupported file type"):
        parse_import_file(image, _library())


def test_parse_import_file_uses_structured_ocr_callback_without_changing_chunks(monkeypatch, tmp_path):
    from app.services import ocr as ocr_service

    pdf_path = tmp_path / "scan.pdf"
    pdf_path.write_bytes(b"pdf")
    page = types.SimpleNamespace(
        extract_text=lambda: "native text with enough characters for the fast path",
        images=[types.SimpleNamespace(data=b"embedded")],
    )
    monkeypatch.setattr(pdf_extract, "_open_reader", lambda _data: types.SimpleNamespace(pages=[page]))
    monkeypatch.setattr(ocr_service, "is_available", lambda: True)
    monkeypatch.setattr(
        ocr_service,
        "ocr_image_blocks",
        lambda _data: [{"text": "ocr text", "confidence": 0.7, "bbox": [1, 2, 3, 4]}],
    )

    result = parse_import_file(pdf_path, _library(ocr_enabled=True))

    assert result.chunks
    assert all("parser_unit" not in chunk for chunk in result.chunks)
    assert result.segments[0]["quality"]["ocr_blocks"][0]["confidence"] == 0.7
