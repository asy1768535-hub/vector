from __future__ import annotations

import asyncio
import hashlib
import inspect
import os
import uuid
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import AsyncMock

import docx
import pytest
from pptx import Presentation

from app.models.embedding_job import EmbeddingJob
from app.models.document_revision import DocumentRevision
from app.schemas.documents import ImportJobRead
from app.services.import_parsing import parse_import_file
from app.services import import_parsing
from app.services import import_uploads as import_uploads_service
from app.services.import_uploads import (
    ImportUploadError,
    folder_path_for_job,
    job_projection,
    job_projections,
    normalize_relative_path,
    source_path_for_job,
)
from app.services.object_storage_local import LocalObjectStorageAdapter
from app.services.revision_files import prepare_managed_file_path
from app.workers.importer import _apply_import_revision_scope
from app.workers import importer


def _library(**overrides):
    values = {
        "chunk_size": 200,
        "chunk_overlap": 20,
        "ocr_enabled": False,
        "docx_table_aware": False,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def test_background_import_parser_remains_synchronous():
    assert inspect.iscoroutinefunction(parse_import_file) is False


@pytest.mark.parametrize("stage", ["converting", "conversion_ready"])
def test_import_job_contract_exposes_doc_conversion_stages(stage):
    payload = _processing_import_job(current_stage=stage)

    assert ImportJobRead.model_validate(payload).current_stage == stage


class _Scalars:
    def __init__(self, value):
        self.value = value

    def first(self):
        return self.value

    def all(self):
        if self.value is None:
            return []
        return self.value if isinstance(self.value, list) else [self.value]


class _Result:
    def __init__(self, value):
        self.value = value

    def scalars(self):
        return _Scalars(self.value)


class _ProjectionDb:
    def __init__(self, *, embedding_job, graph_job=None):
        self.embedding_job = embedding_job
        self.graph_job = graph_job
        self.execute_calls = 0

    async def get(self, model, object_id):
        if model is EmbeddingJob and object_id == self.embedding_job.id:
            return self.embedding_job
        return None

    async def execute(self, statement):
        self.execute_calls += 1
        return _Result(self.graph_job)


class _BatchProjectionDb:
    def __init__(self, *results):
        self.results = list(results)
        self.execute_calls = 0

    async def execute(self, statement):
        self.execute_calls += 1
        return _Result(self.results.pop(0))


class _ScopeDb:
    def __init__(self, *, revision):
        self.revision = revision

    async def get(self, model, object_id):
        if model is DocumentRevision and object_id == self.revision.id:
            return self.revision
        return None


def _processing_import_job(**overrides):
    now = datetime.now(timezone.utc)
    values = {
        "id": uuid.uuid4(),
        "library_id": uuid.uuid4(),
        "batch_id": uuid.uuid4(),
        "file_name": "sample.md",
        "relative_path": None,
        "size_bytes": 12,
        "upload_offset": 12,
        "status": "processing",
        "current_stage": "embedding",
        "attempt_count": 1,
        "last_error": None,
        "result_operation": "created",
        "document_id": uuid.uuid4(),
        "document_revision_id": uuid.uuid4(),
        "embedding_job_id": uuid.uuid4(),
        "created_at": now,
        "updated_at": now,
        "upload_completed_at": now,
        "claimed_at": now,
        "finished_at": None,
        "graph_extraction_requested": True,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def test_job_projections_loads_related_embedding_jobs_once_for_a_visible_batch():
    first = _processing_import_job(graph_extraction_requested=False)
    second = _processing_import_job(graph_extraction_requested=False)
    now = datetime.now(timezone.utc)
    db = _BatchProjectionDb(
        [
            SimpleNamespace(id=first.embedding_job_id, status="done", finished_at=now),
            SimpleNamespace(id=second.embedding_job_id, status="done", finished_at=now),
        ]
    )

    projections = asyncio.run(job_projections(db, [first, second]))

    assert [(row["status"], row["current_stage"]) for row in projections] == [
        ("succeeded", "completed"),
        ("succeeded", "completed"),
    ]
    assert db.execute_calls == 1


def test_relative_path_preserves_nested_folder_hierarchy():
    relative = normalize_relative_path(
        "项目甲\\合同\\主合同.md",
        "主合同.md",
    )
    job = SimpleNamespace(relative_path=relative)

    assert relative == "项目甲/合同/主合同.md"
    assert source_path_for_job(job) == "/项目甲/合同/主合同.md"
    assert folder_path_for_job(job) == "/项目甲/合同"


def test_job_projection_reports_schema_wait_when_graph_job_is_not_created_yet():
    now = datetime.now(timezone.utc)
    job = _processing_import_job()
    embedding_job = SimpleNamespace(
        id=job.embedding_job_id,
        status="done",
        finished_at=now,
        last_error=None,
    )
    db = _ProjectionDb(embedding_job=embedding_job)

    projection = asyncio.run(job_projection(db, job))

    assert projection["status"] == "processing"
    assert projection["current_stage"] == "graph"
    assert projection["schema_discovery_state"] == "waiting_schema"
    assert ImportJobRead.model_validate(projection).schema_discovery_state == (
        "waiting_schema"
    )


@pytest.mark.parametrize("graph_status", ["queued", "processing", "waiting_schema"])
def test_job_projection_keeps_processing_when_graph_job_already_exists(graph_status):
    now = datetime.now(timezone.utc)
    job = _processing_import_job()
    embedding_job = SimpleNamespace(
        id=job.embedding_job_id,
        status="done",
        finished_at=now,
        last_error=None,
    )
    graph_job = SimpleNamespace(
        id=uuid.uuid4(),
        status=graph_status,
        finished_at=None,
        error_message=None,
    )
    db = _ProjectionDb(embedding_job=embedding_job, graph_job=graph_job)

    projection = asyncio.run(job_projection(db, job))

    assert projection["status"] == "processing"
    assert projection["current_stage"] == "graph"


@pytest.mark.parametrize(
    ("graph_status", "expected_status"),
    [("processing", "processing"), ("succeeded", "succeeded")],
)
def test_failed_graph_import_projects_latest_rerun_state(
    graph_status,
    expected_status,
):
    now = datetime.now(timezone.utc)
    job = _processing_import_job(
        status="failed",
        current_stage="graph",
        last_error="graph extraction failed",
        finished_at=now,
    )
    embedding_job = SimpleNamespace(
        id=job.embedding_job_id,
        status="done",
        finished_at=now,
        last_error=None,
    )
    graph_job = SimpleNamespace(
        id=uuid.uuid4(),
        status=graph_status,
        finished_at=now if graph_status == "succeeded" else None,
        error_message=None,
    )
    db = _ProjectionDb(embedding_job=embedding_job, graph_job=graph_job)

    projection = asyncio.run(job_projection(db, job))

    assert projection["status"] == expected_status
    assert projection["current_stage"] == (
        "completed" if graph_status == "succeeded" else "graph"
    )
    assert projection["last_error"] is None
    assert projection["retry_target_type"] is None
    assert projection["retry_target_id"] is None


def test_partially_succeeded_graph_remains_visible_and_routes_graph_retry():
    now = datetime.now(timezone.utc)
    job = _processing_import_job()
    embedding_job = SimpleNamespace(
        id=job.embedding_job_id,
        status="done",
        finished_at=now,
        last_error=None,
    )
    graph_job = SimpleNamespace(
        id=uuid.uuid4(),
        status="partially_succeeded",
        finished_at=now,
        error_message="unit_failures",
    )
    db = _ProjectionDb(embedding_job=embedding_job, graph_job=graph_job)

    projection = asyncio.run(job_projection(db, job))

    assert projection["status"] == "failed"
    assert projection["current_stage"] == "graph"
    assert projection["last_error"] == "unit_failures"
    assert projection["retry_target_type"] == "graph"
    assert projection["retry_target_id"] == graph_job.id
    assert ImportJobRead.model_validate(projection).retry_target_type == "graph"


def test_non_graph_import_failure_is_not_hidden_by_projection():
    job = _processing_import_job(
        status="failed",
        current_stage="parsing",
        last_error="parse failed",
    )
    db = _ProjectionDb(embedding_job=SimpleNamespace(id=job.embedding_job_id))

    projection = asyncio.run(job_projection(db, job))

    assert projection["status"] == "failed"
    assert projection["current_stage"] == "parsing"
    assert projection["last_error"] == "parse failed"
    assert projection["retry_target_type"] == "import"
    assert projection["retry_target_id"] == job.id
    assert db.execute_calls == 0


def test_import_revision_scope_syncs_security_level_and_graph_request():
    revision_id = uuid.uuid4()
    job = SimpleNamespace(
        security_level="internal",
        graph_extraction_requested=True,
    )
    document = SimpleNamespace(security_level=None)
    revision = SimpleNamespace(
        id=revision_id,
        security_level=None,
        parser_config={"parser": "plain"},
    )

    asyncio.run(
        _apply_import_revision_scope(
            _ScopeDb(revision=revision),
            job=job,
            document=document,
            revision_id=revision_id,
        )
    )

    assert document.security_level == "internal"
    assert revision.security_level == "internal"
    assert revision.parser_config == {
        "parser": "plain",
        "graph_extraction_requested": True,
    }


@pytest.mark.parametrize(
    "relative_path",
    [
        "../secret.txt",
        "folder/../secret.txt",
        "/absolute/secret.txt",
        "C:/secret.txt",
        r"\\server\share\secret.txt",
    ],
)
def test_relative_path_rejects_traversal_and_absolute_paths(relative_path):
    with pytest.raises(ImportUploadError):
        normalize_relative_path(relative_path, "secret.txt")


def test_text_and_csv_are_parsed_from_paths(tmp_path):
    text_path = tmp_path / "notes.md"
    text_path.write_text("# 标题\n正文内容", encoding="utf-8")
    csv_path = tmp_path / "records.csv"
    csv_path.write_text("name,value\n甲,1\n乙,2\n", encoding="utf-8")

    text_result = parse_import_file(text_path, _library())
    csv_result = parse_import_file(csv_path, _library())

    assert "# 标题" in text_result.normalized_text
    assert text_result.chunks
    assert "甲 | 1" in csv_result.normalized_text
    assert csv_result.chunks


def test_common_text_formats_and_gb18030_are_parsed(tmp_path):
    html_path = tmp_path / "page.html"
    html_path.write_text(
        "<html><style>hidden</style><body><h1>Visible title</h1><script>bad()</script></body></html>",
        encoding="utf-8",
    )
    text_path = tmp_path / "legacy.txt"
    text_path.write_bytes("中文资料".encode("gb18030"))
    tsv_path = tmp_path / "records.tsv"
    tsv_path.write_text("name\tvalue\n甲\t1\n", encoding="utf-8")

    html_result = parse_import_file(html_path, _library())
    text_result = parse_import_file(text_path, _library())
    tsv_result = parse_import_file(tsv_path, _library())

    assert "Visible title" in html_result.normalized_text
    assert "hidden" not in html_result.normalized_text
    assert "bad()" not in html_result.normalized_text
    assert "中文资料" in text_result.normalized_text
    assert "甲 | 1" in tsv_result.normalized_text


def test_pptx_and_xls_route_to_document_parsers(tmp_path, monkeypatch):
    pptx_path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Project overview"
    presentation.save(pptx_path)

    pptx_result = parse_import_file(pptx_path, _library())
    assert "Project overview" in pptx_result.normalized_text

    xls_path = tmp_path / "legacy.xls"
    xls_path.write_bytes(b"legacy spreadsheet fixture")
    called = {}
    monkeypatch.setattr(
        "app.services.import_parsing.xlsx_extract.extract_xlsx_segments",
        lambda path, *, is_xls=False: called.update(path=path, is_xls=is_xls) or [
            {"kind": "table", "title": "Sheet1", "rows": ["A | B"]}
        ],
    )
    monkeypatch.setattr(
        "app.services.import_parsing.splitter.build_structured_source_from_segments",
        lambda segments, **_kwargs: {
            "normalized_text": "A | B",
            "chunks": [{"text": "A | B"}],
            "segments": segments,
        },
    )

    xls_result = parse_import_file(xls_path, _library())
    assert xls_result.normalized_text == "A | B"
    assert called == {"path": xls_path, "is_xls": True}


def test_append_content_offloads_fsync(tmp_path, monkeypatch):
    calls = []
    real_to_thread = asyncio.to_thread

    async def record_to_thread(function, *args):
        calls.append(function)
        return await real_to_thread(function, *args)

    monkeypatch.setattr(asyncio, "to_thread", record_to_thread)
    config = SimpleNamespace(
        import_staging_dir=str(tmp_path),
        import_upload_chunk_bytes=1024,
    )
    job = SimpleNamespace(
        status="uploading",
        upload_offset=0,
        size_bytes=4,
        staging_key="0" * 32 + ".upload",
    )

    async def body():
        yield b"data"

    db = AsyncMock()
    offset = asyncio.run(
        import_uploads_service.append_content(
            db,
            job=job,
            expected_offset=0,
            body=body(),
            config=config,
        )
    )

    assert offset == 4
    assert os.fsync in calls


def test_wait_for_staging_file_recovers_after_transient_miss(tmp_path, monkeypatch):
    source = tmp_path / "transient.upload"
    sleep_calls = 0

    async def make_file_available(_seconds):
        nonlocal sleep_calls
        sleep_calls += 1
        source.write_bytes(b"data")

    monkeypatch.setattr(importer, "STAGING_READY_RETRY_SECONDS", (1.0,))
    monkeypatch.setattr(importer.asyncio, "sleep", make_file_available)

    actual_size = asyncio.run(importer._wait_for_staging_file(source, 4))

    assert actual_size == 4
    assert sleep_calls == 1


def test_wait_for_staging_file_fails_without_leaking_path(tmp_path, monkeypatch):
    source = tmp_path / "private-name.upload"
    source.write_bytes(b"no")

    async def no_wait(_seconds):
        return None

    monkeypatch.setattr(importer, "STAGING_READY_RETRY_SECONDS", (1.0,))
    monkeypatch.setattr(importer.asyncio, "sleep", no_wait)

    with pytest.raises(
        RuntimeError,
        match=r"exists=true actual_size=2 expected_size=4",
    ) as exc_info:
        asyncio.run(importer._wait_for_staging_file(source, 4))

    assert str(source) not in str(exc_info.value)


def test_docx_is_parsed_directly_from_path(tmp_path):
    path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_paragraph("路径读取的 Word 正文")
    document.save(path)

    result = parse_import_file(path, _library())

    assert "路径读取的 Word 正文" in result.normalized_text
    assert result.chunks


def test_docx_staging_file_uses_original_file_name_for_type(tmp_path):
    path = tmp_path / "00000000000000000000000000000000.upload"
    document = docx.Document()
    document.add_paragraph("DOCX content from a chunked upload")
    document.save(path)

    result = parse_import_file(path, _library(), file_name="nested/sample.docx")

    assert "DOCX content from a chunked upload" in result.normalized_text
    assert result.chunks


def test_converted_doc_parser_units_keep_the_original_doc_source_kind(tmp_path):
    path = tmp_path / "converted.docx"
    document = docx.Document()
    document.add_paragraph("Converted legacy Word content")
    document.save(path)
    parsed = parse_import_file(path, _library(docx_table_aware=True))

    rebound = import_parsing.rebind_converted_doc(
        parsed,
        converter_version="LibreOffice 24.2.0",
    )

    assert rebound.segments
    for segment in rebound.segments:
        assert segment["source_kind"] == "doc"
        assert segment["parser"]["name"] == "libreoffice-doc-to-docx+python-docx"
        assert segment["parser_unit"]["source_kind"] == "doc"
        assert "page" not in segment["parser_unit"].get("source", {})
        for unit in segment.get("structured_units", []):
            assert unit["source_kind"] == "doc"


def test_converted_doc_provenance_preserves_each_downstream_parser():
    parsed = import_parsing.ParsedImport(
        normalized_text="image text",
        chunks=[{"text": "image text"}],
        splitter_name="docx",
        segments=[{
            "source_kind": "docx",
            "parser": {"name": "python-docx", "version": "v1"},
            "parser_unit": {
                "source_kind": "docx",
                "parser": {"name": "python-docx", "version": "v1"},
                "source": {},
            },
            "structured_units": [{
                "source_kind": "docx",
                "parser": {"name": "rapidocr", "version": "1.4"},
                "source": {"bbox": {"x_min": 0, "y_min": 0, "x_max": 1, "y_max": 1}},
            }],
        }],
    )

    rebound = import_parsing.rebind_converted_doc(
        parsed,
        converter_version="LibreOffice 24.2.0",
    )

    assert rebound.segments[0]["parser_unit"]["parser"]["name"].endswith("python-docx")
    assert rebound.segments[0]["structured_units"][0]["parser"]["name"].endswith("rapidocr")


def test_local_object_storage_put_file_streams_and_reuses_content(tmp_path):
    source = tmp_path / "source.bin"
    source.write_bytes(b"streamed-content" * 4096)
    adapter = LocalObjectStorageAdapter(
        root=tmp_path / "objects",
        endpoint_ref="primary",
        max_read_bytes=1024 * 1024,
    )

    first = asyncio.run(
        adapter.put_file("libraries/test/objects/item.bin", source, None)
    )
    second = asyncio.run(
        adapter.put_file("libraries/test/objects/item.bin", source, None)
    )

    assert first == second
    assert adapter.path_for_read("libraries/test/objects/item.bin").read_bytes() == (
        source.read_bytes()
    )


def test_prepare_managed_file_path_validates_expected_hash(tmp_path):
    source = tmp_path / "source.txt"
    source.write_text("large file source", encoding="utf-8")
    adapter = LocalObjectStorageAdapter(
        root=tmp_path / "objects",
        endpoint_ref="primary",
        max_read_bytes=1024,
    )
    digest = hashlib.sha256(source.read_bytes()).hexdigest()

    prepared = asyncio.run(
        prepare_managed_file_path(
            adapter=adapter,
            library_id=SimpleNamespace(hex="library"),
            file_name="source.txt",
            content_type="text/plain",
            source_path=source,
            expected_sha256=digest,
        )
    )

    assert prepared.sha256 == digest
    assert prepared.size_bytes == source.stat().st_size


def test_prepare_managed_file_path_rejects_changed_source(tmp_path):
    source = tmp_path / "source.txt"
    source.write_text("changed", encoding="utf-8")
    adapter = LocalObjectStorageAdapter(
        root=tmp_path / "objects",
        endpoint_ref="primary",
        max_read_bytes=1024,
    )

    with pytest.raises(RuntimeError):
        asyncio.run(
            prepare_managed_file_path(
                adapter=adapter,
                library_id=SimpleNamespace(hex="library"),
                file_name="source.txt",
                content_type="text/plain",
                source_path=source,
                expected_sha256="0" * 64,
            )
        )
