from __future__ import annotations

import asyncio
import hashlib
import json
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import docx
import openpyxl
import pytest

from app.config import settings
from app.models.chunk import Chunk
from app.models.document import Document
from app.models.document_block import DocumentBlock
from app.models.evidence_unit import EvidenceUnit
from app.models.library import Library
from app.services import ingest
from app.services import import_parsing
from app.services import revision_files
from app.services.parser_units import build_parser_unit, parser_provenance
from app.services.import_parsing import ParsedImport
from app.workers import importer


LIBRARY_ID = uuid.UUID("00000000-0000-0000-0000-0000000000aa")
RAW_HASH = "a" * 64
FILE_ID = uuid.UUID("00000000-0000-0000-0000-0000000000bb")


def _library() -> Library:
    return Library(
        id=LIBRARY_ID,
        slug="m2-segmented",
        name="M2 segmented",
        qdrant_collection="m2_segmented",
        embedding_model="fixture",
        embedding_dim=8,
        chunk_size=1000,
        chunk_overlap=0,
    )


def _db(added: list[object], *, first=None, count: int = 0):
    result = MagicMock()
    result.scalars.return_value.first.return_value = first
    result.scalar_one.return_value = count
    db = MagicMock()
    db.execute = AsyncMock(return_value=result)
    db.add = MagicMock(side_effect=added.append)
    db.add_all = MagicMock(side_effect=lambda values: added.extend(values))
    db.flush = AsyncMock()
    nested = MagicMock()
    nested.__aenter__ = AsyncMock(return_value=None)
    nested.__aexit__ = AsyncMock(return_value=False)
    db.begin_nested = MagicMock(return_value=nested)
    return db


def _segment(source_kind: str, *, source: dict | None = None, text: str = "value") -> dict:
    parser = parser_provenance(f"fixture-{source_kind}", "v1", {"fixture": True})
    root = build_parser_unit(
        source_kind=source_kind,
        unit_kind="table" if source_kind == "docx" else "section",
        ordinal=0,
        unit_key=f"{source_kind}:root",
        parser=parser,
        source=source,
        source_text=text,
        source_start=0,
        source_end=len(text),
    )
    return {
        "kind": "table" if source_kind == "docx" else "prose",
        "text": text,
        "source_kind": source_kind,
        "ordinal": 0,
        "unit_key": root["unit_key"],
        "parser": parser,
        "parser_unit": root,
    }


@pytest.mark.parametrize(
    ("source_kind", "source"),
    [
        ("text", {}),
        ("pdf", {"page": {"start": 1, "end": 1}}),
        ("docx", {"table": {"index": 0}}),
        ("xlsx", {"sheet": {"name": "Sheet1"}}),
        ("csv", {"row": {"start": 1, "end": 1}}),
        ("json", {"json_pointer": ""}),
    ],
)
def test_segmented_generation_dual_writes_verified_locator_for_supported_sources(
    monkeypatch, source_kind, source
):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    added: list[object] = []
    segment = _segment(source_kind, source=source)

    result = asyncio.run(
        ingest.ingest_text(
            db=_db(added), library=_library(), text="value", title="source.dat",
            external_id="m2-1", metadata={"owner": "fixture"}, splitter="text",
            created_by=None, chunks=[{"text": "value", "source_start": 0, "source_end": 5}],
            segments=[segment], file_name="source.dat", raw_file_sha256=RAW_HASH,
            document_revision_file_id=FILE_ID,
        )
    )

    _document, job, count, existing = result
    assert job.document_revision_id is not None
    assert count == 1 and existing is False
    blocks = [row for row in added if isinstance(row, DocumentBlock)]
    chunks = [row for row in added if isinstance(row, Chunk)]
    evidence = [row for row in added if isinstance(row, EvidenceUnit)]
    assert len(blocks) == 1
    assert len(chunks) == len(evidence) == 1

    block_locator = blocks[0].content["evidence_locator_v1"]
    chunk_locator = chunks[0].chunk_metadata["evidence_locator_v1"]
    evidence_locator = evidence[0].evidence_metadata["evidence_locator_v1"]
    for locator in (block_locator, chunk_locator, evidence_locator):
        assert locator["provenance_status"] == "verified"
        assert locator["document_revision_file_id"] == str(FILE_ID)
        assert locator["raw_file_sha256"] == RAW_HASH
        assert locator["normalized_content_hash"] == ingest._content_hash("value")
    assert block_locator["unit_id"] == str(blocks[0].id)
    assert chunk_locator["unit_id"] == str(chunks[0].id)
    assert evidence_locator["unit_id"] == str(evidence[0].id)
    assert chunks[0].block_id == blocks[0].id
    assert chunks[0].evidence_id == evidence[0].id
    assert evidence[0].document_block_id == blocks[0].id


def test_segmented_generation_creates_revision_root_for_cross_segment_chunk(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    first = _segment("text", text="one")
    second = _segment("text", text="two")
    second["parser_unit"]["unit_key"] = "text:second"
    second["unit_key"] = "text:second"
    second["parser_unit"]["source"]["text"] = {
        "start": 4,
        "end": 7,
        "ranges": [{"start": 4, "end": 7, "sha256": ""}],
    }
    # Rebuild the second source range through the parser helper to keep its hash exact.
    second["parser_unit"] = build_parser_unit(
        source_kind="text", unit_kind="section", ordinal=1,
        unit_key="text:second", parser=second["parser"], source_text="one\ntwo",
        source_start=4, source_end=7,
    )
    added: list[object] = []
    result = asyncio.run(
        ingest.ingest_text(
            db=_db(added), library=_library(), text="one\ntwo", title="cross.txt",
            external_id=None, metadata=None, splitter="text", created_by=None,
            chunks=[{"text": "one\ntwo", "source_start": 0, "source_end": 7}],
            segments=[first, second], file_name="cross.txt",
            raw_file_sha256=RAW_HASH, document_revision_file_id=FILE_ID,
        )
    )
    assert result[2] == 1
    blocks = [row for row in added if isinstance(row, DocumentBlock)]
    chunks = [row for row in added if isinstance(row, Chunk)]
    assert len(blocks) == 3
    revision_root = next(row for row in blocks if row.position == {"type": "revision_root"})
    assert chunks[0].block_id == revision_root.id


def test_invalid_parser_hierarchy_fails_before_document_mutation(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    segment = _segment("text")
    child = build_parser_unit(
        source_kind="text", unit_kind="structured_unit", ordinal=1,
        unit_key="text:child", parent_key="missing", parser=segment["parser"],
    )
    segment["structured_units"] = [child]
    added: list[object] = []

    with pytest.raises(ValueError, match="dangling"):
        asyncio.run(
            ingest.ingest_text(
                db=_db(added), library=_library(), text="value", title="bad.txt",
                external_id=None, metadata=None, splitter="text", created_by=None,
                chunks=[{"text": "value", "source_start": 0, "source_end": 5}],
                segments=[segment],
            )
        )
    assert added == []


def test_changed_segmented_reingest_keeps_old_revision_records(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    old_revision = uuid.uuid4()
    document = Document(
        id=uuid.uuid4(), library_id=LIBRARY_ID, external_id="replace-me",
        title="old", content_hash="c" * 64, current_revision=1,
        current_revision_id=old_revision, latest_revision_id=old_revision, status="ready",
    )
    segment = _segment("text", text="new")
    added: list[object] = []
    db = _db(added)
    db.execute = AsyncMock(side_effect=[MagicMock(scalar_one=MagicMock(return_value=1))])

    with pytest.MonkeyPatch.context() as local:
        local.setattr("app.services.cleanup.enqueue_delete_before_revision", AsyncMock())
        job, count, changed = asyncio.run(
            ingest.reingest_document(
                db=db, library=_library(), document=document, new_text="new",
                title="new", metadata=None, splitter="text", chunks=[{
                    "text": "new", "source_start": 0, "source_end": 3,
                }], segments=[segment], file_name="new.txt",
                raw_file_sha256=RAW_HASH, document_revision_file_id=FILE_ID,
            )
        )

    assert changed is True and count == 1 and job is not None
    assert document.current_revision == 2
    delete_statements = [str(call.args[0]).lower() for call in db.execute.await_args_list]
    assert not any("delete from chunks" in statement for statement in delete_statements)
    assert old_revision != job.document_revision_id


def test_pdf_page_identity_wins_when_chunk_contains_synthetic_page_header(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    normalized = "PAGE 1\nbody one\n\nPAGE 2\nbody two"
    parser = parser_provenance("fixture-pdf", "v1")
    first = {
        "text": "body one",
        "parser_unit": build_parser_unit(
            source_kind="pdf", unit_kind="section", ordinal=0, unit_key="pdf:page:1",
            parser=parser, source={"page": {"start": 1, "end": 1}},
            source_text=normalized, source_start=7, source_end=15,
        ),
    }
    second = {
        "text": "body two",
        "parser_unit": build_parser_unit(
            source_kind="pdf", unit_kind="section", ordinal=1, unit_key="pdf:page:2",
            parser=parser, source={"page": {"start": 2, "end": 2}},
            source_text=normalized, source_start=24, source_end=32,
        ),
    }
    added: list[object] = []
    result = asyncio.run(
        ingest.ingest_text(
            db=_db(added), library=_library(), text=normalized, title="pages.pdf",
            external_id=None, metadata=None, splitter="text", created_by=None,
            chunks=[{
                "text": "PAGE 1\nbody one", "source_start": 0, "source_end": 15,
                "location": {"type": "page", "page": 1},
            }],
            segments=[first, second], file_name="pages.pdf",
            raw_file_sha256=RAW_HASH, document_revision_file_id=FILE_ID,
        )
    )

    assert result[2] == 1
    blocks = [row for row in added if isinstance(row, DocumentBlock)]
    chunk = next(row for row in added if isinstance(row, Chunk))
    page_one = next(row for row in blocks if row.position.get("page") == {"start": 1, "end": 1})
    assert chunk.block_id == page_one.id


def test_cross_sheet_chunk_uses_coordinate_free_revision_root(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    first = _segment("xlsx", source={"sheet": {"name": "First"}}, text="one")
    second = _segment("xlsx", source={"sheet": {"name": "Second"}}, text="two")
    second["parser_unit"] = build_parser_unit(
        source_kind="xlsx", unit_kind="section", ordinal=1, unit_key="xlsx:second",
        parser=second["parser"], source={"sheet": {"name": "Second"}},
        source_text="one\ntwo", source_start=4, source_end=7,
    )
    second["unit_key"] = "xlsx:second"
    added: list[object] = []

    asyncio.run(
        ingest.ingest_text(
            db=_db(added), library=_library(), text="one\ntwo", title="book.xlsx",
            external_id=None, metadata=None, splitter="text", created_by=None,
            chunks=[{"text": "one\ntwo", "source_start": 0, "source_end": 7}],
            segments=[first, second], file_name="book.xlsx",
            raw_file_sha256=RAW_HASH, document_revision_file_id=FILE_ID,
        )
    )

    revision_root = next(
        row for row in added
        if isinstance(row, DocumentBlock) and row.position == {"type": "revision_root"}
    )
    locator = revision_root.content["evidence_locator_v1"]
    assert locator["source"] == {
        "kind": "xlsx",
        "file_name": "book.xlsx",
        "text": {
            "start": 0,
            "end": 7,
            "ranges": [{"start": 0, "end": 7, "sha256": ingest._content_hash("one\ntwo")}],
        },
    }


@pytest.mark.parametrize(
    "chunk",
    [
        {"text": "wrong", "source_start": 0, "source_end": 5},
        {"text": "value", "source_start": 1, "source_end": 6},
        {
            "text": "value", "source_start": 0, "source_end": 5,
            "source_span_hash": "0" * 16,
        },
    ],
)
def test_invalid_chunk_text_range_or_hash_has_zero_mutation(monkeypatch, chunk):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    added: list[object] = []

    with pytest.raises(ValueError, match="chunk"):
        asyncio.run(
            ingest.ingest_text(
                db=_db(added), library=_library(), text="value", title="bad.txt",
                external_id=None, metadata=None, splitter="text", created_by=None,
                chunks=[chunk], segments=[_segment("text")],
            )
        )
    assert added == []


def test_source_ranges_with_valid_hashes_reject_tampered_chunk_before_mutation(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    normalized = "alpha\nbeta"
    ranges = [
        {
            "start": 0,
            "end": 5,
            "hash": ingest._content_hash(normalized[0:5]),
        },
        {
            "start": 6,
            "end": 10,
            "hash": ingest._content_hash(normalized[6:10]),
        },
    ]
    added: list[object] = []

    with pytest.raises(ValueError, match="source ranges"):
        asyncio.run(
            ingest.ingest_text(
                db=_db(added), library=_library(), text=normalized,
                title="tampered.txt", external_id=None, metadata=None,
                splitter="text", created_by=None,
                chunks=[{
                    "text": "alpha\nBETA",
                    "source_start": 0,
                    "source_end": 10,
                    "source_ranges": ranges,
                    "source_span_hash": ingest._content_hash(normalized),
                }],
                segments=[_segment("text", text=normalized)],
            )
        )
    assert added == []


def test_long_docx_and_xlsx_tables_reconstruct_ranges_and_materialize(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    parse_library = SimpleNamespace(
        chunk_size=110, chunk_overlap=0, ocr_enabled=False, docx_table_aware=True,
    )

    document = docx.Document()
    table = document.add_table(rows=1, cols=3)
    for column, value in enumerate(("name", "value", "note")):
        table.cell(0, column).text = value
    for row_number in range(1, 36):
        values = (
            f"item-{row_number:02d}",
            f"value-{row_number:02d}",
            "x" * 42,
        )
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cells[column].text = value
    docx_path = tmp_path / "long-table.docx"
    document.save(docx_path)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["name", "value", "note"])
    for row_number in range(1, 36):
        sheet.append([
            f"item-{row_number:02d}",
            f"value-{row_number:02d}",
            "x" * 42,
        ])
    xlsx_path = tmp_path / "long-table.xlsx"
    workbook.save(xlsx_path)

    for title, path in (("long-table.docx", docx_path), ("long-table.xlsx", xlsx_path)):
        parsed = import_parsing.parse_import_file(path, parse_library)
        assert len(parsed.chunks) > 1
        assert any(len(chunk.get("source_ranges") or []) > 1 for chunk in parsed.chunks)
        for chunk in parsed.chunks:
            reconstructed = "".join(
                parsed.normalized_text[item["start"]:item["end"]]
                for item in chunk["source_ranges"]
            )
            assert reconstructed == chunk["text"]

        added: list[object] = []
        asyncio.run(
            ingest.ingest_text(
                db=_db(added), library=_library(), text=parsed.normalized_text,
                title=title, external_id=None, metadata=None, splitter="text",
                created_by=None, chunks=parsed.chunks, segments=parsed.segments,
                file_name=title, raw_file_sha256=RAW_HASH,
                document_revision_file_id=FILE_ID,
            )
        )
        materialized_chunks = [row for row in added if isinstance(row, Chunk)]
        assert len(materialized_chunks) == len(parsed.chunks)


def test_structured_block_content_preserves_bounded_parser_payload(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    parser = parser_provenance("fixture-xlsx", "v1")
    root = build_parser_unit(
        source_kind="xlsx", unit_kind="table", ordinal=0, unit_key="xlsx:sheet:0",
        parser=parser, source={"sheet": {"name": "Data"}},
        source_text="42", source_start=0, source_end=2,
    )
    cell = build_parser_unit(
        source_kind="xlsx", unit_kind="cell", ordinal=0, unit_key="xlsx:sheet:0:cell:A1",
        parser=parser, parent_key="xlsx:sheet:0",
        source={
            "sheet": {"name": "Data"}, "row": {"start": 1, "end": 1},
            "column": {"start": 1, "end": 1}, "cell": {"start": "A1", "end": "A1"},
        },
    )
    cell.update({"value": 42, "formula": "=1+1", "text": "42"})
    segment = {
        "kind": "table", "caption": "Data", "header": "Value", "text": "42",
        "parser_unit": root, "structured_units": [cell],
    }
    added: list[object] = []
    asyncio.run(
        ingest.ingest_text(
            db=_db(added), library=_library(), text="42", title="book.xlsx",
            external_id=None, metadata=None, splitter="text", created_by=None,
            chunks=[{"text": "42", "source_start": 0, "source_end": 2}],
            segments=[segment], file_name="book.xlsx",
            raw_file_sha256=RAW_HASH, document_revision_file_id=FILE_ID,
        )
    )

    cell_block = next(row for row in added if isinstance(row, DocumentBlock) and row.block_kind == "cell")
    structural = cell_block.content["parser_unit"]
    assert structural["value"] == 42
    assert structural["formula"] == "=1+1"
    assert structural["segment_caption"] == "Data"
    assert "rows" not in structural
    assert len(json.dumps(cell_block.content, ensure_ascii=False).encode("utf-8")) <= 8192


def test_json_pointer_blocks_keep_bounded_value_and_auditable_text(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    path = tmp_path / "values.json"
    path.write_text(json.dumps({"name": "Alice", "items": [1, 2]}), encoding="utf-8")
    parsed = import_parsing.parse_import_file(path, _library())
    added: list[object] = []

    asyncio.run(
        ingest.ingest_text(
            db=_db(added), library=_library(), text=parsed.normalized_text,
            title="values.json", external_id=None, metadata=None, splitter="text",
            created_by=None, chunks=parsed.chunks, segments=parsed.segments,
            file_name="values.json", raw_file_sha256=RAW_HASH,
            document_revision_file_id=FILE_ID,
        )
    )

    name_block = next(
        row for row in added
        if isinstance(row, DocumentBlock)
        and row.position.get("json_pointer") == "/name"
    )
    structural = name_block.content["parser_unit"]
    assert structural["json_value"] == "Alice"
    assert name_block.text == '"Alice"'


def test_csv_docx_xlsx_parser_values_round_trip_into_block_content(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    parse_library = SimpleNamespace(
        chunk_size=200, chunk_overlap=0, ocr_enabled=False, docx_table_aware=True,
    )

    csv_path = tmp_path / "rows.csv"
    csv_path.write_text("name,value\nAlice,42\n", encoding="utf-8")
    csv_parsed = import_parsing.parse_import_file(csv_path, parse_library)

    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "42"
    docx_path = tmp_path / "table.docx"
    document.save(docx_path)
    docx_parsed = import_parsing.parse_import_file(docx_path, parse_library)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Value"
    sheet["B1"] = "=1+1"
    xlsx_path = tmp_path / "formula.xlsx"
    workbook.save(xlsx_path)
    xlsx_parsed = import_parsing.parse_import_file(xlsx_path, parse_library)

    for title, parsed, expected in (
        ("rows.csv", csv_parsed, ("Alice", "cell", "A2")),
        ("table.docx", docx_parsed, ("Alice", "cell", "A2")),
        ("formula.xlsx", xlsx_parsed, ("=1+1", "cell", "B1")),
    ):
        added: list[object] = []
        asyncio.run(
            ingest.ingest_text(
                db=_db(added), library=_library(), text=parsed.normalized_text,
                title=title, external_id=None, metadata=None, splitter="text",
                created_by=None, chunks=parsed.chunks, segments=parsed.segments,
                file_name=title, raw_file_sha256=RAW_HASH,
                document_revision_file_id=FILE_ID,
            )
        )
        block = next(
            row for row in added
            if isinstance(row, DocumentBlock)
            and row.position.get("cell", {}).get("start") == expected[2]
        )
        structural = block.content["parser_unit"]
        if title == "formula.xlsx":
            assert structural["formula"] == expected[0]
        else:
            assert structural["value"] == expected[0]
        assert structural["unit_kind"] == expected[1]


def test_unchanged_segmented_reingest_does_not_mutate_ready_document(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    old_revision = uuid.uuid4()
    document = Document(
        id=uuid.uuid4(), library_id=LIBRARY_ID, external_id="same",
        title="ready.txt", content_hash=ingest._content_hash("value"), current_revision=4,
        current_revision_id=old_revision, latest_revision_id=old_revision, status="ready",
        doc_metadata={"stable": True}, source_path="/ready.txt",
    )
    added: list[object] = []
    db = _db(added, count=1)
    result = asyncio.run(
        ingest.reingest_document(
            db=db, library=_library(), document=document, new_text="value",
            title="ready.txt", metadata={"stable": True}, splitter="text",
            chunks=[{"text": "value", "source_start": 0, "source_end": 5}],
            segments=[_segment("text")], file_name="ready.txt",
            raw_file_sha256="b" * 64, document_revision_file_id=uuid.uuid4(),
        )
    )

    assert result == (None, 1, False)
    assert document.current_revision == 4
    assert document.current_revision_id == old_revision
    assert document.latest_revision_id == old_revision
    assert document.status == "ready"
    assert added == []


class _WorkerSessionContext:
    def __init__(self, db):
        self.db = db

    async def __aenter__(self):
        return self.db

    async def __aexit__(self, exc_type, exc, traceback):
        return False


class _WorkerTransactionContext:
    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, traceback):
        return False


def test_importer_passes_segments_and_one_preallocated_file_id(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    monkeypatch.setattr(settings, "revision_file_storage_enabled", True)
    job_id = uuid.uuid4()
    revision_id = uuid.uuid4()
    job = SimpleNamespace(
        id=job_id,
        library_id=LIBRARY_ID,
        file_name="source.txt",
        content_type="text/plain",
        size_bytes=5,
        sha256=RAW_HASH,
        staging_key="staged/source.txt",
        external_id=None,
        requested_by_user_id=None,
        security_level=None,
        graph_extraction_requested=False,
        replace_document_id=None,
        status="processing",
        result_operation=None,
        document_id=None,
        document_revision_id=None,
        embedding_job_id=None,
        current_stage="parsing",
        worker_id="fixture-worker",
        claimed_at=None,
        finished_at=None,
    )
    library = _library()
    source = tmp_path / "source.txt"
    source.write_text("value", encoding="utf-8")
    segments = [_segment("text")]
    parsed = ParsedImport(
        normalized_text="value",
        chunks=[{"text": "value", "source_start": 0, "source_end": 5}],
        splitter_name="text",
        segments=segments,
    )
    document = SimpleNamespace(
        id=uuid.uuid4(),
        current_revision=1,
        current_revision_id=revision_id,
        latest_revision_id=revision_id,
        content_hash=ingest._content_hash("value"),
        source_path=None,
        folder_id=None,
    )
    embedding_job = SimpleNamespace(id=uuid.uuid4(), document_revision_id=revision_id)

    class _ReadDb:
        async def get(self, model, object_id):
            if model is importer.DocumentImportJob:
                return job
            if model is importer.Library:
                return library
            return None

    class _WriteResult:
        def scalars(self):
            return self

        def first(self):
            return job

    class _WriteDb:
        def begin(self):
            return _WorkerTransactionContext()

        async def execute(self, statement):
            return _WriteResult()

        async def get(self, model, object_id):
            if model is importer.Library:
                return library
            return None

    read_db = _ReadDb()
    write_db = _WriteDb()
    sessions = iter([read_db, write_db])
    monkeypatch.setattr(
        importer,
        "async_session_factory",
        lambda: _WorkerSessionContext(next(sessions)),
    )
    monkeypatch.setattr(importer, "staging_path", lambda _key: source)
    monkeypatch.setattr(importer, "source_path_for_job", lambda _job: None)
    monkeypatch.setattr(importer, "folder_path_for_job", lambda _job: None)
    monkeypatch.setattr(importer, "parse_import_file", lambda *args, **kwargs: parsed)
    monkeypatch.setattr(importer, "_set_stage", AsyncMock())
    monkeypatch.setattr(importer, "_apply_import_revision_scope", AsyncMock())
    monkeypatch.setattr(importer.folders_service, "ensure_folder_path", AsyncMock(return_value=None))
    ingest_call = AsyncMock(return_value=(document, embedding_job, 1, False))
    monkeypatch.setattr(importer.ingest_service, "ingest_text", ingest_call)
    prepared = SimpleNamespace()
    monkeypatch.setattr(importer, "_prepare_revision_file", AsyncMock(return_value=prepared))
    persist_call = AsyncMock(return_value=SimpleNamespace(id=uuid.uuid4()))
    monkeypatch.setattr(importer, "_persist_revision_file", persist_call)

    asyncio.run(importer._process_claimed_job(job_id))

    assert ingest_call.await_args.kwargs["segments"] is segments
    assert persist_call.await_args.kwargs["file_id"] is not None
    assert persist_call.await_args.kwargs["revision_id"] == revision_id
    assert persist_call.await_args.kwargs["job"] is job
    assert persist_call.await_args.kwargs["prepared"] is prepared
    assert job.status == "processing"
    assert job.document_revision_id == revision_id


def test_doc_importer_parses_converted_docx_but_persists_original_doc(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "enable_evidence_write_path", True)
    monkeypatch.setattr(settings, "revision_file_storage_enabled", True)
    job_id = uuid.uuid4()
    revision_id = uuid.uuid4()
    original = tmp_path / "original.upload"
    original.write_bytes(b"original-doc")
    converted = tmp_path / "original.converted.docx"
    converted.write_bytes(b"converted-docx")
    conversion_sha256 = hashlib.sha256(converted.read_bytes()).hexdigest()
    job = SimpleNamespace(
        id=job_id, library_id=LIBRARY_ID, file_name="legacy.doc",
        content_type="application/msword", size_bytes=original.stat().st_size,
        sha256=RAW_HASH, staging_key="0" * 32 + ".upload", external_id=None,
        requested_by_user_id=None, security_level=None, graph_extraction_requested=False,
        replace_document_id=None, status="processing", result_operation=None,
        document_id=None, document_revision_id=None, embedding_job_id=None,
        current_stage="parsing", worker_id="fixture-worker", claimed_at=None,
        finished_at=None, conversion_sha256=conversion_sha256,
        converter_version="LibreOffice 24.2.0",
    )
    library = _library()
    parsed = ParsedImport(
        normalized_text="value",
        chunks=[{"text": "value", "source_start": 0, "source_end": 5}],
        splitter_name="docx",
        segments=[_segment("docx")],
    )
    document = SimpleNamespace(
        id=uuid.uuid4(), current_revision=1, current_revision_id=revision_id,
        latest_revision_id=revision_id, content_hash=ingest._content_hash("value"),
        source_path=None, folder_id=None,
    )
    embedding_job = SimpleNamespace(id=uuid.uuid4(), document_revision_id=revision_id)

    class _ReadDb:
        async def get(self, model, object_id):
            if model is importer.DocumentImportJob:
                return job
            if model is importer.Library:
                return library
            return None

    class _WriteResult:
        def scalars(self):
            return self

        def first(self):
            return job

    class _WriteDb:
        def begin(self):
            return _WorkerTransactionContext()

        async def execute(self, statement):
            return _WriteResult()

        async def get(self, model, object_id):
            if model is importer.Library:
                return library
            return None

    sessions = iter([_ReadDb(), _WriteDb()])
    monkeypatch.setattr(importer, "async_session_factory", lambda: _WorkerSessionContext(next(sessions)))
    monkeypatch.setattr(importer, "staging_path", lambda _key: original)
    monkeypatch.setattr(importer, "converted_staging_path", lambda _key: converted, raising=False)
    monkeypatch.setattr(importer, "source_path_for_job", lambda _job: None)
    monkeypatch.setattr(importer, "folder_path_for_job", lambda _job: None)
    parse_call = MagicMock(return_value=parsed)
    monkeypatch.setattr(importer, "parse_import_file", parse_call)
    monkeypatch.setattr(importer, "_set_stage", AsyncMock())
    monkeypatch.setattr(importer, "_apply_import_revision_scope", AsyncMock())
    monkeypatch.setattr(importer.folders_service, "ensure_folder_path", AsyncMock(return_value=None))
    ingest_call = AsyncMock(return_value=(document, embedding_job, 1, False))
    monkeypatch.setattr(importer.ingest_service, "ingest_text", ingest_call)
    prepared = SimpleNamespace()
    monkeypatch.setattr(importer, "_prepare_revision_file", AsyncMock(return_value=prepared))
    persist_call = AsyncMock(return_value=SimpleNamespace(id=uuid.uuid4()))
    monkeypatch.setattr(importer, "_persist_revision_file", persist_call)

    asyncio.run(importer._process_claimed_job(job_id))

    assert parse_call.call_args.args[0] == converted
    assert parse_call.call_args.kwargs.get("file_name") is None
    assert ingest_call.await_args.kwargs["raw_file_sha256"] == RAW_HASH
    assert ingest_call.await_args.kwargs["segments"][0]["source_kind"] == "doc"
    assert persist_call.await_args.kwargs["prepared"] is prepared


def test_persist_revision_file_binds_preallocated_id(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "revision_file_storage_enabled", True)
    file_id = uuid.uuid4()
    job = SimpleNamespace(
        library_id=LIBRARY_ID,
        file_name="source.txt",
        content_type="text/plain",
        sha256=RAW_HASH,
    )
    document = SimpleNamespace(id=uuid.uuid4())
    prepared = SimpleNamespace(
        library_id=LIBRARY_ID,
        file_name="source.txt",
        content_type="text/plain",
        size_bytes=5,
        sha256=RAW_HASH,
        locator=SimpleNamespace(),
        managed_snapshot=True,
        source_locator=None,
        verified_at=None,
    )
    persist_call = AsyncMock(return_value=SimpleNamespace(id=file_id))
    monkeypatch.setattr(importer, "persist_revision_file_capture", persist_call)
    monkeypatch.setattr(importer, "bind_prepared_file_object", revision_files.bind_prepared_file_object)

    result = asyncio.run(
        importer._persist_revision_file(
            MagicMock(),
            job=job,
            document=document,
            prepared=prepared,
            revision_id=uuid.uuid4(),
            file_id=file_id,
        )
    )

    assert result.id == file_id
    assert persist_call.await_args.kwargs["prepared"].file_id == file_id
