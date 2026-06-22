"""docx 抽取单测：表格内容应被抽取，且与段落按文档顺序拼接。"""
from __future__ import annotations

import io

import docx

from app.services.docx_extract import extract_docx_text


def _build_docx() -> bytes:
    d = docx.Document()
    d.add_paragraph("引言段落")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "技术类别"
    t.rows[0].cells[1].text = "拟选方案"
    t.rows[1].cells[0].text = "数据访问"
    t.rows[1].cells[1].text = "MyBatis-Plus、MySQL"
    d.add_paragraph("结尾段落")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_extract_includes_paragraphs_and_table():
    text = extract_docx_text(_build_docx())
    assert "引言段落" in text and "结尾段落" in text
    assert "MyBatis-Plus" in text          # 表格内容被抽到（原先会丢）
    assert "技术类别 | 拟选方案" in text     # 行内单元格用 | 拼接


def test_extract_preserves_document_order():
    text = extract_docx_text(_build_docx())
    # 段落—表格—段落 的顺序应保留
    assert text.index("引言段落") < text.index("MyBatis-Plus") < text.index("结尾段落")


def _build_docx_with_image() -> bytes:
    from PIL import Image  # 随 rapidocr 装入；无则跳过该测试

    d = docx.Document()
    d.add_paragraph("图前段落")
    img = io.BytesIO()
    Image.new("RGB", (40, 20), (255, 255, 255)).save(img, format="PNG")
    img.seek(0)
    d.add_picture(img)
    d.add_paragraph("图后段落")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_extract_runs_ocr_on_images_in_order():
    import pytest
    pytest.importorskip("PIL")
    # 用假 ocr 回调（不依赖真引擎）：任何图片都返回固定标记
    text = extract_docx_text(_build_docx_with_image(), ocr=lambda _blob: "图片识别文字XYZ")
    assert "图片识别文字XYZ" in text
    # 顺序：图前段落 < 图片OCR文字 < 图后段落
    assert text.index("图前段落") < text.index("图片识别文字XYZ") < text.index("图后段落")


def test_no_ocr_callback_skips_images():
    import pytest
    pytest.importorskip("PIL")
    text = extract_docx_text(_build_docx_with_image())  # 不传 ocr
    assert "图前段落" in text and "图后段落" in text  # 图片被忽略，不报错


def test_segments_track_heading_and_split_table():
    d = docx.Document()
    d.add_heading("1 概述", level=1)
    d.add_paragraph("概述正文")
    d.add_heading("2 技术", level=1)
    d.add_paragraph("技术正文")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "类别"
    t.rows[0].cells[1].text = "方案"
    t.rows[1].cells[0].text = "数据库"
    t.rows[1].cells[1].text = "MySQL"
    buf = io.BytesIO()
    d.save(buf)

    from app.services.docx_extract import extract_docx_segments
    segs = extract_docx_segments(buf.getvalue())
    kinds = [s["kind"] for s in segs]
    assert "prose" in kinds and "table" in kinds
    tbl = next(s for s in segs if s["kind"] == "table")
    assert "2 技术" in tbl["heading"]           # 表归属第 2 节
    assert tbl["header"] == "类别 | 方案"
    # 第二段散文挂在第 2 节标题下
    prose2 = next(s for s in segs if s["kind"] == "prose" and "技术正文" in s["text"])
    assert "2 技术" in prose2["heading"]


def test_table_dedups_merged_cells():
    d = docx.Document()
    row = d.add_table(rows=1, cols=3).rows[0]
    row.cells[0].text = "A"
    row.cells[1].text = "A"   # 模拟合并单元格重复文本
    row.cells[2].text = "B"
    buf = io.BytesIO()
    d.save(buf)
    text = extract_docx_text(buf.getvalue())
    assert text == "A | B"    # 相邻重复被折叠
