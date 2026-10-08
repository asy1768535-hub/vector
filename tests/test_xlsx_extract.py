"""xlsx 摄入单测：工作表 → 表格 segment。"""
from __future__ import annotations

import datetime
import io
import zipfile
from types import SimpleNamespace
from unittest.mock import patch
from xml.etree import ElementTree

import openpyxl
import pytest
from openpyxl.worksheet import _reader as openpyxl_reader

from app.config import settings
from app.services import xlsx_extract
from app.services.xlsx_extract import XlsxLimitError, extract_xlsx_segments


_NATIVE_NUMBER_CAST = openpyxl_reader._cast_number


def _xlsx_bytes(*rows):
    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fake_zip(infos):
    class FakeZip:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def infolist(self):
            return infos

    return FakeZip()


def test_xlsx_sheet_becomes_table_segment():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "技术选型"
    ws.append(["类别", "方案"])
    ws.append(["数据库", "MySQL"])
    ws.append(["缓存", "Redis"])
    buf = io.BytesIO()
    wb.save(buf)

    segs = extract_xlsx_segments(buf.getvalue())
    assert len(segs) == 1
    s = segs[0]
    assert s["kind"] == "table"
    assert s["heading"] == "技术选型"          # 工作表名作章节/表名
    assert s["header"] == "类别 | 方案"         # 首行作表头
    assert "数据库 | MySQL" in s["rows"]
    assert "缓存 | Redis" in s["rows"]
    assert s["row_numbers"] == [1, 2, 3]
    assert s["location"] == {"type": "sheet", "sheet": "技术选型"}


def test_xlsx_empty_sheet_skipped():
    wb = openpyxl.Workbook()
    wb.active.title = "空表"
    buf = io.BytesIO()
    wb.save(buf)
    assert extract_xlsx_segments(buf.getvalue()) == []


def test_xlsx_staging_path_does_not_require_xlsx_suffix(tmp_path):
    path = tmp_path / "00000000000000000000000000000000.upload"
    wb = openpyxl.Workbook()
    wb.active.append(["name", "value"])
    wb.active.append(["asset", "42"])
    wb.save(path)

    segments = extract_xlsx_segments(path)

    assert segments[0]["rows"] == ["name | value", "asset | 42"]


def test_xls_rejects_oversized_input_before_reading_parser(tmp_path, monkeypatch):
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"toolarge")
    monkeypatch.setattr(settings, "xls_max_input_bytes", 3)

    with pytest.raises(XlsxLimitError, match="XLS input exceeds"):
        extract_xlsx_segments(path, is_xls=True)


def test_xlsx_rejects_zip_entry_count_before_openpyxl(monkeypatch):
    monkeypatch.setattr(settings, "xlsx_max_zip_entries", 1)
    archive = _fake_zip([SimpleNamespace(file_size=1, compress_size=1)] * 2)

    with patch.object(xlsx_extract.zipfile, "ZipFile", return_value=archive):
        with pytest.raises(XlsxLimitError, match="entry limit"):
            xlsx_extract._validate_xlsx_archive(b"fixture")


@pytest.mark.parametrize(
    ("info", "attribute", "limit", "message"),
    [
        (SimpleNamespace(file_size=1, compress_size=11), "xlsx_max_zip_compressed_bytes", 10, "compressed-size"),
        (SimpleNamespace(file_size=11, compress_size=1), "xlsx_max_zip_uncompressed_bytes", 10, "uncompressed-size"),
        (SimpleNamespace(file_size=11, compress_size=1), "xlsx_max_zip_entry_bytes", 10, "entry exceeds"),
        (SimpleNamespace(file_size=11, compress_size=1), "xlsx_max_zip_compression_ratio", 10, "compression-ratio"),
    ],
)
def test_xlsx_rejects_zip_bomb_metadata(monkeypatch, info, attribute, limit, message):
    monkeypatch.setattr(settings, attribute, limit)
    archive = _fake_zip([info])

    with patch.object(xlsx_extract.zipfile, "ZipFile", return_value=archive):
        with pytest.raises(XlsxLimitError, match=message):
            xlsx_extract._validate_xlsx_archive(b"fixture")


def test_xlsx_rejects_declared_sheet_dimensions_without_large_fixture(monkeypatch):
    monkeypatch.setattr(settings, "xlsx_max_rows", 1)

    with pytest.raises(XlsxLimitError, match="row limit"):
        extract_xlsx_segments(_xlsx_bytes(["header"], ["value"]))


def test_xlsx_rejects_sheet_count_without_large_fixture(monkeypatch):
    monkeypatch.setattr(settings, "xlsx_max_sheets", 1)
    wb = openpyxl.Workbook()
    wb.create_sheet("second")
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(XlsxLimitError, match="sheet limit"):
        extract_xlsx_segments(buf.getvalue())


def test_xlsx_rejects_cell_and_text_limits_without_large_fixture(monkeypatch):
    monkeypatch.setattr(settings, "xlsx_max_cells", 1)
    with pytest.raises(XlsxLimitError, match="cell limit"):
        extract_xlsx_segments(_xlsx_bytes(["header", "second"]))

    monkeypatch.setattr(settings, "xlsx_max_cells", 1_000_000)
    monkeypatch.setattr(settings, "xlsx_max_cell_text_chars", 3)
    with pytest.raises(XlsxLimitError, match="cell text"):
        extract_xlsx_segments(_xlsx_bytes(["secret-content"]))


def test_xlsx_ignores_phantom_styled_columns_and_consecutive_empty_rows():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "StyledSheet"
    ws.append(["Name", "Age"])
    ws.append(["Alice", "30"])
    # Create phantom column width up to column 16384 (XFD)
    ws.cell(row=1, column=16384, value=None)
    buf = io.BytesIO()
    wb.save(buf)

    segs = extract_xlsx_segments(buf.getvalue())
    assert len(segs) == 1
    assert segs[0]["rows"] == ["Name | Age", "Alice | 30"]

def test_xlsx_sparse_far_column_compresses_delimiters_and_preserves_locator():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "花名册"
    ws.cell(row=1, column=1, value="序号")
    ws.cell(row=1, column=2, value="姓名")
    ws.cell(row=1, column=16381, value="备注")
    ws.cell(row=2, column=1, value="1")
    ws.cell(row=2, column=2, value="张三")
    ws.cell(row=2, column=16381, value="已转正")
    buf = io.BytesIO()
    wb.save(buf)

    segs = extract_xlsx_segments(buf.getvalue())
    assert len(segs) == 1
    s = segs[0]
    # Delimiter compression: no 16000+ " | " strings
    assert s["header"] == "序号 | 姓名 | [第16381列/XFA] 备注"
    assert s["rows"][1] == "1 | 张三 | [第16381列/XFA] 已转正"
    # Length is small and bounded
    assert len(s["header"]) < 100

    # Structured units preserve exact coordinates and cell names for EvidenceLocatorV1
    units = s["structured_units"]
    far_units = [u for u in units if u["source"]["column"]["start"] == 16381]
    assert len(far_units) == 2
    assert far_units[0]["unit_key"] == "xlsx:sheet:0:cell:XFA1"
    assert far_units[0]["value"] == "备注"
    assert far_units[1]["unit_key"] == "xlsx:sheet:0:cell:XFA2"
    assert far_units[1]["value"] == "已转正"


def test_xlsx_normal_empty_cells_retain_standard_delimiters():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "常规表"
    ws.append(["Col1", None, "Col3", None, None, "Col6"])
    buf = io.BytesIO()
    wb.save(buf)

    segs = extract_xlsx_segments(buf.getvalue())
    assert len(segs) == 1
    # Empty span <= 8 retains standard empty " |  | " format
    assert segs[0]["rows"][0] == "Col1 |  | Col3 |  |  | Col6"


def _xlsx_with_numeric_values(workbook, replacements):
    """Represent malformed exporter values in actual worksheet XML."""
    source = io.BytesIO()
    workbook.save(source)
    output = io.BytesIO()
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as modified:
        for entry in original.infolist():
            content = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                worksheet = ElementTree.fromstring(content)
                for cell in worksheet.iter(f"{namespace}c"):
                    coordinate = cell.get("r")
                    if coordinate in replacements:
                        # Numeric is the default cell type, including cached formula values.
                        cell.attrib.pop("t", None)
                        value = cell.find(f"{namespace}v")
                        if value is None:
                            value = ElementTree.SubElement(cell, f"{namespace}v")
                        value.text = replacements[coordinate]
                content = ElementTree.tostring(worksheet)
            modified.writestr(entry, content)
    return output.getvalue()


@pytest.mark.parametrize("raw_value", ["NULL", "invalid", "  NULL  ", "1x"])
@pytest.mark.parametrize("date_style", [False, True])
def test_xlsx_invalid_numeric_values_retain_raw_text_and_coordinates(raw_value, date_style):
    workbook = openpyxl.Workbook()
    workbook.active["C3"] = 1
    if date_style:
        workbook.active["C3"].number_format = "yyyy-mm-dd"
    data = _xlsx_with_numeric_values(workbook, {"C3": raw_value})

    segment = extract_xlsx_segments(data)[0]

    assert segment["rows"] == [raw_value.strip()]
    assert segment["row_numbers"] == [3]
    unit = segment["structured_units"][0]
    assert unit["value"] == raw_value
    assert unit["formula"] is None
    assert unit["source"]["cell"] == {"start": "C3", "end": "C3"}


def test_xlsx_invalid_numeric_formula_cache_retains_formula_and_raw_value():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "=1+1"
    workbook.active["A1"].number_format = "yyyy-mm-dd"
    data = _xlsx_with_numeric_values(workbook, {"A1": "NULL"})

    unit = extract_xlsx_segments(data)[0]["structured_units"][0]

    assert unit["value"] == "NULL"
    assert unit["formula"] == "=1+1"


def test_xlsx_numeric_compatibility_does_not_change_other_openpyxl_calls():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = 1
    data = _xlsx_with_numeric_values(workbook, {"A1": "NULL"})

    assert extract_xlsx_segments(data)[0]["rows"] == ["NULL"]
    assert openpyxl_reader._cast_number is _NATIVE_NUMBER_CAST
    native_workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    try:
        with pytest.raises(ValueError):
            list(native_workbook.active.iter_rows(values_only=True))
    finally:
        native_workbook.close()


def test_xlsx_numeric_compatibility_preserves_numbers_dates_and_formula_cache():
    workbook = openpyxl.Workbook()
    workbook.active.append([42, 1.25, datetime.date(2026, 10, 5), "=1+1"])
    data = _xlsx_with_numeric_values(workbook, {"D1": "2"})

    units = extract_xlsx_segments(data)[0]["structured_units"]

    assert [unit["value"] for unit in units] == [42, 1.25, datetime.datetime(2026, 10, 5), 2]
    assert [unit["formula"] for unit in units] == [None, None, None, "=1+1"]


@pytest.mark.parametrize("attribute,limit", [("xlsx_max_cell_text_chars", 3), ("xlsx_max_text_chars", 5)])
def test_xlsx_invalid_numeric_values_still_obey_text_limits(monkeypatch, attribute, limit):
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = 1
    data = _xlsx_with_numeric_values(workbook, {"A1": "NULL"})
    monkeypatch.setattr(settings, attribute, limit)

    with pytest.raises(XlsxLimitError, match="text"):
        extract_xlsx_segments(data)


def test_xlsx_keeps_content_after_a_long_empty_row_gap():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "header"
    workbook.active["A100"] = "late value"
    source = io.BytesIO()
    workbook.save(source)

    segment = extract_xlsx_segments(source.getvalue())[0]

    assert segment["rows"] == ["header", "late value"]
    assert segment["row_numbers"] == [1, 100]
    late_unit = segment["structured_units"][1]
    assert late_unit["value"] == "late value"
    assert late_unit["source"]["cell"] == {"start": "A100", "end": "A100"}


def test_xlsx_long_empty_row_gap_does_not_bypass_content_row_limit(monkeypatch):
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "header"
    workbook.active["A100"] = "late value"
    source = io.BytesIO()
    workbook.save(source)
    monkeypatch.setattr(settings, "xlsx_max_rows", 1)

    with pytest.raises(XlsxLimitError, match="row limit"):
        extract_xlsx_segments(source.getvalue())


def test_styled_phantom_columns_do_not_expand_every_empty_row():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "header"
    workbook.active["XFD2"].number_format = "@"
    workbook.active["B1000"] = "late value"
    source = io.BytesIO()
    workbook.save(source)
    data = source.getvalue()
    reader = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        rows = list(xlsx_extract._iter_xlsx_rows(reader.active))
        # Only two real values need slots; phantom grid width must not multiply
        # the work for the 998 empty rows. Physical row/column identity remains.
        assert sum(len(row) for row in rows) <= 10
        assert rows[0][0] == "header" and rows[999][1] == "late value"
    finally:
        reader.close()
    segment = extract_xlsx_segments(data)[0]
    assert segment["row_numbers"] == [1, 1000]
    assert segment["structured_units"][-1]["source"]["cell"] == {"start": "B1000", "end": "B1000"}


@pytest.mark.parametrize("raw_value", ["NaN", "Infinity", "-Infinity", "1e309"])
@pytest.mark.parametrize("date_style", [False, True])
def test_nonfinite_numeric_exports_preserve_raw_text_without_date_conversion(raw_value, date_style):
    workbook = openpyxl.Workbook()
    workbook.active["C3"] = 1
    if date_style:
        workbook.active["C3"].number_format = "yyyy-mm-dd"
    data = _xlsx_with_numeric_values(workbook, {"C3": raw_value})
    segment = extract_xlsx_segments(data)[0]
    assert segment["rows"] == [raw_value]
    assert segment["structured_units"][0]["value"] == raw_value
    assert segment["structured_units"][0]["source"]["cell"] == {"start": "C3", "end": "C3"}


def test_understated_sheet_dimensions_do_not_drop_real_late_rows():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "header"
    workbook.active["C100"] = "late value"
    source = io.BytesIO()
    workbook.save(source)
    output = io.BytesIO()
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as modified:
        for entry in original.infolist():
            content = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                root.find(f"{namespace}dimension").set("ref", "A1")
                content = ElementTree.tostring(root)
            modified.writestr(entry, content)
    segment = extract_xlsx_segments(output.getvalue())[0]
    assert segment["row_numbers"] == [1, 100]
    assert segment["structured_units"][-1]["value"] == "late value"
    assert segment["structured_units"][-1]["source"]["cell"] == {"start": "C100", "end": "C100"}


def test_actual_row_coordinate_obeys_bounds_even_when_dimension_is_small():
    source = _xlsx_bytes(["late value"])
    output = io.BytesIO()
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with zipfile.ZipFile(io.BytesIO(source)) as original, zipfile.ZipFile(output, "w") as modified:
        for entry in original.infolist():
            content = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                root.find(f"{namespace}sheetData/{namespace}row").set("r", "2000001")
                root.find(f"{namespace}sheetData/{namespace}row/{namespace}c").set("r", "A2000001")
                content = ElementTree.tostring(root)
            modified.writestr(entry, content)
    with pytest.raises(XlsxLimitError, match="physical row"):
        extract_xlsx_segments(output.getvalue())


@pytest.mark.parametrize("raw_value", ["100000000", "-100000000"])
def test_out_of_range_date_serial_retains_raw_text(raw_value):
    workbook = openpyxl.Workbook()
    workbook.active["C3"] = 1
    workbook.active["C3"].number_format = "yyyy-mm-dd"
    segment = extract_xlsx_segments(_xlsx_with_numeric_values(workbook, {"C3": raw_value}))[0]
    assert segment["rows"] == [raw_value]
    assert segment["structured_units"][0]["value"] == raw_value
