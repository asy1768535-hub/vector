from __future__ import annotations

import uuid
from types import SimpleNamespace

import pytest

from app.services import chat_evidence as E
from app.services.evidence_locator_projection import parse_locator
from app.services.splitter import _segment_flat_text
from app.services.xlsx_extract import _row_to_numbered_line
from tests.test_chat_evidence_pg import _locator


def table(rows):
    document_id, revision_id, parent_id = (uuid.uuid4() for _ in range(3))
    text = _segment_flat_text({"kind": "table", "heading": "Data", "caption": "Data",
                               "rows": [_row_to_numbered_line(tuple(row)) for row in rows]})
    parent = SimpleNamespace(id=parent_id, document_id=document_id, library_id=uuid.uuid4(),
        document_revision_id=revision_id, parent_block_id=None, text=text,
        content={"parser_unit": {"segment_heading": "Data", "segment_caption": "Data"}})
    raw = _locator(document_id, revision_id, parent_id, None, text, 1, len(rows), "table")
    locator = parse_locator(raw, document_id=document_id, document_revision_id=revision_id, revision_no=3,
                            unit_id=parent_id, parent_unit_id=None)
    blocks = []
    for row_number, row in enumerate(rows, 1):
        for column, value in enumerate(row, 1):
            if value is None:
                continue
            unit_id = uuid.uuid4()
            cell_text = str(value)
            raw_cell = _locator(document_id, revision_id, unit_id, parent_id, cell_text, row_number, row_number, "cell")
            raw_cell["source"]["column"] = {"start": column, "end": column}
            blocks.append(SimpleNamespace(id=unit_id, document_id=document_id, library_id=parent.library_id,
                document_revision_id=revision_id, parent_block_id=parent_id, block_kind="cell", text=cell_text,
                content={"evidence_locator_v1": raw_cell, "parser_unit": {"value": value}}))
    return parent, locator, blocks


def test_exact_decimal_sum_and_distinct_values_have_separate_meanings():
    parent, locator, blocks = table([["对象", "费用(元)"], ["重复对象", 0.1], ["重复对象", 0.2]])
    result, notices = E.calculate_table_snapshot(parent, locator, blocks)
    assert not notices
    assert result["row_count"] == 2
    assert result["columns"][0]["distinct_count"] == 1
    assert result["columns"][1]["sum"] == "0.3"
    assert result["columns"][1]["unit"] == "元"


def test_saved_total_row_is_not_added_a_second_time():
    parent, locator, blocks = table([["对象", "数量(件)"], ["甲", 2], ["乙", 3], ["合计", 5]])
    result, notices = E.calculate_table_snapshot(parent, locator, blocks)
    assert not notices
    assert result["row_count"] == 2
    assert result["columns"][1]["sum"] == "5"
    assert result["excluded_summary_rows"] == [4]


@pytest.mark.parametrize("change", ["duplicate_coordinate", "forged_identity", "bad_hash", "missing_snapshot_row", "header_duplicate"])
def test_untrusted_or_incomplete_snapshot_cannot_produce_a_full_calculation(change):
    rows = [["对象", "数量(件)"], ["甲", 2], ["乙", 3]]
    if change == "header_duplicate":
        rows[0] = ["数量", "数量"]
    parent, locator, blocks = table(rows)
    if change == "duplicate_coordinate":
        blocks.append(blocks[-1])
    elif change == "forged_identity":
        blocks[-1].document_revision_id = uuid.uuid4()
    elif change == "bad_hash":
        blocks[-1].text = "changed"
    elif change == "missing_snapshot_row":
        blocks = blocks[:-2]
    result, notices = E.calculate_table_snapshot(parent, locator, blocks)
    assert result is None and notices


@pytest.mark.parametrize("value", [True, "NULL", "=1+1", None, "3千件", "Infinity"])
def test_ambiguous_missing_or_non_numeric_values_limit_only_the_affected_field(value):
    parent, locator, blocks = table([["对象", "数量(件)", "费用(元)"], ["甲", 2, 0.1], ["乙", value, 0.2]])
    result, notices = E.calculate_table_snapshot(parent, locator, blocks)
    assert result["row_count"] == 2
    assert result["columns"][1]["sum"] is None
    assert result["columns"][2]["sum"] == "0.3"
    assert notices


def test_mixed_units_are_not_summed_or_silently_converted():
    parent, locator, blocks = table([["对象", "质量(kg)"], ["甲", "2 kg"], ["乙", "3 g"]])
    result, notices = E.calculate_table_snapshot(parent, locator, blocks)
    assert result["columns"][1]["sum"] is None and notices


def test_formula_without_saved_value_is_not_evaluated():
    parent, locator, blocks = table([["对象", "数量(件)"], ["甲", 2], ["乙", "=1+1"]])
    blocks[-1].content["parser_unit"] = {"formula": "=1+1"}
    result, notices = E.calculate_table_snapshot(parent, locator, blocks)
    assert result["columns"][1]["sum"] is None and notices


def test_cell_and_column_limits_are_frozen_before_calculation():
    parent, locator, blocks = table([["对象", "数量(件)"], ["甲", 2]])
    result, notices = E.calculate_table_snapshot(parent, locator, blocks * 513)
    assert result is None and notices


def test_calculation_text_is_reserved_after_required_sources_before_optional_text():
    from app.schemas.dify import DifyRecord
    from app.services.chat_answer import build_context
    source = DifyRecord(content="对象 | 费用(元)\n甲 | 0.1\n乙 | 0.2", title="Data.xlsx", score=1,
                        metadata={"chunk_id": "required"})
    optional = DifyRecord(content="optional" * 300, title="Other.pdf", score=1, metadata={"chunk_id": "optional"})
    calculation = {"title": "Data.xlsx", "location": "Data，第 1–3 行", "chunk_ids": ["required"],
                   "result": {"row_count": 2, "excluded_summary_rows": [], "columns": [
                       {"field": "费用(元)", "unit": "元", "sum": "0.3", "distinct_count": 2, "nonempty_count": 2}]}}
    records, notices, scope = E.pack_calculated_evidence([[source]], [optional], [calculation], 1000)
    assert not notices and records == [source]
    assert "合计=0.3" in scope["calculation_context"] and "[1]" in scope["calculation_context"]
    context, _used = build_context(records, scope["context_chars"])
    assert len(context) + len(scope["calculation_context"]) <= 1000


def test_calculation_budget_never_displaces_required_evidence():
    from app.schemas.dify import DifyRecord
    from app.services.chat_answer import build_context
    record = DifyRecord(content="the necessary full table", title="Data", score=1, metadata={"chunk_id": "required"})
    context, _ = build_context([record], 1000)
    calculation = {"title": "Data", "location": "", "chunk_ids": ["required"],
                   "result": {"row_count": 2, "excluded_summary_rows": [], "columns": []}}
    records, notices, scope = E.pack_calculated_evidence([[record]], [], [calculation], len(context))
    assert records == [record] and notices
    assert not scope["calculation_context"]


def test_calculation_is_removed_if_any_required_group_member_is_outside_context():
    from app.schemas.dify import DifyRecord
    records = [DifyRecord(content="x" * 100, title="Data", score=1, metadata={"chunk_id": str(index)}) for index in range(2)]
    calculation = {"title": "Data", "location": "", "chunk_ids": ["0", "1"],
                   "result": {"row_count": 2, "excluded_summary_rows": [], "columns": []}}
    packed, notices, scope = E.pack_calculated_evidence([records], [], [calculation], 150)
    assert not packed and notices and not scope["calculation_context"]


def test_cell_outside_parent_row_range_is_not_reclassified_as_covered_evidence():
    parent, locator, blocks = table([["对象", "数量(件)"], ["甲", 2], ["乙", 3]])
    for block in blocks[-2:]:
        block.content["evidence_locator_v1"]["source"]["row"] = {"start": 99, "end": 99}
    result, notices = E.calculate_table_snapshot(parent, locator, blocks)
    assert result is None and notices


def test_real_docx_row_and_cell_hierarchy_supports_the_same_field_calculation():
    from io import BytesIO
    import docx
    from app.services.docx_extract import extract_docx_segments
    from app.services.evidence_write_path import _structural_payload
    document = docx.Document()
    grid = document.add_table(rows=3, cols=2)
    for r, row in enumerate([["对象", "数量(件)"], ["甲", "2"], ["乙", "4"]]):
        for c, text in enumerate(row):
            grid.cell(r, c).text = text
    buffer = BytesIO()
    document.save(buffer)
    segment = next(item for item in extract_docx_segments(buffer.getvalue()) if item["kind"] == "table")
    parent, _previous, _blocks = table([["对象", "数量(件)"], ["甲", "2"], ["乙", "4"]])
    parent.text = _segment_flat_text(segment)
    parent.content["parser_unit"] = {f"segment_{key}": segment.get(key) for key in ("heading", "caption")}
    raw = _locator(parent.document_id, parent.document_revision_id, parent.id, None, parent.text, 1, 3, "table")
    raw["source"].update(kind="docx", table={"index": 0})
    raw["source"].pop("sheet")
    locator = parse_locator(raw, document_id=parent.document_id, document_revision_id=parent.document_revision_id,
                            revision_no=3, unit_id=parent.id, parent_unit_id=None)
    ids = {segment["unit_key"]: parent.id}
    blocks = []
    for unit in segment["structured_units"]:
        unit_id = uuid.uuid4()
        ids[unit["unit_key"]] = unit_id
        parent_id = ids[unit["parent_key"]]
        row = unit["source"]["row"]["start"]
        text = unit.get("text", str(unit.get("value", "")))
        raw = _locator(parent.document_id, parent.document_revision_id, unit_id, parent_id, text, row, row, unit["unit_kind"])
        raw["source"] = {"kind": "docx", "file_name": "Data.docx", **unit["source"]}
        blocks.append(SimpleNamespace(id=unit_id, document_id=parent.document_id, library_id=parent.library_id,
            document_revision_id=parent.document_revision_id, parent_block_id=parent_id, block_kind=unit["unit_kind"], text=text,
            content={"evidence_locator_v1": raw, "parser_unit": _structural_payload(unit, segment)}))
    result, notices = E.calculate_table_snapshot(parent, locator, blocks)
    assert not notices and result["columns"][1]["sum"] == "6"


def test_summary_word_in_a_regular_field_does_not_delete_a_data_row():
    parent, locator, blocks = table([["对象", "备注", "数量(件)"], ["甲", "合计", 2], ["乙", "正常", 3]])
    result, notices = E.calculate_table_snapshot(parent, locator, blocks)
    assert not notices
    assert result["row_count"] == 2 and result["excluded_summary_rows"] == []
    assert result["columns"][2]["sum"] == "5"
