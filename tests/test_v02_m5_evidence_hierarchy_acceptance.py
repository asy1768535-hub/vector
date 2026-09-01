from __future__ import annotations

import asyncio
import hashlib
import json
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import docx
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.auth.backend import current_active_user
from app.api.health import _evidence_locator_payload
from app.config import Settings, settings
from app.db import get_db
from app.main import app
from app.models.chunk import Chunk
from app.models.document import Document
from app.models.document_revision import DocumentRevision
from app.models.evidence_unit import EvidenceUnit
from app.models.library import Library
from app.models.user import User
from app.schemas.dify import DifyRetrievalRequest, RetrievalSetting
from app.schemas.evidence_locator import EvidenceLocatorV1, sha256_text
from app.services import retrieval, xlsx_extract
from app.services.federated_retrieval import _source_projection
from app.services.import_parsing import parse_import_file
from app.services.parser_units import validate_parser_unit_hierarchy
from app.workers.embedder import _build_payload


API_LIBRARY_ID = uuid.UUID("00000000-0000-0000-0000-000000000101")
API_DOCUMENT_ID = uuid.UUID("00000000-0000-0000-0000-000000000102")
API_REVISION_ID = uuid.UUID("00000000-0000-0000-0000-000000000103")
API_OLD_REVISION_ID = uuid.UUID("00000000-0000-0000-0000-000000000104")
API_BLOCK_ID = uuid.UUID("00000000-0000-0000-0000-000000000105")
API_EVIDENCE_ID = uuid.UUID("00000000-0000-0000-0000-000000000106")
API_OLD_EVIDENCE_ID = uuid.UUID("00000000-0000-0000-0000-000000000107")
API_LEGACY_EVIDENCE_ID = uuid.UUID("00000000-0000-0000-0000-000000000108")
API_CHUNK_ID = uuid.UUID("00000000-0000-0000-0000-000000000109")
API_FILE_ID = uuid.UUID("00000000-0000-0000-0000-00000000010a")
API_USER_ID = uuid.UUID("00000000-0000-0000-0000-00000000010b")


def _library() -> SimpleNamespace:
    return SimpleNamespace(
        chunk_size=160,
        chunk_overlap=0,
        ocr_enabled=False,
        docx_table_aware=True,
    )


def _flatten(segments: list[dict]) -> list[dict]:
    units: list[dict] = []
    for segment in segments:
        units.append(segment["parser_unit"])
        units.extend(segment.get("structured_units") or [])
    return units


def _assert_segment_ranges(result, *, exact_segment_text: bool = True) -> None:
    assert result.segments
    assert result.chunks
    validate_parser_unit_hierarchy(_flatten(result.segments))
    assert hashlib.sha256(result.normalized_text.encode("utf-8")).hexdigest()
    for segment in result.segments:
        unit = segment["parser_unit"]
        assert len(unit["parser"]["config_hash"]) == 64
        source_text = unit.get("source", {}).get("text")
        if source_text is not None:
            reconstructed = result.normalized_text[source_text["start"] : source_text["end"]]
            if exact_segment_text and segment.get("text"):
                assert reconstructed == segment["text"]
            elif segment.get("text"):
                assert segment["text"] in reconstructed
            for item in source_text.get("ranges", []):
                value = result.normalized_text[item["start"] : item["end"]]
                assert item["sha256"] == hashlib.sha256(value.encode("utf-8")).hexdigest()


def _write_docx(path: Path) -> None:
    document = docx.Document()
    document.add_heading("Acceptance heading", level=1)
    document.add_paragraph("A fresh DOCX paragraph with provenance.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "1"
    document.save(path)


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["Name", "Formula"])
    first.append(["A", "=1+1"])
    second = workbook.create_sheet("Second")
    second.append(["Name", "Value"])
    second.append(["B", "2"])
    workbook.save(path)


def _minimal_text_pdf(text: str) -> bytes:
    """Build a tiny PDF using only the PDF syntax consumed by the existing pypdf parser."""
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 20 50 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 100] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_start = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF\n".encode(
            "ascii"
        )
    )
    return bytes(output)


def test_m5_fresh_isolated_format_matrix_round_trips_provenance(tmp_path, monkeypatch):
    library = _library()

    txt = tmp_path / "fixture.txt"
    txt.write_text("plain text fixture", encoding="utf-8")
    markdown = tmp_path / "fixture.md"
    markdown.write_text("# Markdown\n\nbody", encoding="utf-8")
    for path in (txt, markdown):
        result = parse_import_file(path, library)
        _assert_segment_ranges(result)
        assert result.segments[0]["source_kind"] == "text"

    pdf = tmp_path / "fixture.pdf"
    pdf.write_bytes(_minimal_text_pdf("PDF acceptance fixture"))
    pdf_result = parse_import_file(pdf, library)
    _assert_segment_ranges(pdf_result)
    assert pdf_result.segments[0]["source_kind"] == "pdf"
    assert pdf_result.segments[0]["parser_unit"]["source"]["page"] == {"start": 1, "end": 1}
    assert "PDF acceptance fixture" in pdf_result.normalized_text

    docx_path = tmp_path / "fixture.docx"
    _write_docx(docx_path)
    docx_result = parse_import_file(docx_path, library)
    _assert_segment_ranges(docx_result, exact_segment_text=False)
    assert any(
        unit.get("source", {}).get("table") == {"index": 0}
        for unit in _flatten(docx_result.segments)
    )
    assert any(unit["unit_kind"] == "cell" for unit in _flatten(docx_result.segments))

    xlsx_path = tmp_path / "fixture.xlsx"
    _write_xlsx(xlsx_path)
    xlsx_result = parse_import_file(xlsx_path, library)
    _assert_segment_ranges(xlsx_result)
    sheets = [segment["parser_unit"]["source"]["sheet"]["name"] for segment in xlsx_result.segments]
    assert sheets == ["First", "Second"]
    assert any(unit.get("formula") == "=1+1" for unit in _flatten(xlsx_result.segments))

    csv_path = tmp_path / "fixture.csv"
    csv_path.write_text('name,value\n\n"A\n1",two\n', encoding="utf-8")
    csv_result = parse_import_file(csv_path, library)
    _assert_segment_ranges(csv_result)
    assert csv_result.segments[1]["parser_unit"]["source"]["row"] == {"start": 3, "end": 3}
    assert csv_result.segments[1]["structured_units"][0]["source"]["cell"] == {
        "start": "A3",
        "end": "A3",
    }

    json_path = tmp_path / "fixture.json"
    json_path.write_text(json.dumps({"records": [{"name": "A", "value": 1}]}), encoding="utf-8")
    json_result = parse_import_file(json_path, library)
    _assert_segment_ranges(json_result)
    pointers = {
        unit["source"]["json_pointer"]
        for unit in _flatten(json_result.segments)
        if "json_pointer" in unit.get("source", {})
    }
    assert {"/records/0/name", "/records/0/value"} <= pointers


def test_m5_xls_optional_dependency_is_an_explicit_skip_or_real_parser_contract(monkeypatch):
    pytest.importorskip("xlrd", reason="xlrd is optional for .xls parser support")
    parser = {"name": "fixture-xls", "version": "v1", "config_hash": "a" * 64}
    segment = {
        "kind": "table",
        "heading": "Legacy",
        "caption": "Legacy",
        "header": "Name",
        "rows": ["Name", "A"],
        "row_numbers": [1, 2],
        "source_kind": "xls",
        "ordinal": 0,
        "unit_key": "xls:sheet:0",
        "parser": parser,
        "location": {"type": "sheet", "sheet": "Legacy"},
        "structured_units": [],
    }
    monkeypatch.setattr(xlsx_extract, "_read_xls", lambda _data: [segment])
    parsed = xlsx_extract.extract_xlsx_segments(b"fixture", is_xls=True)
    assert parsed[0]["source_kind"] == "xls"
    assert parsed[0]["unit_key"] == "xls:sheet:0"


def test_m5_locator_read_rollout_off_uses_legacy_federated_projection(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_locator_read", False)
    metadata = {
        "document_id": "legacy-doc",
        "document_revision_id": "legacy-rev",
        "document_revision": 4,
        "chunk_id": "legacy-chunk",
        "page": 7,
        "title_path": ["Legacy"],
        "evidence_locator_v1_projection": {"document_id": "forged"},
    }
    source = _source_projection(metadata)
    assert source.document_id == "legacy-doc"
    assert source.document_revision_id == "legacy-rev"
    assert source.document_revision == 4
    assert source.chunk_id == "legacy-chunk"
    assert source.page == 7
    assert source.title_path == ("Legacy",)


def test_m5_retrieval_rollout_off_preserves_legacy_payload_shape(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_locator_read", False)
    hits = [
        {
            "id": "legacy-chunk",
            "score": 0.9,
            "payload": {
                "document_id": "legacy-doc",
                "chunk_id": "legacy-chunk",
                "text": "legacy text",
                "title": "Legacy",
                "page": 7,
                "evidence_locator_v1_projection": {"document_id": "forged"},
            },
        }
    ]

    async def search(*_args, **_kwargs):
        return hits

    async def embed_one(*_args, **_kwargs):
        return [0.1, 0.2, 0.3]

    with patch.object(retrieval.qdrant, "search", new=search), patch.object(
        retrieval.embedding, "embed_one", new=embed_one
    ):
        response = asyncio.run(
            retrieval.run_retrieval(
                collection="m5",
                embedding_model="fixture",
                embedding_base_url=None,
                request=DifyRetrievalRequest(
                    knowledge_id="m5",
                    query="legacy",
                    retrieval_setting=RetrievalSetting(top_k=1, score_threshold=0),
                ),
            )
        )
    assert len(response.records) == 1
    assert response.records[0].metadata["document_id"] == "legacy-doc"
    assert "evidence_locator_v1_projection" not in response.records[0].metadata


def test_m5_projection_rollout_off_preserves_existing_embedding_payload(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_locator_projection", False)
    library = SimpleNamespace(id="library", qdrant_collection="m5")
    document = SimpleNamespace(
        id="document",
        title="Fixture",
        external_id="fixture",
        current_revision=1,
        doc_metadata={"legacy": "value"},
        visibility_scope=None,
        security_level=None,
    )
    chunk = SimpleNamespace(
        id="chunk",
        seq=0,
        text="body",
        document_revision_id=None,
        evidence_id=None,
        block_id=None,
        chunk_kind="text",
        page_start=None,
        page_end=None,
        title_path=None,
        source_start=0,
        source_end=4,
        position=None,
    )
    payload = _build_payload(library, document, chunk)
    assert payload["legacy"] == "value"
    assert payload["text"] == "body"
    assert "evidence_locator_v1_projection" not in payload


def test_m5_health_projection_is_bounded_and_contains_no_source_data(monkeypatch):
    monkeypatch.setattr(settings, "enable_evidence_locator_read", True)
    monkeypatch.setattr(settings, "enable_evidence_locator_projection", True)
    payload = _evidence_locator_payload()
    assert set(payload) == {
        "locator_contract_version",
        "parser_unit_contract_version",
        "projection_version",
        "read_enabled",
        "projection_enabled",
        "write_path_enabled",
    }
    serialized = json.dumps(payload, ensure_ascii=False)
    assert "raw_file_sha256" not in serialized
    assert "quote" not in serialized
    assert "storage" not in serialized
    assert all(isinstance(value, (str, bool)) for value in payload.values())


def test_m5_env_example_matches_locator_settings_defaults():
    defaults = Settings(_env_file=None)
    env_example = Path(".env.example").read_text(encoding="utf-8")
    assert defaults.enable_evidence_write_path is False
    assert defaults.enable_evidence_locator_read is True
    assert defaults.enable_evidence_locator_projection is True
    assert "ENABLE_EVIDENCE_WRITE_PATH=false" in env_example
    assert "ENABLE_EVIDENCE_LOCATOR_READ=true" in env_example
    assert "ENABLE_EVIDENCE_LOCATOR_PROJECTION=true" in env_example


def _api_library() -> Library:
    return Library(
        id=API_LIBRARY_ID,
        slug="m5-api",
        name="M5 API",
        qdrant_collection="m5-api",
        embedding_model="fixture",
        embedding_dim=3,
        chunk_size=160,
        chunk_overlap=0,
        lifecycle_mode="managed",
        index_state="ready",
    )


def _api_document() -> Document:
    return Document(
        id=API_DOCUMENT_ID,
        library_id=API_LIBRARY_ID,
        title="M5 API fixture",
        external_id="m5-api",
        content_hash=sha256_text("before quote after"),
        current_revision=2,
        current_revision_id=API_REVISION_ID,
        latest_revision_id=API_REVISION_ID,
        status="ready",
    )


def _api_revision(revision_id, revision_no: int, text: str, status: str) -> DocumentRevision:
    return DocumentRevision(
        id=revision_id,
        document_id=API_DOCUMENT_ID,
        library_id=API_LIBRARY_ID,
        revision_no=revision_no,
        title="M5 API fixture",
        content_hash=sha256_text(text),
        normalized_text=text,
        parser_name="fixture",
        parser_version="v1",
        chunking_strategy="text",
        chunking_strategy_version="v1",
        status=status,
    )


def _api_locator(*, evidence_id, revision_id, revision_no: int, text: str, page: int) -> dict:
    quote = text
    start = 0 if text.startswith("old") else 7
    return EvidenceLocatorV1(
        document_id=API_DOCUMENT_ID,
        document_revision_id=revision_id,
        revision_no=revision_no,
        document_revision_file_id=API_FILE_ID,
        raw_file_sha256="a" * 64,
        normalized_content_hash=sha256_text("old quote" if revision_no == 1 else "before quote after"),
        unit_id=evidence_id,
        parent_unit_id=API_BLOCK_ID,
        unit_kind="chunk",
        ordinal=0,
        parser={"name": "fixture", "version": "v1", "config_hash": "b" * 64},
        source={
            "kind": "pdf",
            "file_name": "fixture.pdf",
            "page": {"start": page, "end": page},
            "text": {
                "start": start,
                "end": start + len(quote),
                "ranges": [{
                    "start": start,
                    "end": start + len(quote),
                    "sha256": sha256_text(quote),
                }],
            },
        },
        unit_text_sha256=sha256_text(quote),
        quote_sha256=sha256_text(quote),
    ).model_dump(mode="json")


def _api_rows():
    current_text = "before quote after"
    old_text = "old quote"
    current_revision = _api_revision(API_REVISION_ID, 2, current_text, "ready")
    old_revision = _api_revision(API_OLD_REVISION_ID, 1, old_text, "superseded")
    current = EvidenceUnit(
        id=API_EVIDENCE_ID,
        library_id=API_LIBRARY_ID,
        document_id=API_DOCUMENT_ID,
        document_revision_id=API_REVISION_ID,
        document_block_id=API_BLOCK_ID,
        evidence_kind="chunk",
        source_start=1,
        source_end=5,
        page_start=8,
        text_quote="quote",
        status="active",
        evidence_metadata={
            "evidence_locator_v1": _api_locator(
                evidence_id=API_EVIDENCE_ID,
                revision_id=API_REVISION_ID,
                revision_no=2,
                text="quote",
                page=2,
            )
        },
    )
    old_locator = _api_locator(
        evidence_id=API_OLD_EVIDENCE_ID,
        revision_id=API_OLD_REVISION_ID,
        revision_no=1,
        text=old_text,
        page=1,
    )
    old_locator.pop("document_revision_file_id")
    old_locator.pop("raw_file_sha256")
    old = EvidenceUnit(
        id=API_OLD_EVIDENCE_ID,
        library_id=API_LIBRARY_ID,
        document_id=API_DOCUMENT_ID,
        document_revision_id=API_OLD_REVISION_ID,
        document_block_id=API_BLOCK_ID,
        evidence_kind="chunk",
        source_start=0,
        source_end=len(old_text),
        text_quote=old_text,
        status="active",
        evidence_metadata={"evidence_locator_v1": old_locator},
    )
    legacy = EvidenceUnit(
        id=API_LEGACY_EVIDENCE_ID,
        library_id=API_LIBRARY_ID,
        document_id=API_DOCUMENT_ID,
        document_revision_id=API_REVISION_ID,
        document_block_id=API_BLOCK_ID,
        evidence_kind="chunk",
        source_start=1,
        source_end=5,
        page_start=8,
        text_quote="quote",
        status="active",
        evidence_metadata={"legacy": True},
    )
    chunk = Chunk(
        id=API_CHUNK_ID,
        document_id=API_DOCUMENT_ID,
        library_id=API_LIBRARY_ID,
        document_revision_id=API_REVISION_ID,
        block_id=API_BLOCK_ID,
        evidence_id=API_EVIDENCE_ID,
        seq=0,
        chunk_kind="text",
        text="quote",
        token_count=1,
        source_start=1,
        source_end=5,
        chunk_metadata={"evidence_locator_v1": _api_locator(
            evidence_id=API_CHUNK_ID,
            revision_id=API_REVISION_ID,
            revision_no=2,
            text="quote",
            page=2,
        )},
    )
    revision_file = SimpleNamespace(
        id=API_FILE_ID,
        document_id=API_DOCUMENT_ID,
        document_revision_id=API_REVISION_ID,
        library_id=API_LIBRARY_ID,
        sha256="a" * 64,
    )
    return current_revision, old_revision, current, old, legacy, chunk, revision_file


def test_m5_http_evidence_source_window_current_history_and_legacy_round_trip():
    current_revision, old_revision, current, old, legacy, chunk, revision_file = _api_rows()
    document = _api_document()
    library = _api_library()
    rows = {
        (EvidenceUnit, current.id): current,
        (EvidenceUnit, old.id): old,
        (EvidenceUnit, legacy.id): legacy,
        (Chunk, chunk.id): chunk,
        (Document, document.id): document,
        (DocumentRevision, current_revision.id): current_revision,
        (DocumentRevision, old_revision.id): old_revision,
    }

    async def get_model(model, ident):
        if model.__name__ == "DocumentRevisionFile":
            return revision_file if ident == API_FILE_ID else None
        return rows.get((model, ident))

    db = AsyncMock()
    db.get = AsyncMock(side_effect=get_model)
    user = User(id=API_USER_ID, email="m5-api@example.com", is_superuser=True, is_active=True)

    async def override_db():
        return db

    async def override_user():
        return user

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_active_user] = override_user
    try:
        with patch("app.deps.load_active_library", new=AsyncMock(return_value=library)):
            client = TestClient(app)
            current_response = client.get(f"/libraries/{library.slug}/evidence/{current.id}")
            source_response = client.get(f"/libraries/{library.slug}/chunks/{chunk.id}/source")
            history_response = client.get(
                f"/libraries/{library.slug}/evidence/{old.id}",
                params={"revision_id": str(old_revision.id)},
            )
            legacy_response = client.get(f"/libraries/{library.slug}/evidence/{legacy.id}")
    finally:
        app.dependency_overrides.clear()

    assert current_response.status_code == 200
    current_payload = current_response.json()
    assert current_payload["document_revision_id"] == str(API_REVISION_ID)
    assert current_payload["source_start"] == 7
    assert current_payload["page_start"] == 2
    assert current_payload["text_quote"] == "quote"
    assert "evidence_locator_v1" not in current_payload["evidence_metadata"]
    assert "raw_file_sha256" not in json.dumps(current_payload)
    assert "fixture.pdf" in json.dumps(current_payload)

    assert source_response.status_code == 200
    source_payload = source_response.json()
    assert source_payload["chunk_id"] == str(API_CHUNK_ID)
    assert source_payload["source_start"] == 7
    assert source_payload["page_start"] == 2
    assert "raw_file_sha256" not in json.dumps(source_payload)

    assert history_response.status_code == 200
    history_payload = history_response.json()
    assert history_payload["document_revision_id"] == str(API_OLD_REVISION_ID)
    assert history_payload["source_start"] == 0
    assert history_payload["page_start"] == 1

    assert legacy_response.status_code == 200
    legacy_payload = legacy_response.json()
    assert legacy_payload["source_start"] == 1
    assert legacy_payload["source_end"] == 5
    assert legacy_payload["page_start"] == 8
    assert legacy_payload["evidence_metadata"] == {"legacy": True}


def test_m5_http_invalid_locator_hash_falls_back_to_scalar_without_leaking_envelope():
    current_revision, _old_revision, current, _old, _legacy, _chunk, revision_file = _api_rows()
    document = _api_document()
    library = _api_library()
    invalid_locator = dict(current.evidence_metadata["evidence_locator_v1"])
    invalid_locator["quote_sha256"] = sha256_text("wrong")
    current.evidence_metadata = {"evidence_locator_v1": invalid_locator}
    rows = {
        (EvidenceUnit, current.id): current,
        (Document, document.id): document,
        (DocumentRevision, current_revision.id): current_revision,
    }

    async def get_model(model, ident):
        if model.__name__ == "DocumentRevisionFile":
            return revision_file if ident == API_FILE_ID else None
        return rows.get((model, ident))

    db = AsyncMock()
    db.get = AsyncMock(side_effect=get_model)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[current_active_user] = lambda: User(
        id=API_USER_ID, email="m5-api@example.com", is_superuser=True, is_active=True
    )
    try:
        with patch("app.deps.load_active_library", new=AsyncMock(return_value=library)):
            response = TestClient(app).get(f"/libraries/{library.slug}/evidence/{current.id}")
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    payload = response.json()
    assert payload["source_start"] == 1
    assert payload["source_end"] == 5
    assert payload["page_start"] == 8
    assert "raw_file_sha256" not in response.text
    assert "evidence_locator_v1" not in response.text
