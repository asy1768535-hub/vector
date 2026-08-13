from __future__ import annotations

import asyncio
import json
import hashlib
import uuid
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch
import pytest
from fastapi.testclient import TestClient
from fastapi import status
from sqlalchemy.exc import IntegrityError

from app.main import app
from app.auth.backend import current_active_user
from app.config import settings
from app.db import get_db
from app.models.user import User
from app.models.library import Library
from app.models.document import Document
from app.models.chunk import Chunk
from app.models.document_file import DocumentFile
from app.models.document_source import DocumentSource
from app.models.document_revision import DocumentRevision
from app.models.embedding_job import EmbeddingJob
from app.api.documents import _persist_graph_extraction_request
from app.schemas.documents import QueryRequest
from app.services import ingest as ingest_service


# Mock user and library
mock_user = User(
    id="00000000-0000-0000-0000-000000000001",
    email="test@example.com",
    is_superuser=True,
    is_active=True,
)
mock_library = Library(
    id="00000000-0000-0000-0000-000000000002",
    slug="testlib",
    name="Test Library",
    qdrant_collection="testlib_col",
    embedding_model="bge-m3",
    embedding_base_url="http://mock-embeddings",
    embedding_dim=1024,
    chunk_size=1000,
    chunk_overlap=120,
    lifecycle_mode="managed",
    index_state="ready",
)


def test_structured_chunk_metadata_includes_source_offsets():
    text, metadata = ingest_service._chunk_text_and_metadata(
        {"text": "第二行", "source_start": 4, "source_end": 7, "location": {"type": "line", "start_line": 2, "end_line": 2}},
        title="demo.txt",
        external_id="ext-1",
        revision=3,
    )

    assert text == "第二行"
    assert metadata == {
        "title": "demo.txt",
        "external_id": "ext-1",
        "source_start": 4,
        "source_end": 7,
        "location": {"type": "line", "start_line": 2, "end_line": 2},
        "source_revision": 3,
    }


def test_string_chunk_metadata_does_not_fabricate_offsets():
    text, metadata = ingest_service._chunk_text_and_metadata(
        "legacy chunk",
        title=None,
        external_id=None,
        revision=1,
    )

    assert text == "legacy chunk"
    assert metadata is None


def test_upload_graph_request_is_persisted_on_the_target_revision():
    job_id = uuid.uuid4()
    revision_id = uuid.uuid4()
    job = MagicMock(document_revision_id=revision_id)
    revision = MagicMock(parser_config={"parser": "plain"})
    db = AsyncMock()

    async def get_row(model, object_id):
        return {
            (EmbeddingJob, job_id): job,
            (DocumentRevision, revision_id): revision,
        }.get((model, object_id))

    db.get = AsyncMock(side_effect=get_row)
    persisted = asyncio.run(_persist_graph_extraction_request(
        db,
        {"operation": "created", "job_id": str(job_id)},
    ))

    assert persisted is True
    assert revision.parser_config == {
        "parser": "plain",
        "graph_extraction_requested": True,
    }


def test_unchanged_upload_does_not_request_graph_extraction():
    db = AsyncMock()

    persisted = asyncio.run(_persist_graph_extraction_request(
        db,
        {"operation": "unchanged", "job_id": str(uuid.uuid4())},
    ))

    assert persisted is False
    db.get.assert_not_awaited()

async def override_user():
    return mock_user


def make_db_mock(*, existing=None, scalar_list=None, get_return=None, library=None):
    """构造贴近真实 AsyncSession 的 mock。

    关键点：`await db.execute(...)` 返回的是同步 Result，其 .scalars()/.first()/.all()
    都是同步方法——裸 AsyncMock 会把它们变成协程导致链式调用炸。这里显式建一个同步
    Result，让 external_id 查重等 `(await db.execute(...)).scalars().first()` 正常工作。

    针对 #6 写守卫：_lock_writable 会 `SELECT sys_libraries ... FOR KEY SHARE`，期望拿到
    Library 行（拿不到则 404）。这里按语句是否命中 sys_libraries 派发——库查询返回 library
    （默认 mock_library），其余查询走 existing 语义，避免库锁查到 Document 或误 404。
    """
    db = AsyncMock()  # commit/flush/get/rollback 是 async → AsyncMock 合适
    db.add = MagicMock()
    db.add_all = MagicMock()
    lib_obj = library if library is not None else mock_library

    result = MagicMock()
    result.scalars.return_value.first.return_value = existing
    result.scalars.return_value.all.return_value = scalar_list or []
    result.scalar_one.return_value = 0
    result.scalar_one_or_none.return_value = existing

    lib_result = MagicMock()
    lib_result.scalars.return_value.first.return_value = lib_obj
    lib_result.scalars.return_value.all.return_value = [lib_obj]
    lib_result.scalar_one.return_value = lib_obj
    lib_result.scalar_one_or_none.return_value = lib_obj

    async def _execute(stmt, *a, **k):
        return lib_result if "sys_libraries" in str(stmt).lower() else result

    db.execute = _execute
    if get_return is not None:
        db.get = AsyncMock(return_value=get_return)
    else:
        db.get = AsyncMock(return_value=None)
    return db


async def override_db():
    return make_db_mock()

@pytest.fixture
def client():
    # Default mock user to superuser
    mock_user.is_superuser = True
    app.dependency_overrides[current_active_user] = override_user
    app.dependency_overrides[get_db] = override_db
    
    with patch("app.deps.load_active_library", new_callable=AsyncMock) as mock_load:
        mock_load.return_value = mock_library
        yield TestClient(app)
        
    app.dependency_overrides.clear()

def test_query_request_validation():
    # Test valid limit
    req = QueryRequest(query="test query", limit=10)
    assert req.query == "test query"
    assert req.limit == 10

    # Test limit validation constraints (min 1, max 20)
    with pytest.raises(Exception):
        QueryRequest(query="test", limit=0)
    with pytest.raises(Exception):
        QueryRequest(query="test", limit=21)
    with pytest.raises(Exception):
        QueryRequest(query="", limit=5)

@patch("app.services.embedding.embed_one", new_callable=AsyncMock)
@patch("app.services.qdrant.search", new_callable=AsyncMock)
def test_query_library_endpoint(mock_search, mock_embed, client, monkeypatch):
    # 显式关掉 rerank，使该用例不依赖 .env 的全局开关、也不联网（只验 dense 路径）
    monkeypatch.setattr(settings, "rerank_enabled", False)
    monkeypatch.setattr(settings, "rerank_base_url", "")
    # 关掉可见性回查（本用例验 dense/enrich 路径，不验 #6 过滤；过滤有独立单测+E2E）
    monkeypatch.setattr(settings, "retrieval_consistency_filter", False)
    mock_embed.return_value = [0.1] * 1024
    mock_search.return_value = [
        {
            "id": "123",
            "score": 0.85,
            "payload": {
                "text": "matched chunk text",
                "document_id": "doc123",
                "chunk_id": "chunk123",
                "title": "doc title",
                "extra": "value"
            }
        }
    ]

    response = client.post(
        "/libraries/testlib/query",
        json={"query": "test query", "limit": 5}
    )
    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert "results" in data
    assert len(data["results"]) == 1
    assert data["results"][0]["text"] == "matched chunk text"
    assert data["results"][0]["similarity"] == 0.85
    assert data["results"][0]["document_id"] == "doc123"
    # metadata 保留业务字段 + 排查用的 vector_score（无 rerank 时不含 rerank_score）
    assert data["results"][0]["metadata"] == {"extra": "value", "vector_score": 0.85}


@patch("app.services.rerank.rerank", new_callable=AsyncMock)
@patch("app.services.qdrant.search", new_callable=AsyncMock)
@patch("app.services.embedding.embed_one", new_callable=AsyncMock)
def test_query_rerank_backfills_and_keeps_both_scores(mock_embed, mock_search, mock_rerank, client, monkeypatch):
    """rerank 生效路径端到端：召回 3 条、reranker 只返回 1 条 →
    断言补满到 3 条、命中项带 rerank_score、补满项只带 vector_score。"""
    # 全局开启 rerank 并配好地址（lib.rerank_enabled=None → 回退全局）
    monkeypatch.setattr(settings, "rerank_enabled", True)
    monkeypatch.setattr(settings, "rerank_base_url", "http://mock-rerank/rerank")
    monkeypatch.setattr(settings, "rerank_model", "bge-reranker-v2-m3")
    monkeypatch.setattr(settings, "retrieval_consistency_filter", False)  # 本用例验 rerank，不验 #6 过滤

    mock_embed.return_value = [0.1] * 1024
    mock_search.return_value = [
        {"id": "1", "score": 0.90, "payload": {"text": "A", "document_id": "d1", "chunk_id": "c1", "title": "tA"}},
        {"id": "2", "score": 0.80, "payload": {"text": "B", "document_id": "d2", "chunk_id": "c2", "title": "tB"}},
        {"id": "3", "score": 0.70, "payload": {"text": "C", "document_id": "d3", "chunk_id": "c3", "title": "tC"}},
    ]
    # reranker 只返回向量序最差的第 3 条（idx=2），高分上浮到首位
    mock_rerank.return_value = [(2, 0.99)]

    response = client.post("/libraries/testlib/query", json={"query": "q", "limit": 5})
    assert response.status_code == status.HTTP_200_OK
    results = response.json()["results"]

    # 补满：3 条召回全部返回，顺序 = 重排命中 [2] + 原向量序补满 [0,1]
    assert [r["text"] for r in results] == ["C", "A", "B"]

    # 命中项：similarity=rerank 分，metadata 同时含 vector_score 与 rerank_score
    assert results[0]["similarity"] == 0.99
    assert results[0]["metadata"]["vector_score"] == 0.70
    assert results[0]["metadata"]["rerank_score"] == 0.99

    # 补满项：similarity 回退向量分，metadata 只有 vector_score、无 rerank_score
    assert results[1]["similarity"] == 0.90
    assert results[1]["metadata"]["vector_score"] == 0.90
    assert "rerank_score" not in results[1]["metadata"]
    assert results[2]["metadata"]["vector_score"] == 0.80
    assert "rerank_score" not in results[2]["metadata"]

@patch("app.deps.has_permission")
def test_query_library_unauthorized(mock_has_perm, client):
    # Set user to non-superuser to trigger permission checks
    mock_user.is_superuser = False
    mock_has_perm.return_value = False

    response = client.post(
        "/libraries/testlib/query",
        json={"query": "test query", "limit": 5}
    )
    assert response.status_code == status.HTTP_403_FORBIDDEN

def _safe_tmp_dir(name: str) -> Path:
    root = Path.cwd() / "pytest_tmp_files"
    root.mkdir(exist_ok=True)
    path = root / f"vector-db-{name}-{uuid.uuid4().hex}"
    path.mkdir(parents=True, exist_ok=False)
    return path.resolve()


@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_import_txt_file_persists_original_upload(mock_ingest, monkeypatch):
    mock_user.is_superuser = True
    tmp_path = _safe_tmp_dir("document-file-import")
    monkeypatch.setattr(settings, "document_files_dir", str(tmp_path))
    doc = MagicMock()
    doc.id = uuid.uuid4()
    doc.status = "pending"
    doc.title = "keep.txt"
    doc.external_id = None
    doc.current_revision = 1
    job = MagicMock()
    job.id = uuid.uuid4()
    mock_ingest.return_value = (doc, job, 1, False)
    db = make_db_mock()
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).post(
                "/libraries/testlib/import-file",
                files={"file": ("keep.txt", b"original bytes", "text/plain")},
            )
            assert resp.status_code == status.HTTP_201_CREATED
            added_files = [
                call.args[0] for call in db.add.call_args_list
                if call.args and isinstance(call.args[0], DocumentFile)
            ]
            assert len(added_files) == 1
            row = added_files[0]
            assert row.document_id == doc.id
            assert row.revision == 1
            assert row.file_name == "keep.txt"
            assert row.content_type == "text/plain"
            assert row.size_bytes == len(b"original bytes")
            assert (tmp_path / row.storage_path).read_bytes() == b"original bytes"
    finally:
        app.dependency_overrides.clear()


@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_import_txt_file(mock_ingest, client):
    mock_doc = AsyncMock()
    mock_doc.id = "00000000-0000-0000-0000-000000000003"
    mock_doc.status = "pending"
    mock_job = AsyncMock()
    mock_ingest.return_value = (mock_doc, mock_job, 1, False)

    file_content = b"This is plain text content."
    files = {"file": ("test.txt", file_content, "text/plain")}
    response = client.post("/libraries/testlib/import-file", files=files)
    
    assert response.status_code == status.HTTP_201_CREATED
    data = response.json()
    assert data["status"] == "success"
    assert data["imported_count"] == 1
    assert data["documents"][0]["title"] == "test.txt"
    assert data["documents"][0]["chunk_count"] == 1
    
    mock_ingest.assert_called_once()
    kwargs = mock_ingest.call_args[1]
    assert kwargs["text"] == "This is plain text content."
    assert kwargs["title"] == "test.txt"
    chunks = kwargs.get("chunks")
    assert chunks and all(isinstance(chunk, dict) for chunk in chunks)
    assert chunks[0]["text"] == "This is plain text content."
    assert kwargs["text"][chunks[0]["source_start"]:chunks[0]["source_end"]] == chunks[0]["text"]
    assert chunks[0]["location"]["type"] == "line"

@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_import_xlsx_passes_table_aware_chunks(mock_ingest, client):
    """上传 .xlsx → 表格感知切分 → chunks= 传进 ingest_text（防上传分支被改断）。"""
    import io as _io

    import openpyxl

    mock_doc = AsyncMock()
    mock_doc.id = "00000000-0000-0000-0000-000000000005"
    mock_doc.status = "pending"
    mock_ingest.return_value = (mock_doc, AsyncMock(), 1, False)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "技术选型"
    ws.append(["类别", "方案"])
    ws.append(["数据库", "MySQL"])
    ws.append(["缓存", "Redis"])
    buf = _io.BytesIO()
    wb.save(buf)

    files = {"file": ("台账.xlsx", buf.getvalue(),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    from app.services import xlsx_extract

    with patch("app.api.documents.asyncio.to_thread", new_callable=AsyncMock) as offload:
        offload.side_effect = lambda func, *args, **kwargs: func(*args, **kwargs)
        response = client.post("/libraries/testlib/import-file", files=files)

    assert response.status_code == status.HTTP_201_CREATED
    offload.assert_awaited_once()
    assert offload.await_args.args[0] is xlsx_extract.extract_xlsx_segments
    mock_ingest.assert_called_once()
    kwargs = mock_ingest.call_args[1]
    chunks = kwargs.get("chunks")
    assert chunks is not None and len(chunks) >= 1        # 关键：chunks 真传进去了
    assert all(isinstance(chunk, dict) for chunk in chunks)
    text = kwargs["text"]
    for chunk in chunks:
        assert text[chunk["source_start"]:chunk["source_end"]] == chunk["text"]
        assert chunk["location"]["type"] == "sheet_row"
        assert chunk["location"]["sheet"] == "技术选型"
    joined = "\n".join(chunk["text"] for chunk in chunks)
    assert "【章节】技术选型" in joined                     # 工作表名作上下文（caption==heading 不重复）
    assert "类别 | 方案" in joined                          # 首行作表头
    assert "数据库 | MySQL" in joined and "缓存 | Redis" in joined


@patch("app.services.xlsx_extract.extract_xlsx_segments")
def test_import_xlsx_limit_returns_413_without_echoing_parser_details(mock_extract, client):
    from app.services.xlsx_extract import XlsxLimitError

    mock_extract.side_effect = XlsxLimitError("secret-cell-content")
    response = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("limited.xlsx", b"fixture", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == status.HTTP_413_CONTENT_TOO_LARGE
    assert response.json()["detail"] == "spreadsheet exceeds parser safety limits"
    assert "secret-cell-content" not in response.text


@patch("app.services.xlsx_extract.extract_xlsx_segments")
def test_import_xlsx_parse_error_is_clear_without_echoing_exception(mock_extract, client):
    mock_extract.side_effect = ValueError("secret-cell-content")
    response = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("invalid.xlsx", b"fixture", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == status.HTTP_400_BAD_REQUEST
    assert response.json()["detail"] == "Invalid spreadsheet format"
    assert "secret-cell-content" not in response.text

@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_import_json_file_list(mock_ingest, client):
    mock_doc = AsyncMock()
    mock_doc.id = "00000000-0000-0000-0000-000000000004"
    mock_doc.status = "pending"
    mock_job = AsyncMock()
    mock_ingest.return_value = (mock_doc, mock_job, 2, False)

    json_data = [
        {"title": "Doc 1", "text": "First doc text", "external_id": "ext1"},
        {"title": "Doc 2", "text": "Second doc text", "external_id": "ext2"}
    ]
    file_content = json.dumps(json_data).encode("utf-8")
    files = {"file": ("docs.json", file_content, "application/json")}
    response = client.post("/libraries/testlib/import-file", files=files)
    
    assert response.status_code == status.HTTP_201_CREATED
    data = response.json()
    assert data["status"] == "success"
    assert data["imported_count"] == 2
    assert data["documents"][0]["title"] == "Doc 1"
    assert data["documents"][1]["title"] == "Doc 2"
    assert mock_ingest.call_count == 2

@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_import_csv_file(mock_ingest, client):
    mock_doc = AsyncMock()
    mock_doc.id = "00000000-0000-0000-0000-000000000005"
    mock_doc.status = "pending"
    mock_job = AsyncMock()
    mock_ingest.return_value = (mock_doc, mock_job, 3, False)

    csv_data = "title,text,external_id\nTitle A,Text A,extA\nTitle B,Text B,extB"
    file_content = csv_data.encode("utf-8")
    files = {"file": ("data.csv", file_content, "text/csv")}
    response = client.post("/libraries/testlib/import-file", files=files)
    
    assert response.status_code == status.HTTP_201_CREATED
    data = response.json()
    assert data["status"] == "success"
    assert data["imported_count"] == 2
    assert data["documents"][0]["title"] == "Title A"
    assert data["documents"][1]["title"] == "Title B"
    assert mock_ingest.call_count == 2


@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_import_json_fanout_limit_returns_stable_413(mock_ingest, client, monkeypatch):
    from app.services import import_parsing

    monkeypatch.setattr(import_parsing, "MAX_FANOUT_DOCUMENTS", 1)
    payload = json.dumps([
        {"title": "secret one", "text": "private one"},
        {"title": "secret two", "text": "private two"},
    ]).encode("utf-8")

    response = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("many.json", payload, "application/json")},
    )

    assert response.status_code == status.HTTP_413_CONTENT_TOO_LARGE
    assert response.json()["detail"] == import_parsing.RESOURCE_LIMIT_ERROR
    assert "secret" not in response.text
    mock_ingest.assert_not_awaited()


@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_import_json_text_budget_returns_stable_413(mock_ingest, client, monkeypatch):
    from app.services import import_parsing

    monkeypatch.setattr(import_parsing, "MAX_NORMALIZED_TEXT_CHARS", 4)
    response = client.post(
        "/libraries/testlib/import-file",
        files={
            "file": (
                "large-text.json",
                json.dumps({"title": "secret", "text": "private text"}).encode("utf-8"),
                "application/json",
            )
        },
    )

    assert response.status_code == status.HTTP_413_CONTENT_TOO_LARGE
    assert response.json()["detail"] == import_parsing.RESOURCE_LIMIT_ERROR
    assert "secret" not in response.text
    mock_ingest.assert_not_awaited()


@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_import_json_input_bytes_rejects_before_json_loads(mock_ingest, client, monkeypatch):
    from app.api import documents
    from app.services import import_parsing

    payload = b'{"title":"secret","text":"private"}'
    monkeypatch.setattr(import_parsing, "MAX_JSON_INPUT_BYTES", len(payload) - 1)
    load_calls = 0

    def fail_json_loads(*_args, **_kwargs):
        nonlocal load_calls
        load_calls += 1
        raise AssertionError("json.loads must not run")

    monkeypatch.setattr(documents.json, "loads", fail_json_loads)
    response = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("large-input.json", payload, "application/json")},
    )

    assert response.status_code == status.HTTP_413_CONTENT_TOO_LARGE
    assert import_parsing.RESOURCE_LIMIT_ERROR in response.text
    assert "secret" not in response.text
    assert load_calls == 0
    mock_ingest.assert_not_awaited()


@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_import_csv_row_and_cell_limits_return_stable_413(mock_ingest, client, monkeypatch):
    from app.services import import_parsing

    monkeypatch.setattr(import_parsing, "MAX_CSV_ROWS", 2)
    payload = "text,title\nsecret one,one\nsecret two,two\n".encode("utf-8")

    response = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("many.csv", payload, "text/csv")},
    )

    assert response.status_code == status.HTTP_413_CONTENT_TOO_LARGE
    assert response.json()["detail"] == import_parsing.RESOURCE_LIMIT_ERROR
    assert "secret" not in response.text
    mock_ingest.assert_not_awaited()

    monkeypatch.setattr(import_parsing, "MAX_CSV_ROWS", 100_000)
    monkeypatch.setattr(import_parsing, "MAX_CSV_CELLS", 2)
    response = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("wide.csv", payload, "text/csv")},
    )
    assert response.status_code == status.HTTP_413_CONTENT_TOO_LARGE
    assert response.json()["detail"] == import_parsing.RESOURCE_LIMIT_ERROR


# ── 3A：API Key（无 cookie）上传闭环 ──────────────────────────────────────

def test_import_file_with_api_key_no_cookie():
    """API Key (Bearer) 能调用上传接口，不依赖浏览器 cookie；返回含 job_id。"""
    mock_user.is_superuser = True
    app.dependency_overrides[get_db] = override_db
    doc = MagicMock()
    doc.id = uuid.uuid4()
    doc.status = "pending"
    job = MagicMock()
    job.id = uuid.uuid4()
    with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml, \
         patch("app.auth.api_key.APIKeyStrategy.read_token", new_callable=AsyncMock) as rt, \
         patch("app.services.ingest.ingest_text", new_callable=AsyncMock) as mi:
        ml.return_value = mock_library
        rt.return_value = mock_user          # Bearer key → 解析出用户（不连 DB）
        mi.return_value = (doc, job, 1, False)
        client = TestClient(app)
        resp = client.post(
            "/libraries/testlib/import-file",
            files={"file": ("a.txt", b"hello world", "text/plain")},
            headers={"Authorization": "Bearer vk_fake_key"},   # 仅 Bearer，无 cookie
        )
        assert resp.status_code == status.HTTP_201_CREATED
        data = resp.json()
        assert data["imported_count"] == 1
        assert data["documents"][0]["job_id"] == str(job.id)
        rt.assert_awaited()                  # 证明确实走了 API Key 后端
    app.dependency_overrides.clear()


def test_import_file_api_key_without_insert_forbidden():
    """API Key 用户没有 insert 权限 → 上传返回 403。"""
    mock_user.is_superuser = False
    app.dependency_overrides[get_db] = override_db
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml, \
             patch("app.auth.api_key.APIKeyStrategy.read_token", new_callable=AsyncMock) as rt, \
             patch("app.deps.has_permission") as hp:
            ml.return_value = mock_library
            rt.return_value = mock_user
            hp.return_value = False
            client = TestClient(app)
            resp = client.post(
                "/libraries/testlib/import-file",
                files={"file": ("a.txt", b"hi", "text/plain")},
                headers={"Authorization": "Bearer vk_fake_key"},
            )
            assert resp.status_code == status.HTTP_403_FORBIDDEN
    finally:
        mock_user.is_superuser = True
        app.dependency_overrides.clear()


# ── 3A：文件上传 external_id 自动 upsert ──────────────────────────────────

def test_import_file_external_id_upsert():
    """带 external_id 重复上传 → 命中库内已存在文档走 reingest 覆盖，不新建。"""
    existing = Document(
        id=uuid.uuid4(), library_id=mock_library.id,
        external_id="ext-1", status="pending", content_hash="oldhash",
    )
    db = make_db_mock(existing=existing)

    async def _ov_db():
        return db

    app.dependency_overrides[current_active_user] = override_user
    app.dependency_overrides[get_db] = _ov_db
    job = MagicMock()
    job.id = uuid.uuid4()
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml, \
             patch("app.services.ingest.reingest_document", new_callable=AsyncMock) as rr, \
             patch("app.services.ingest.ingest_text", new_callable=AsyncMock) as it:
            ml.return_value = mock_library
            rr.return_value = (job, 3, False)   # changed=False → 不触发 qdrant 清理
            client = TestClient(app)
            resp = client.post(
                "/libraries/testlib/import-file",
                files={"file": ("note.txt", b"new content", "text/plain")},
                data={"external_id": "ext-1"},
            )
            assert resp.status_code == status.HTTP_201_CREATED
            d = resp.json()["documents"][0]
            assert d["document_id"] == str(existing.id)
            assert d["chunk_count"] == 3
            assert d["external_id"] == "ext-1"
            assert d["job_id"] == str(job.id)
            rr.assert_awaited_once()
            assert rr.call_args.kwargs["force"] is False     # upsert 用 no-op 判定
            it.assert_not_called()                           # 命中 upsert，不应新建
    finally:
        app.dependency_overrides.clear()


# ── 3A/3B：库级任务状态查询（普通用户）─────────────────────────────────────

def _job(**kw) -> EmbeddingJob:
    kw.setdefault("id", uuid.uuid4())
    kw.setdefault("library_id", mock_library.id)
    kw.setdefault("document_id", uuid.uuid4())
    kw.setdefault("status", "done")
    kw.setdefault("attempt_count", 0)
    kw.setdefault("document_revision", 1)
    kw.setdefault("created_at", datetime.now(timezone.utc))
    return EmbeddingJob(**kw)


def _client_with_db(db):
    app.dependency_overrides[current_active_user] = override_user

    async def _ov_db():
        return db

    app.dependency_overrides[get_db] = _ov_db
    return TestClient(app)


def test_get_job_status():
    mock_user.is_superuser = True
    job = _job(status="done", attempt_count=1)
    db = make_db_mock(get_return=job)
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/jobs/{job.id}")
            assert resp.status_code == status.HTTP_200_OK
            body = resp.json()
            assert body["id"] == str(job.id)
            assert body["status"] == "done"
            assert body["attempt_count"] == 1
    finally:
        app.dependency_overrides.clear()


def test_get_job_status_wrong_library_404():
    """job 属于别的库 → 404（库级隔离）。"""
    mock_user.is_superuser = True
    job = _job(library_id=uuid.uuid4())     # 与 mock_library.id 不同
    db = make_db_mock(get_return=job)
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/jobs/{job.id}")
            assert resp.status_code == status.HTTP_404_NOT_FOUND
    finally:
        app.dependency_overrides.clear()


def test_list_document_jobs():
    mock_user.is_superuser = True
    doc = Document(id=uuid.uuid4(), library_id=mock_library.id, status="ready", content_hash="h")
    job = _job(document_id=doc.id, status="processing")
    db = make_db_mock(get_return=doc, scalar_list=[job])
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/jobs")
            assert resp.status_code == status.HTTP_200_OK
            rows = resp.json()
            assert len(rows) == 1
            assert rows[0]["status"] == "processing"
            assert rows[0]["document_id"] == str(doc.id)
    finally:
        app.dependency_overrides.clear()


# ── #2：库级统计补 done_jobs / total_jobs ─────────────────────────────────

def test_library_stats_includes_done_and_total():
    mock_user.is_superuser = True
    job_rows = [("pending", 1), ("processing", 2), ("done", 3), ("failed", 4)]
    db = AsyncMock()
    # doc_count / chunk_count 用 scalar_one；job 聚合用 .all()——同一 result mock 都能满足
    db.execute = AsyncMock(return_value=MagicMock(
        scalar_one=MagicMock(return_value=7),
        all=MagicMock(return_value=job_rows),
    ))
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get("/libraries/testlib/stats")
            assert resp.status_code == status.HTTP_200_OK
            s = resp.json()
            assert s["document_count"] == 7 and s["chunk_count"] == 7
            assert s["pending_jobs"] == 1 and s["processing_jobs"] == 2
            assert s["done_jobs"] == 3 and s["failed_jobs"] == 4
            assert s["total_jobs"] == 10
    finally:
        app.dependency_overrides.clear()


# ── #3：docx 表格感知开关接入上传 ─────────────────────────────────────────

def test_import_docx_table_aware_passes_structured_segment_chunks():
    """库开 docx_table_aware → docx 走 segments，structured chunks 传进 ingest_text。"""
    mock_user.is_superuser = True
    app.dependency_overrides[current_active_user] = override_user
    app.dependency_overrides[get_db] = override_db
    mock_library.docx_table_aware = True
    doc = MagicMock()
    doc.id = uuid.uuid4()
    doc.status = "pending"
    segs = [{"kind": "table", "heading": "技术选型", "caption": "技术选型",
             "header": "类别 | 方案", "rows": ["类别 | 方案", "数据库 | MySQL"]}]
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml, \
              patch("app.services.docx_extract.extract_docx_segments", return_value=segs) as seg, \
              patch("app.services.ingest.ingest_text", new_callable=AsyncMock) as mi:
            ml.return_value = mock_library
            mi.return_value = (doc, MagicMock(id=uuid.uuid4()), 1, False)
            client = TestClient(app)
            resp = client.post(
                "/libraries/testlib/import-file",
                files={"file": ("台账.docx", b"PK-fake-docx-bytes",
                                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            assert resp.status_code == status.HTTP_201_CREATED
            seg.assert_called_once()
            chunks = mi.call_args.kwargs.get("chunks")
            assert chunks is not None and all(isinstance(chunk, dict) for chunk in chunks)
            text = mi.call_args.kwargs["text"]
            assert "数据库 | MySQL" in "\n".join(chunk["text"] for chunk in chunks)
            for chunk in chunks:
                assert text[chunk["source_start"]:chunk["source_end"]] == chunk["text"]
    finally:
        mock_library.docx_table_aware = None
        app.dependency_overrides.clear()


# ── #4：旧 Office 格式不在严格白名单内 ─────────────────────────────────────

def test_import_xls_rejected(client):
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("旧台账.xls", b"\xd0\xcf\x11\xe0fake-xls", "application/vnd.ms-excel")},
    )
    assert resp.status_code == status.HTTP_415_UNSUPPORTED_MEDIA_TYPE


def test_import_doc_rejected(client):
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("旧合同.doc", b"\xd0\xcf\x11\xe0fake-doc", "application/msword")},
    )
    assert resp.status_code == status.HTTP_415_UNSUPPORTED_MEDIA_TYPE


# ── docs/23：PDF 文字层 / 扫描页 OCR 接入上传 ─────────────────────────────────
from app.services.pdf_extract import PdfExtractError, PdfOcrUnavailableError  # noqa: E402


@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
@patch("app.services.pdf_extract.build_pdf_source")
def test_import_text_pdf_ocr_off_ingests(mock_extract, mock_ingest, client):
    """OCR 关 + 文字 PDF：提取出的带页码正文照常入库，splitter=text，ocr_enabled=False。"""
    mock_extract.return_value = {
        "normalized_text": "【第 1 页】\n文字版PDF内容",
        "chunks": [{
            "text": "【第 1 页】\n文字版PDF内容",
            "source_start": 0,
            "source_end": 16,
            "location": {"type": "page", "page": 1},
        }],
    }
    doc = MagicMock()
    doc.id = uuid.uuid4()
    doc.status = "pending"
    mock_ingest.return_value = (doc, MagicMock(id=uuid.uuid4()), 1, False)

    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("doc.pdf", b"%PDF-1.4 fake", "application/pdf")},
    )
    assert resp.status_code == status.HTTP_201_CREATED
    kw = mock_ingest.call_args[1]
    assert kw["text"] == "【第 1 页】\n文字版PDF内容"
    assert kw["splitter"] == "text"
    assert kw["chunks"][0]["location"] == {"type": "page", "page": 1}
    assert mock_extract.call_args.kwargs["ocr_enabled"] is False   # 库未开 + 全局默认关


@patch("app.services.pdf_extract.build_pdf_source")
def test_import_scanned_pdf_ocr_off_returns_enable_hint(mock_extract, client):
    """OCR 关 + 扫描 PDF：服务抛 PdfExtractError → 400，提示去开启 OCR。"""
    mock_extract.side_effect = PdfExtractError(
        "PDF 无可提取文本；如为扫描件，请在知识库开启图片 OCR"
    )
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("scan.pdf", b"%PDF-1.4 fake", "application/pdf")},
    )
    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert "OCR" in resp.json()["detail"]


@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
@patch("app.services.pdf_extract.build_pdf_source")
def test_import_pdf_ocr_on_passes_flag_and_params(mock_extract, mock_ingest, client, monkeypatch):
    """库开 ocr_enabled：把库级开关、ocr_image 回调与三个 pdf_ocr_* 参数传给服务。"""
    monkeypatch.setattr(mock_library, "ocr_enabled", True)
    sentinel = object()
    monkeypatch.setattr("app.services.ocr.is_available", lambda: True)
    monkeypatch.setattr("app.services.ocr.ocr_image", sentinel)
    monkeypatch.setattr(settings, "pdf_ocr_min_text_chars", 20)
    monkeypatch.setattr(settings, "pdf_ocr_render_dpi", 200)
    monkeypatch.setattr(settings, "pdf_ocr_max_pages", 50)
    mock_extract.return_value = {
        "normalized_text": "【第 1 页】\nx",
        "chunks": [{"text": "【第 1 页】\nx", "source_start": 0, "source_end": 9, "location": {"type": "page", "page": 1}}],
    }
    doc = MagicMock()
    doc.id = uuid.uuid4()
    doc.status = "pending"
    mock_ingest.return_value = (doc, MagicMock(id=uuid.uuid4()), 1, False)

    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("scan.pdf", b"%PDF-1.4 fake", "application/pdf")},
    )
    assert resp.status_code == status.HTTP_201_CREATED
    kw = mock_extract.call_args.kwargs
    assert kw["ocr_enabled"] is True
    assert kw["ocr"] is sentinel
    assert kw["min_text_chars"] == 20
    assert kw["render_dpi"] == 200
    assert kw["max_ocr_pages"] == 50


@patch("app.services.pdf_extract.build_pdf_source")
def test_import_pdf_ocr_unavailable_returns_install_hint(mock_extract, client, monkeypatch):
    """OCR 开但引擎缺：服务抛 PdfOcrUnavailableError → 400，提示安装 .[ocr]。"""
    monkeypatch.setattr(mock_library, "ocr_enabled", True)
    monkeypatch.setattr("app.services.ocr.is_available", lambda: True)
    mock_extract.side_effect = PdfOcrUnavailableError(
        'PDF 含扫描页需要 OCR，但 OCR 依赖未安装；请执行 pip install -e ".[ocr]"'
    )
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("scan.pdf", b"%PDF-1.4 fake", "application/pdf")},
    )
    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert ".[ocr]" in resp.json()["detail"]


# ───────── 新增 / 替换上传：operation 字段 + 按 document ID 替换闭环 ─────────

def _doc(*, external_id="ext-keep", library_id=None, deleted_at=None, doc_metadata=None):
    """构造一个替换目标文档（_replace_document 只读这几个属性）。"""
    d = MagicMock()
    d.id = uuid.uuid4()
    d.library_id = library_id if library_id is not None else mock_library.id
    d.external_id = external_id
    d.deleted_at = deleted_at
    d.status = "pending"
    d.doc_metadata = doc_metadata
    return d


def _lib(**over):
    base = dict(
        id=mock_library.id, slug="testlib", name="Test", qdrant_collection="testlib_col",
        embedding_model="bge-m3", embedding_base_url="http://m", embedding_dim=1024,
        chunk_size=1000, chunk_overlap=120, lifecycle_mode="managed", index_state="ready",
    )
    base.update(over)
    return Library(**base)


@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_add_mode_returns_created(mock_ingest, client):
    """新增模式：新建文档 → operation=created。"""
    doc = MagicMock()
    doc.id = uuid.uuid4()
    doc.status = "pending"
    mock_ingest.return_value = (doc, MagicMock(id="j"), 1, False)
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("a.txt", b"hello world", "text/plain")},
    )
    assert resp.status_code == status.HTTP_201_CREATED
    assert resp.json()["documents"][0]["operation"] == "created"


@patch("app.services.ingest.ingest_text", new_callable=AsyncMock)
def test_add_mode_identical_returns_unchanged_with_real_doc_info(mock_ingest, client):
    """新增模式：内容命中去重（was_existing=True）→ operation=unchanged；
    回显数据库里真实保留的旧文件名/external_id，而非本次上传的文件名。"""
    doc = MagicMock()
    doc.id = uuid.uuid4()
    doc.status = "pending"
    doc.title = "original_name.txt"      # 库里真实保留的旧文件名
    doc.external_id = "orig-ext"
    mock_ingest.return_value = (doc, MagicMock(id="j"), 1, True)
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("newly_uploaded.txt", b"dup content", "text/plain")},  # 故意用不同文件名
    )
    assert resp.status_code == status.HTTP_201_CREATED
    body = resp.json()["documents"][0]
    assert body["operation"] == "unchanged"
    assert body["title"] == "original_name.txt"   # 真实旧文件名，不是 newly_uploaded.txt
    assert body["external_id"] == "orig-ext"


@patch("app.services.ingest.reingest_document", new_callable=AsyncMock)
def test_replace_keeps_id_external_id_and_forces_reingest(mock_re, client):
    """替换：document_id 不变、external_id 保留、走 reingest(force=True)、标题更新为新文件名。"""
    target = _doc(external_id="keep-me")
    app.dependency_overrides[get_db] = lambda: make_db_mock(existing=target)
    mock_re.return_value = (MagicMock(id="job-x"), 3, True)
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("new.txt", b"new body", "text/plain")},
        data={"replace_document_id": str(target.id)},
    )
    assert resp.status_code == status.HTTP_201_CREATED
    body = resp.json()["documents"][0]
    assert body["document_id"] == str(target.id)
    assert body["external_id"] == "keep-me"
    assert body["operation"] == "updated"
    assert body["chunk_count"] == 3
    # force=True 保证 revision+1 / supersede 旧 job / cleanup outbox（不另写更新逻辑）
    assert mock_re.await_args.kwargs["force"] is True
    assert mock_re.await_args.kwargs["title"] == "new.txt"


@patch("app.services.ingest.reingest_document", new_callable=AsyncMock)
def test_replace_preserves_metadata_when_file_has_none(mock_re, client):
    """替换：上传文件未带 metadata 时，保留目标文档原 metadata（不清空作者/分类等）。"""
    target = _doc(doc_metadata={"author": "alice", "category": "legal"})
    app.dependency_overrides[get_db] = lambda: make_db_mock(existing=target)
    mock_re.return_value = (MagicMock(id="job-z"), 1, True)
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("new.txt", b"new body without metadata", "text/plain")},
        data={"replace_document_id": str(target.id)},
    )
    assert resp.status_code == status.HTTP_201_CREATED
    # 传给 reingest 的 metadata 应是目标原 metadata，而非 None
    assert mock_re.await_args.kwargs["metadata"] == {"author": "alice", "category": "legal"}


@patch("app.services.ingest.reingest_document", new_callable=AsyncMock)
def test_replace_applies_management_scope_to_new_revision(mock_re, client):
    target = _doc()
    target.visibility_scope = None
    target.security_level = None
    app.dependency_overrides[get_db] = lambda: make_db_mock(existing=target)
    mock_re.return_value = (MagicMock(id="job-scope"), 1, True)

    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("new.txt", b"managed content", "text/plain")},
        data={
            "replace_document_id": str(target.id),
            "visibility_scope": " internal-users ",
            "security_level": " internal ",
        },
    )

    assert resp.status_code == status.HTTP_201_CREATED
    assert target.visibility_scope == "internal-users"
    assert target.security_level == "internal"
    assert mock_re.await_args.kwargs["document"] is target


def test_import_rejects_blank_management_scope(client):
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("new.txt", b"managed content", "text/plain")},
        data={"security_level": "   "},
    )

    assert resp.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    assert resp.json()["detail"] == "security_level must not be blank"


@patch("app.services.ingest.reingest_document", new_callable=AsyncMock)
def test_replace_duplicate_content_returns_409(mock_re, client):
    """替换：新内容与同库另一篇无 external_id 文档相同 → 撞唯一约束 → 回滚并 409（不 500）。"""
    target = _doc()
    db = make_db_mock(existing=target)
    app.dependency_overrides[get_db] = lambda: db
    mock_re.side_effect = IntegrityError("stmt", {}, Exception("duplicate content_hash"))
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("dup.txt", b"identical to another doc", "text/plain")},
        data={"replace_document_id": str(target.id)},
    )
    assert resp.status_code == status.HTTP_409_CONFLICT
    db.rollback.assert_awaited()


@patch("app.services.ingest.reingest_document", new_callable=AsyncMock)
def test_replace_doc_without_external_id(mock_re, client):
    """替换：目标没有 external_id 也能按 document ID 替换。"""
    target = _doc(external_id=None)
    app.dependency_overrides[get_db] = lambda: make_db_mock(existing=target)
    mock_re.return_value = (MagicMock(id="job-y"), 1, True)
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("x.txt", b"abc def ghi", "text/plain")},
        data={"replace_document_id": str(target.id)},
    )
    assert resp.status_code == status.HTTP_201_CREATED
    body = resp.json()["documents"][0]
    assert body["document_id"] == str(target.id)
    assert body["external_id"] is None
    assert body["operation"] == "updated"


def test_replace_target_not_found_404(client):
    """替换：目标不存在 → 404。"""
    app.dependency_overrides[get_db] = lambda: make_db_mock(existing=None)
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("x.txt", b"abc def", "text/plain")},
        data={"replace_document_id": str(uuid.uuid4())},
    )
    assert resp.status_code == status.HTTP_404_NOT_FOUND


def test_replace_target_other_library_404(client):
    """替换：目标属于别的库 → 404。"""
    target = _doc(library_id=uuid.uuid4())
    app.dependency_overrides[get_db] = lambda: make_db_mock(existing=target)
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("x.txt", b"abc def", "text/plain")},
        data={"replace_document_id": str(target.id)},
    )
    assert resp.status_code == status.HTTP_404_NOT_FOUND


def test_replace_target_deleted_404(client):
    """替换：目标已删除 → 404，不能复活 tombstone 文档。"""
    target = _doc(deleted_at=datetime.now(timezone.utc))
    app.dependency_overrides[get_db] = lambda: make_db_mock(existing=target)
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("x.txt", b"abc def", "text/plain")},
        data={"replace_document_id": str(target.id)},
    )
    assert resp.status_code == status.HTTP_404_NOT_FOUND


def test_replace_invalid_document_id_returns_422(client):
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("x.txt", b"abc def", "text/plain")},
        data={"replace_document_id": "not-a-uuid"},
    )
    assert resp.status_code == 422


def test_replace_multi_doc_file_returns_400(client):
    """替换：多文档文件（JSON 数组）→ 400，不允许一对多。"""
    payload = json.dumps([{"text": "doc one"}, {"text": "doc two"}]).encode("utf-8")
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("multi.json", payload, "application/json")},
        data={"replace_document_id": str(uuid.uuid4())},
    )
    assert resp.status_code == status.HTTP_400_BAD_REQUEST


def test_replace_on_external_library_rejected(client):
    """替换：external 库仍按现有规则拒绝写入 → 409。"""
    target = _doc()
    app.dependency_overrides[get_db] = lambda: make_db_mock(existing=target, library=_lib(lifecycle_mode="external"))
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("x.txt", b"abc def", "text/plain")},
        data={"replace_document_id": str(target.id)},
    )
    assert resp.status_code == status.HTTP_409_CONFLICT


def test_replace_on_rebuilding_library_rejected(client):
    """替换：rebuilding 库仍按现有规则拒绝写入 → 503。"""
    target = _doc()
    app.dependency_overrides[get_db] = lambda: make_db_mock(existing=target, library=_lib(index_state="rebuilding"))
    resp = client.post(
        "/libraries/testlib/import-file",
        files={"file": ("x.txt", b"abc def", "text/plain")},
        data={"replace_document_id": str(target.id)},
    )
    assert resp.status_code == status.HTTP_503_SERVICE_UNAVAILABLE


@patch("app.deps.has_permission")
def test_import_file_requires_insert_permission(mock_has_perm):
    mock_user.is_superuser = False
    mock_has_perm.return_value = False
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(make_db_mock()).post(
                "/libraries/testlib/import-file",
                files={"file": ("x.txt", b"abc def", "text/plain")},
            )
            assert resp.status_code == status.HTTP_403_FORBIDDEN
    finally:
        mock_user.is_superuser = True
        app.dependency_overrides.clear()


@patch("app.deps.has_permission")
def test_replace_document_requires_insert_permission(mock_has_perm):
    mock_user.is_superuser = False
    mock_has_perm.return_value = False
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(make_db_mock()).post(
                "/libraries/testlib/import-file",
                files={"file": ("x.txt", b"abc def", "text/plain")},
                data={"replace_document_id": str(uuid.uuid4())},
            )
            assert resp.status_code == status.HTTP_403_FORBIDDEN
    finally:
        mock_user.is_superuser = True
        app.dependency_overrides.clear()


def _get_source_response_with_span_hash(stored_hash):
    mock_user.is_superuser = True
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="source.txt",
        content_hash="h",
        current_revision=2,
        status="ready",
    )
    chunk = Chunk(
        id=uuid.uuid4(),
        document_id=doc.id,
        library_id=mock_library.id,
        seq=1,
        text="needle text",
        token_count=11,
        chunk_metadata={
            "source_start": 7,
            "source_end": 18,
            "location": {"type": "line", "start_line": 2, "end_line": 2},
            "source_revision": 2,
            **({} if stored_hash is None else {"source_span_hash": stored_hash}),
        },
    )
    source = DocumentSource(
        document_id=doc.id,
        revision=2,
        file_name="source.txt",
        file_type="txt",
        normalized_text="before needle text after",
    )
    db = AsyncMock()

    async def _get(model, ident):
        if model is Document:
            return doc
        if model is Chunk:
            return chunk
        if model is DocumentSource:
            return source
        return None

    db.get = _get
    with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
        ml.return_value = mock_library
        resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/source?chunk_id={chunk.id}")
        assert resp.status_code == status.HTTP_200_OK
        return resp.json(), chunk


def test_get_document_source_legacy_when_span_hash_missing():
    try:
        data, chunk = _get_source_response_with_span_hash(None)
        assert data["legacy"] is True
        assert data["fallback_chunk"] == chunk.text
    finally:
        app.dependency_overrides.clear()


def test_get_document_source_legacy_when_span_hash_wrong():
    try:
        data, chunk = _get_source_response_with_span_hash("0" * 16)
        assert data["legacy"] is True
        assert data["fallback_chunk"] == chunk.text
    finally:
        app.dependency_overrides.clear()


def test_get_document_source_precise_window_when_span_hash_matches():
    span_hash = hashlib.sha256("needle text".encode("utf-8")).hexdigest()[:16]
    try:
        data, chunk = _get_source_response_with_span_hash(span_hash)
        assert data["legacy"] is False
        assert data["document_title"] == "source.txt"
        assert data["file_type"] == "txt"
        assert data["chunk_seq"] == 1
        local_start = data["source_start"] - data["window_start"]
        local_end = data["source_end"] - data["window_start"]
        assert data["text_window"][local_start:local_end] == chunk.text
    finally:
        app.dependency_overrides.clear()


def test_get_document_source_legacy_when_any_source_range_hash_wrong():
    mock_user.is_superuser = True
    source_text = "first middle second"
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="source.md",
        content_hash="h",
        current_revision=2,
        status="ready",
    )
    ranges = [
        {"start": 0, "end": 5, "hash": hashlib.sha256(b"first").hexdigest()[:16]},
        {"start": 13, "end": 19, "hash": "0" * 16},
    ]
    chunk = Chunk(
        id=uuid.uuid4(),
        document_id=doc.id,
        library_id=mock_library.id,
        seq=1,
        text="first  \nsecond",
        token_count=14,
        chunk_metadata={
            "source_start": 0,
            "source_end": 19,
            "source_ranges": ranges,
            "location": {"type": "line", "start_line": 1, "end_line": 1},
            "source_revision": 2,
        },
    )
    source = DocumentSource(
        document_id=doc.id,
        revision=2,
        file_name="source.md",
        file_type="md",
        normalized_text=source_text,
    )
    db = AsyncMock()

    async def _get(model, ident):
        if model is Document:
            return doc
        if model is Chunk:
            return chunk
        if model is DocumentSource:
            return source
        return None

    db.get = _get
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/source?chunk_id={chunk.id}")
            assert resp.status_code == status.HTTP_200_OK
            data = resp.json()
            assert data["legacy"] is True
            assert data["fallback_chunk"] == "first  \nsecond"
    finally:
        app.dependency_overrides.clear()


def test_get_document_source_precise_returns_source_ranges_when_hashes_match():
    mock_user.is_superuser = True
    source_text = "first middle second"
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="source.md",
        content_hash="h",
        current_revision=2,
        status="ready",
    )
    ranges = [
        {"start": 0, "end": 5, "hash": hashlib.sha256(b"first").hexdigest()[:16]},
        {"start": 13, "end": 19, "hash": hashlib.sha256(b"second").hexdigest()[:16]},
    ]
    chunk = Chunk(
        id=uuid.uuid4(),
        document_id=doc.id,
        library_id=mock_library.id,
        seq=1,
        text="first  \nsecond",
        token_count=14,
        chunk_metadata={
            "source_start": 0,
            "source_end": 19,
            "source_ranges": ranges,
            "location": {"type": "line", "start_line": 1, "end_line": 1},
            "source_revision": 2,
        },
    )
    source = DocumentSource(
        document_id=doc.id,
        revision=2,
        file_name="source.md",
        file_type="md",
        normalized_text=source_text,
    )
    db = AsyncMock()

    async def _get(model, ident):
        if model is Document:
            return doc
        if model is Chunk:
            return chunk
        if model is DocumentSource:
            return source
        return None

    db.get = _get
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/source?chunk_id={chunk.id}")
            assert resp.status_code == status.HTTP_200_OK
            data = resp.json()
            assert data["legacy"] is False
            assert [(r["start"], r["end"]) for r in data["source_ranges"]] == [(0, 5), (13, 19)]
    finally:
        app.dependency_overrides.clear()


def test_get_document_source_legacy_without_source_does_not_fabricate_offsets():
    mock_user.is_superuser = True
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="legacy.txt",
        content_hash="h",
        current_revision=1,
        status="ready",
    )
    chunk = Chunk(
        id=uuid.uuid4(),
        document_id=doc.id,
        library_id=mock_library.id,
        seq=0,
        text="legacy chunk",
        token_count=12,
        chunk_metadata=None,
    )
    db = AsyncMock()

    async def _get(model, ident):
        if model is Document:
            return doc
        if model is Chunk:
            return chunk
        if model is DocumentSource:
            return None
        return None

    db.get = _get
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/source?chunk_id={chunk.id}")
            assert resp.status_code == status.HTTP_200_OK
            data = resp.json()
            assert data["legacy"] is True
            assert data["fallback_chunk"] == "legacy chunk"
            assert data["source_start"] is None
            assert data["source_end"] is None
    finally:
        app.dependency_overrides.clear()


def test_get_document_full_source_returns_normalized_text():
    mock_user.is_superuser = True
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="full-source.txt",
        content_hash="h",
        current_revision=3,
        status="ready",
    )
    now = datetime.now(timezone.utc)
    source = DocumentSource(
        document_id=doc.id,
        revision=3,
        file_name="full-source.txt",
        file_type="txt",
        normalized_text="第一行\n第二行 needle",
        created_at=now,
        updated_at=now,
    )
    db = AsyncMock()

    async def _get(model, ident):
        if model is Document:
            return doc
        if model is DocumentSource:
            return source
        return None

    db.get = _get
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/source/full")
            assert resp.status_code == status.HTTP_200_OK
            data = resp.json()
            assert data["document_id"] == str(doc.id)
            assert data["document_title"] == "full-source.txt"
            assert data["file_name"] == "full-source.txt"
            assert data["file_type"] == "txt"
            assert data["revision"] == 3
            assert data["normalized_text"] == "第一行\n第二行 needle"
            assert data["text_length"] == len("第一行\n第二行 needle")
    finally:
        app.dependency_overrides.clear()


def test_get_document_full_source_without_snapshot_returns_clear_404():
    mock_user.is_superuser = True
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="legacy.txt",
        content_hash="h",
        current_revision=1,
        status="ready",
    )
    db = AsyncMock()

    async def _get(model, ident):
        if model is Document:
            return doc
        if model is DocumentSource:
            return None
        return None

    db.get = _get
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/source/full")
            assert resp.status_code == status.HTTP_404_NOT_FOUND
            assert "重新导入" in resp.json()["detail"]
    finally:
        app.dependency_overrides.clear()


@patch("app.deps.has_permission")
def test_get_document_full_source_requires_read_permission(mock_has_perm):
    mock_user.is_superuser = False
    mock_has_perm.return_value = False
    doc_id = uuid.uuid4()
    db = AsyncMock()
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc_id}/source/full")
            assert resp.status_code == status.HTTP_403_FORBIDDEN
    finally:
        mock_user.is_superuser = True
        app.dependency_overrides.clear()


def test_download_document_file_returns_persisted_bytes(monkeypatch):
    tmp_path = _safe_tmp_dir("document-file-download")
    mock_user.is_superuser = True
    monkeypatch.setattr(settings, "document_files_dir", str(tmp_path))
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="download.txt",
        content_hash="h",
        current_revision=1,
        status="ready",
    )
    stored = tmp_path / "lib" / "doc" / "file.txt"
    stored.parent.mkdir(parents=True)
    stored.write_bytes(b"download bytes")
    row = DocumentFile(
        document_id=doc.id,
        revision=1,
        file_name="download.txt",
        content_type="text/plain",
        storage_path=str(stored.relative_to(tmp_path)),
        size_bytes=len(b"download bytes"),
        sha256=hashlib.sha256(b"download bytes").hexdigest(),
    )
    db = AsyncMock()

    async def _get(model, ident):
        if model is Document:
            return doc
        if model is DocumentFile:
            return row
        return None

    db.get = _get
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/file")
            assert resp.status_code == status.HTTP_200_OK
            assert resp.content == b"download bytes"
            assert "download.txt" in resp.headers.get("content-disposition", "")
    finally:
        app.dependency_overrides.clear()


def test_download_document_file_missing_original_returns_clear_404():
    mock_user.is_superuser = True
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="missing.txt",
        content_hash="h",
        current_revision=1,
        status="ready",
    )
    db = AsyncMock()

    async def _get(model, ident):
        if model is Document:
            return doc
        if model is DocumentFile:
            return None
        return None

    db.get = _get
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/file")
            assert resp.status_code == status.HTTP_404_NOT_FOUND
            assert "缺少原始文件" in resp.json()["detail"]
    finally:
        app.dependency_overrides.clear()


def test_download_document_file_rejects_relative_escape(monkeypatch):
    tmp_path = _safe_tmp_dir("document-file-relative-escape")
    outside = tmp_path.parent / f"outside-{uuid.uuid4().hex}.txt"
    outside.write_bytes(b"outside")
    mock_user.is_superuser = True
    monkeypatch.setattr(settings, "document_files_dir", str(tmp_path))
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="escape.txt",
        content_hash="h",
        current_revision=1,
        status="ready",
    )
    row = DocumentFile(
        document_id=doc.id,
        revision=1,
        file_name="escape.txt",
        content_type="text/plain",
        storage_path=f"../{outside.name}",
        size_bytes=len(b"outside"),
        sha256=hashlib.sha256(b"outside").hexdigest(),
    )
    db = AsyncMock()

    async def _get(model, ident):
        if model is Document:
            return doc
        if model is DocumentFile:
            return row
        return None

    db.get = _get
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/file")
            assert resp.status_code == status.HTTP_404_NOT_FOUND
    finally:
        app.dependency_overrides.clear()


def test_download_document_file_rejects_absolute_path_outside_root(monkeypatch):
    tmp_path = _safe_tmp_dir("document-file-absolute-escape")
    outside = tmp_path.parent / f"absolute-{uuid.uuid4().hex}.txt"
    outside.write_bytes(b"outside")
    mock_user.is_superuser = True
    monkeypatch.setattr(settings, "document_files_dir", str(tmp_path))
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="absolute.txt",
        content_hash="h",
        current_revision=1,
        status="ready",
    )
    row = DocumentFile(
        document_id=doc.id,
        revision=1,
        file_name="absolute.txt",
        content_type="text/plain",
        storage_path=str(outside),
        size_bytes=len(b"outside"),
        sha256=hashlib.sha256(b"outside").hexdigest(),
    )
    db = AsyncMock()

    async def _get(model, ident):
        if model is Document:
            return doc
        if model is DocumentFile:
            return row
        return None

    db.get = _get
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc.id}/file")
            assert resp.status_code == status.HTTP_404_NOT_FOUND
    finally:
        app.dependency_overrides.clear()


@patch("app.deps.has_permission")
def test_download_document_file_requires_read_permission(mock_has_perm):
    mock_user.is_superuser = False
    mock_has_perm.return_value = False
    doc_id = uuid.uuid4()
    db = AsyncMock()
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).get(f"/libraries/testlib/documents/{doc_id}/file")
            assert resp.status_code == status.HTTP_403_FORBIDDEN
    finally:
        mock_user.is_superuser = True
        app.dependency_overrides.clear()


@patch("app.services.cleanup.enqueue_delete_document", new_callable=AsyncMock)
def test_delete_document_tombstones_supersedes_jobs_and_enqueues_cleanup(mock_enqueue):
    mock_user.is_superuser = True
    doc = Document(
        id=uuid.uuid4(),
        library_id=mock_library.id,
        title="delete-me.txt",
        content_hash="h",
        current_revision=2,
        status="ready",
    )
    db = make_db_mock(existing=doc)
    statements = []
    original_execute = db.execute

    async def _execute(stmt, *args, **kwargs):
        statements.append(str(stmt))
        return await original_execute(stmt, *args, **kwargs)

    db.execute = _execute
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(db).delete(f"/libraries/testlib/documents/{doc.id}")
            assert resp.status_code == status.HTTP_204_NO_CONTENT
            assert any("UPDATE documents" in stmt and "deleted_at" in stmt and "status" in stmt for stmt in statements)
            assert any("UPDATE embedding_jobs" in stmt and "status" in stmt and "finished_at" in stmt for stmt in statements)
            mock_enqueue.assert_awaited_once_with(db, mock_library, doc.id)
            db.commit.assert_awaited_once()
    finally:
        app.dependency_overrides.clear()


@patch("app.deps.has_permission")
def test_delete_document_requires_delete_permission(mock_has_perm):
    mock_user.is_superuser = False
    mock_has_perm.return_value = False
    try:
        with patch("app.deps.load_active_library", new_callable=AsyncMock) as ml:
            ml.return_value = mock_library
            resp = _client_with_db(make_db_mock()).delete(f"/libraries/testlib/documents/{uuid.uuid4()}")
            assert resp.status_code == status.HTTP_403_FORBIDDEN
    finally:
        mock_user.is_superuser = True
        app.dependency_overrides.clear()
